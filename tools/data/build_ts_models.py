# -*- coding: utf-8 -*-
"""Модели машин → вариант поиска «Вида объекта» (база, вид спецтехники, шаблон).

Задача пользователя 06.10.2026: «в поиск добавляем по классификации с выжимки,
что есть что: буквально чтобы можно было модель вбить, и нам уже выбралась
база. А ещё чтобы можно было и модули, если что, подтянуть». Источник — книга
пар docs/pary-poiska-ts.xlsx (tools/docs/build_pary_poiska_ts.py по
обезличенной книге): лист «Все сочетания» — сколько раз модель записана с
каким типом ТС и кузовом, лист «Тип и кузов» — к какому варианту поиска
привязана пара «тип + кузов» (привязку правит человек).

Модель получает варианты, к которым привязаны её сочетания, по убыванию числа
записей: ГАЗ 53 — самосвал, бортовой, фургон… Шаблон несёт модули, поэтому
выбор модели собирает и базу, и модули. Редкие варианты (меньше SHARE записей
модели) отбрасываются, но первый остаётся всегда; не больше LIMIT на модель.

    python tools/data/build_ts_models.py   # app/modules/vehicle/data/tsModels.js
"""
import collections
import io
import json
import os
import re
import subprocess
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
BOOK = os.path.join(ROOT, 'docs', 'pary-poiska-ts.xlsx')
OUT = os.path.join(ROOT, 'app', 'modules', 'vehicle', 'data', 'tsModels.js')
EMPTY = '—'
SHARE = 0.05
LIMIT = 4

# Подписи вариантов — как в книге пар (лист «Варианты»): «вид: название».
JS = r"""
const m = await import('./app/modules/vehicle/data/tsCatalog.js');
console.log(JSON.stringify([
  ...m.TS_BASES.filter((b) => b.name !== 'Прочее').map((b) => 'база: ' + b.name),
  ...m.TS_SELF_GROUPS.flatMap((g) => g.items.map((i) => 'спецтехника: ' + i.name)),
  ...m.TS_TEMPLATES.map((t) => 'шаблон: ' + t.name)]));
"""


def key(v):
    return re.sub(r'\s+', ' ', str(v)).strip().lower() if v not in (None, EMPTY) else ''


def build():
    wb = load_workbook(BOOK, read_only=True)
    bind = {(key(r[0]), key(r[1])): r[4] for r in wb['Тип и кузов'].iter_rows(min_row=2, values_only=True)
            if r[0] is not None and r[4]}
    models = collections.defaultdict(collections.Counter)
    spell = {}
    for t, mk, mo, k, n in wb['Все сочетания'].iter_rows(min_row=2, values_only=True):
        if key(mo) == '':
            continue
        mid = (key(mk), key(mo))
        spell.setdefault(mid, (mk if mk != EMPTY else '', mo))
        target = bind.get((key(t), key(k)))
        if target:
            models[mid][target] += n
    wb.close()

    labels = set(json.loads(subprocess.run(['node', '--input-type=module', '-e', JS], cwd=ROOT, capture_output=True,
                                           encoding='utf-8', check=True).stdout))
    stale = sorted({x for c in models.values() for x in c} - labels)
    assert not stale, 'привязки к вариантам, которых нет в каталоге (пересобрать книгу пар): %s' % stale[:10]

    out = []
    for mid, c in sorted(models.items(), key=lambda x: -sum(x[1].values())):
        total = sum(c.values())
        keep = [x for i, (x, n) in enumerate(c.most_common()) if i == 0 or n >= SHARE * total][:LIMIT]
        out.append([spell[mid][0], spell[mid][1], keep])

    head = ('// Модели машин → варианты поиска «Вида объекта»: [марка, модель, [вариант, …]],\n'
            '// вариант — «база: …», «спецтехника: …» или «шаблон: …», по убыванию числа записей.\n'
            '//\n'
            '// ФАЙЛ СОБИРАЕТСЯ СКРИПТОМ — руками не править:\n'
            '//     python tools/data/build_ts_models.py\n'
            '// Источник — книга пар docs/pary-poiska-ts.xlsx (обезличенные записи техпаспортов).\n'
            '//\n'
            '// ДЛЯ СЕРВЕРНОЙ ВЕРСИИ: справочник моделей, который пополняется из карточек и правится\n'
            '// без программиста; здесь он в коде макета.\n')
    body = 'export const TS_MODELS = [\n' + ''.join('  %s,\n' % json.dumps(r, ensure_ascii=False) for r in out) + '];\n'
    io.open(OUT, 'w', encoding='utf-8', newline='\n').write(head + '\n' + body)
    return len(out), sum(len(r[2]) for r in out), sum(1 for r in out if len(r[2]) > 1)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    n, leaves, many = build()
    print('tsModels.js собран: моделей %d, вариантов %d, моделей с несколькими вариантами %d' % (n, leaves, many))
