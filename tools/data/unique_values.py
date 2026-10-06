# -*- coding: utf-8 -*-
"""Уникальные значения по столбцам книги Excel — со счётчиком, без первой строки.

Задача пользователя 06.10.2026: «скрипт, который пройдётся и соберёт все
уникальные значения по столбцам F G H I, игнорируя первую строку». Скрипт
читает ТОЛЬКО названные столбцы (правило проекта: назван столбец — остальные не
читаются) и пишет книгу рядом с исходной: лист на столбец — значение и сколько
раз оно встречается, по убыванию; шапка закреплена, есть автофильтр.

Значения сравниваются без учёта регистра и лишних пробелов («Легковой» и
« легковой » — одно), в книге — самая частая запись. Пустые ячейки — строкой
«(пусто)».

    python tools/data/unique_values.py "книга.xlsx"
    python tools/data/unique_values.py "книга.xlsx" --sheet "Лист1" --cols F G H I --skip 1
    python tools/data/unique_values.py "книга.xlsx" --out "уникальные.xlsx"
"""
import argparse
import collections
import os
import re
import sys

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import column_index_from_string

EMPTY = '(пусто)'


def key(value):
    return re.sub(r'\s+', ' ', str(value)).strip().lower()


def collect(path, sheet, cols, skip):
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    out = {}
    for col in cols:
        idx = column_index_from_string(col)
        counts = collections.Counter()
        spell = collections.defaultdict(collections.Counter)
        for (v,) in ws.iter_rows(min_row=skip + 1, min_col=idx, max_col=idx, values_only=True):
            if v is None or not str(v).strip():
                counts[EMPTY] += 1
                continue
            k = key(v)
            counts[k] += 1
            spell[k][re.sub(r'\s+', ' ', str(v)).strip()] += 1
        out[col] = [(spell[k].most_common(1)[0][0] if k != EMPTY else EMPTY, n) for k, n in counts.most_common()]
    wb.close()
    return ws.title, out


def write(path, sheet_name, data):
    wb = Workbook()
    wb.remove(wb.active)
    head = PatternFill('solid', fgColor='DCEBF8')
    for col, rows in data.items():
        ws = wb.create_sheet('Столбец %s' % col)
        ws.append(['Значение', 'Сколько раз'])
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = head
        for value, n in rows:
            ws.append([value, n])
        ws.column_dimensions['A'].width = 48
        ws.column_dimensions['B'].width = 14
        for (c,) in ws.iter_rows(min_row=2, max_col=1):
            c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
    wb.save(path)


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('book', help='книга .xlsx')
    ap.add_argument('--sheet', default='', help='лист (по умолчанию первый)')
    ap.add_argument('--cols', nargs='+', default=['F', 'G', 'H', 'I'], help='столбцы (по умолчанию F G H I)')
    ap.add_argument('--skip', type=int, default=1, help='сколько первых строк пропустить (по умолчанию 1)')
    ap.add_argument('--out', default='', help='куда записать (по умолчанию «<книга> — уникальные.xlsx» рядом)')
    a = ap.parse_args()

    cols = [c.upper() for c in a.cols]
    title, data = collect(a.book, a.sheet, cols, a.skip)
    out = a.out or os.path.splitext(a.book)[0] + ' — уникальные.xlsx'
    write(out, title, data)
    sys.stdout.reconfigure(encoding='utf-8')
    for col, rows in data.items():
        print('Столбец %s: уникальных %d, строк %d' % (col, len(rows), sum(n for _, n in rows)))
    print('Записано:', out)


if __name__ == '__main__':
    main()
