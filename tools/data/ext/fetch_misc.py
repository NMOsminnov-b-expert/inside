# -*- coding: utf-8 -*-
"""Каталоги моделей: cars-base (легковые, русские названия) и vPIC (грузовики, автобусы, мото, прицепы).

Задача пользователя 07.10.2026: «собираем по маркам и моделям ряд по
интернету. Ищем наиболее полные каталоги и расширяемся».

  * cars-base.ru (github.com/blanzh/carsBase, cars.csv) — марки и модели
    легковых с названием кириллицей и классом; бесплатная часть каталога.
  * vPIC NHTSA (vpic.nhtsa.dot.gov, открытый API) — модели по виду машины:
    Truck, Bus, Motorcycle, Trailer, Incomplete Vehicle (шасси). Каталог
    американский и огромный (тысячи мелких марок прицепов), поэтому берутся
    только марки, которые есть в записях техпаспортов (лист «Марка и модель»
    книги пар docs/pary-poiska-ts.xlsx).

Кэш — tools/data/ext/cache/carsbase.csv и vpic.json: [{make, type, model}].

    python tools/data/ext/fetch_misc.py
"""
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request

from openpyxl import load_workbook

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(os.path.dirname(HERE)))
CACHE = os.path.join(HERE, 'cache')
BOOK = os.path.join(ROOT, 'docs', 'pary-poiska-ts.xlsx')
UA = 'inside-mockup-catalog/1.0 (local research; python urllib)'
TYPES = ['Truck', 'Bus', 'Motorcycle', 'Trailer', 'Incomplete Vehicle']


def get(url, binary=False):
    for attempt in range(4):
        try:
            data = urllib.request.urlopen(urllib.request.Request(url, headers={'User-Agent': UA}), timeout=120).read()
            return data if binary else json.loads(data)
        except Exception as e:  # noqa: BLE001 — повтор при сбое сети
            print('  повтор', url, e, flush=True)
            time.sleep(5 * (attempt + 1))
    return None


def bkey(s):
    return re.sub(r'[^A-ZА-Я0-9]', '', str(s).upper())


def main():
    os.makedirs(CACHE, exist_ok=True)
    csv = get('https://raw.githubusercontent.com/blanzh/carsBase/master/cars.csv', binary=True)
    open(os.path.join(CACHE, 'carsbase.csv'), 'wb').write(csv)
    print('cars-base: строк', csv.count(b'\n'), flush=True)

    wb = load_workbook(BOOK, read_only=True)
    ours = {bkey(r[0]) for r in wb['Марка и модель'].iter_rows(min_row=2, values_only=True) if r[0]}
    wb.close()
    rows = []
    for t in TYPES:
        makes = get('https://vpic.nhtsa.dot.gov/api/vehicles/GetMakesForVehicleType/%s?format=json'
                    % urllib.parse.quote(t))['Results']
        hit = sorted({m['MakeName'] for m in makes if bkey(m['MakeName']) in ours})
        print('vPIC %-18s марок %5d, из них в наших записях %d' % (t, len(makes), len(hit)), flush=True)
        for mk in hit:
            time.sleep(0.4)
            r = get('https://vpic.nhtsa.dot.gov/api/vehicles/GetModelsForMakeYear/make/%s/vehicletype/%s?format=json'
                    % (urllib.parse.quote(mk), urllib.parse.quote(t)))
            for m in (r or {}).get('Results', []):
                rows.append({'make': mk, 'type': t, 'model': m['Model_Name']})
    json.dump(rows, open(os.path.join(CACHE, 'vpic.json'), 'w', encoding='utf-8'), ensure_ascii=False, indent=0)
    print('vPIC: моделей', len(rows), flush=True)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    main()
