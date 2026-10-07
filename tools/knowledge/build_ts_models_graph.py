# -*- coding: utf-8 -*-
"""Модели ТС в граф: запись на марку — модели и чем они записаны в техпаспортах.

Задача пользователя 07.10.2026: «все модели с описанием, что это за типы, в
граф». Источник — книга пар docs/pary-poiska-ts.xlsx (по обезличенной книге
записей техпаспортов): лист «Все сочетания» — сколько раз модель записана с
каким типом ТС и кузовом, лист «Тип и кузов» — к какому варианту поиска
привязана пара «тип + кузов». Поверх — модели из внешних каталогов
(tools/data/ts_model_catalog.json, если есть): у них источник и вид машины по
каталогу.

Запись — knowledge/terms/marka_ts_<марка>.py (метка «модели-ТС»), оглавление —
knowledge/terms/modeli_ts_po_markam.py со связями «содержит» на все марки; на
оглавление ведёт якорь карты раздела «Механизмы и транспортные средства».
Записи с меткой «модели-ТС» пересобираются целиком: руками не править, марка,
которой больше нет в данных, удаляется.

Марки сводятся без учёта регистра, пробелов, дефисов и точек
(«MERCEDES-BENZ» и «MERCEDES BENZ» — одна марка), в заголовке — самая частая
запись.

    python tools/knowledge/build_ts_models_graph.py
"""
import collections
import io
import json
import os
import re
import sys

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
BOOK = os.path.join(ROOT, 'docs', 'pary-poiska-ts.xlsx')
CATALOG = os.path.join(ROOT, 'tools', 'data', 'ts_model_catalog.json')
TERMS = os.path.join(ROOT, 'knowledge', 'terms')
TAG = 'модели-ТС'
INDEX_ID = 'modeli-ts-po-markam'
DATE = '2026-10-07'
EMPTY = '—'
SHOW = 6  # сочетаний «тип, кузов» в пункте модели

TRANSLIT = dict(zip('абвгдеёжзийклмнопрстуфхцчшщъыьэюя',
                    ['a', 'b', 'v', 'g', 'd', 'e', 'e', 'zh', 'z', 'i', 'y', 'k', 'l', 'm', 'n', 'o', 'p', 'r',
                     's', 't', 'u', 'f', 'h', 'c', 'ch', 'sh', 'sch', '', 'y', '', 'e', 'yu', 'ya']))


def slug(s):
    s = ''.join(TRANSLIT.get(ch, ch) for ch in s.lower())
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-') or 'bez-nazvaniya'


def brand_key(s):
    return re.sub(r'[\s\-_.]+', '', str(s)).upper()


def key(v):
    return re.sub(r'\s+', ' ', str(v)).strip().lower() if v not in (None, EMPTY) else ''


def share(n, total):
    return '%d%%' % round(100 * n / total)


def read():
    wb = load_workbook(BOOK, read_only=True)
    bind = {(key(r[0]), key(r[1])): r[4] for r in wb['Тип и кузов'].iter_rows(min_row=2, values_only=True)
            if r[0] is not None and r[4]}
    brands = collections.defaultdict(lambda: {'spell': collections.Counter(), 'models': {}})
    for t, mk, mo, k, n in wb['Все сочетания'].iter_rows(min_row=2, values_only=True):
        if mk in (None, EMPTY) and mo in (None, EMPTY):
            continue
        b = brands[brand_key(mk if mk != EMPTY else 'без марки')]
        b['spell'][mk if mk != EMPTY else 'Без марки'] += n
        m = b['models'].setdefault(key(mo), {'spell': collections.Counter(), 'combos': collections.Counter(),
                                              'targets': collections.Counter(), 'catalog': []})
        m['spell'][mo if mo != EMPTY else 'модель не записана'] += n
        m['combos'][', '.join(x for x in (t, k) if x not in (None, EMPTY)) or 'тип и кузов не записаны'] += n
        target = bind.get((key(t), key(k)))
        m['targets'][target or 'без привязки (неоднозначно или нет в справочнике)'] += n
    wb.close()
    # Внешние каталоги: [{"brand", "model", "kind", "run", "source"}].
    if os.path.exists(CATALOG):
        for c in json.load(io.open(CATALOG, encoding='utf-8')):
            b = brands[brand_key(c['brand'])]
            if not b['spell']:
                b['spell'][c['brand']] += 0
            m = b['models'].setdefault(key(c['model']), {'spell': collections.Counter({c['model']: 0}),
                                                         'combos': collections.Counter(), 'targets': collections.Counter(),
                                                         'catalog': []})
            m['catalog'].append(c)
    return brands


def model_point(m):
    name = m['spell'].most_common(1)[0][0]
    total = sum(m['combos'].values())
    parts = []
    if total:
        combos = m['combos'].most_common()
        txt = '; '.join('%s — %s' % (c, share(n, total)) for c, n in combos[:SHOW])
        if len(combos) > SHOW:
            txt += '; ещё сочетаний: %d' % (len(combos) - SHOW)
        targets = ', '.join('%s (%s)' % (re.sub(r'^шаблон: ', 'шаблон ', x), share(n, total))
                            for x, n in m['targets'].most_common(4))
        parts.append('%s — записей в техпаспортах %d: %s. В поиске карточки: %s.' % (name, total, txt, targets))
    else:
        parts.append('%s — в записях техпаспортов нет.' % name)
    for c in m['catalog']:
        run = ', ходовая: %s' % c['run'] if c.get('run') else ''
        parts.append('По каталогу %s: %s%s.' % (c['source'], c['kind'], run))
    return ' '.join(parts)


def write_record(path, rid, title, tags, source, points, links):
    body = ('# -*- coding: utf-8 -*-\n"""%s\n\nЗапись графа знаний проекта (knowledge/terms). Файл — данные, не код:\n'
            'читается разбором (tools/knowledge/graph.py), не исполняется.\n\nСОБИРАЕТСЯ СКРИПТОМ — руками не править:\n'
            '    python tools/knowledge/build_ts_models_graph.py\n"""\n' % title)
    body += 'ID = %r\nKIND = %r\nTITLE = %r\nTAGS = %r\nSTATUS = %r\nDATE = %r\nSOURCE = %r\n' % (
        rid, 'понятие', title, tags, 'актуально', DATE, source)
    body += 'POINTS = [\n' + ''.join('    %r,\n' % p for p in points) + ']\n'
    body += 'LINKS = [\n' + ''.join('    %r,\n' % l for l in links) + ']\n'
    io.open(path, 'w', encoding='utf-8', newline='\n').write(body)


def build():
    brands = read()
    old = [f for f in os.listdir(TERMS) if f.startswith('marka_ts_') and f.endswith('.py')]
    ids, made = set(), []
    source = ('обезличенные записи техпаспортов ТС (книга «Поиск по названию и  ИНН (УНА).xlsx», лист «Основной»), '
              'сведены в docs/pary-poiska-ts.xlsx')
    if os.path.exists(CATALOG):
        source += '; внешние каталоги моделей — tools/data/ts_model_catalog.json'
    for bk, b in sorted(brands.items(), key=lambda x: -sum(sum(m['combos'].values()) for m in x[1]['models'].values())):
        name = b['spell'].most_common(1)[0][0]
        rid = 'marka-ts-' + slug(name)
        while rid in ids:
            rid += '-2'
        ids.add(rid)
        models = sorted(b['models'].values(), key=lambda m: (-sum(m['combos'].values()), m['spell'].most_common(1)[0][0]))
        total = sum(sum(m['combos'].values()) for m in models)
        spells = [s for s, _ in b['spell'].most_common() if s != name]
        head = 'Марка %s: моделей %d, записей в техпаспортах %d.' % (name, len(models), total)
        if spells:
            head += ' Другие написания марки: %s.' % ', '.join(spells)
        points = [head] + [model_point(m) for m in models]
        title = 'Марка ТС %s: модели и что это за машины' % name
        fname = rid.replace('-', '_') + '.py'
        write_record(os.path.join(TERMS, fname), rid, title, ['понятие', 'ТС', TAG, 'марка'], source, points,
                     [])
        made.append((rid, name, len(models), total, fname))
    for f in old:
        if f not in {m[4] for m in made}:
            os.remove(os.path.join(TERMS, f))

    index_points = [
        'Модели ТС по маркам: что за машина каждая модель — какими типом ТС и кузовом она записана в техпаспортах, '
        'с долей записей, и к какому варианту поиска «Вида объекта» ведёт в карточке ТС. Марок %d, моделей %d.'
        % (len(made), sum(m[2] for m in made)),
        'Собирается скриптом tools/knowledge/build_ts_models_graph.py по книге пар docs/pary-poiska-ts.xlsx '
        '(tools/docs/build_pary_poiska_ts.py) и внешним каталогам моделей; тот же источник у справочника моделей '
        'карточки app/modules/vehicle/data/tsModels.js (tools/data/build_ts_models.py).',
        'Задача пользователя 07.10.2026: «все модели с описанием, что это за типы, в граф».',
        'Крупнейшие марки: ' + ', '.join('%s (%d моделей)' % (m[1], m[2]) for m in sorted(made, key=lambda x: -x[2])[:20]) + '.',
    ]
    write_record(os.path.join(TERMS, INDEX_ID.replace('-', '_') + '.py'), INDEX_ID,
                 'Модели ТС по маркам: что это за машины', ['понятие', 'ТС', TAG, 'оглавление'], source, index_points,
                 [{'тип': 'содержит', 'куда': m[0], 'папка': 'terms'} for m in made]
                 + [{'тип': 'относится к', 'куда': 'privyazat-zapisi-tehpasportov-k-variantam-poiska-ts-po-knige-par',
                     'папка': 'tasks'}])
    return len(made), sum(m[2] for m in made)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    n, models = build()
    print('граф: марок %d, моделей %d (knowledge/terms/marka_ts_*.py, modeli_ts_po_markam.py)' % (n, models))
