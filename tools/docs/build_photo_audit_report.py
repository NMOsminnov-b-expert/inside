# -*- coding: utf-8 -*-
"""Отчёт по областям: раскладка фото регионалов, что разбирается автоматически, что — руками.

Данные — тот же разбор перечня папок, что и в аудите (build_photo_audit.py, photo_audit/).
Счёт ведётся и по папкам, и по объектам: объект — папка с адресом (или названием), к
которой относятся фото; ручная работа планируется по объектам.

    python tools/docs/build_photo_audit_report.py

На выходе — «Анализ по фото регионалов/отчёт_по_областям.docx» (локальный файл).
"""
import os
import sys
import random
import collections as C

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, Cm

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'photo_audit'))
from common import SRC_DIR, BS  # noqa: E402
from parse import parse_all, DEV, MANUAL, CHECK, manual_reason, check_reasons, object_key  # noqa: E402
from template import template  # noqa: E402

OUT = os.path.join(SRC_DIR, 'отчёт_по_областям.docx')
AUTO = {'A', 'A*', 'B', 'B*'}
RE_TIERS = {'A', 'A*', 'B', 'B*', 'D-ЗАК', 'C', 'D'}
MAN = {k: (label, do) for k, label, do in MANUAL}
CHK = {k: (label, do) for k, label, do in CHECK}
# отклонения, по которым область сравнивается со средним
PROFILE = ['WRAP', 'DUP_OBJ', 'LIT_SPELL', 'MERGED', 'INVERTED', 'FLOOR_AS_CAT', 'CAT_SPELL', 'NO_CAT', 'NO_LIT',
           'BLD_NO_LIT', 'NODOC', 'MULTI_LIT', 'ADDR_PLACE', 'NO_ADDR', 'NESTED_OBJ', 'PHOTO_IN_MID']

# Проверка разбора чтением путей (23.09.2026). Цифры — из выборок, прочитанных вручную.
CHECKS = [
    ('Авторазбор (A, B), первая выборка: по 20 случайных путей на область', '140', '5 ошибок, 2 неоднозначных',
     'Все пять ошибок — повторяющиеся шаблоны, исправлены в правилах: папка другой школы внутри папки школы '
     'склеивалась с ней в один объект; «Литер В-2» читалось как «В»; снимки из папки «Поэтажные чертежи» '
     'считались фото литеры; «Тер. больница» (территориальная) принималась за адрес, и настоящий адрес уровнем '
     'выше терялся.'),
    ('Авторазбор (A, B), вторая выборка после исправлений: по 20 других путей на область', '140',
     '2 ошибки, 2 неоднозначных',
     'Исправлены: «Котельная» на участке уходила в участок, а не в строение без литеры; «Литер Н (комната '
     'охраны)» получала категорию «Внутр. помещения». Неоднозначные — «лит а 2», «Литер Б 1»: литера с цифрой '
     'или этаж, помечены для сверки.'),
    ('Ручной разбор (C, D): по 6 путей на область', '42', '3 ложно ручных',
     'Оборудование Бишкекской ТЭЦ без инвентарного номера («тисы», «272-277. Резервуар») стояло в D — теперь '
     'весь заказчик с описью оборудования уходит в оборудование. Участок в группе «земля» — остаётся ручным, '
     'пока не решено, считать ли всё в группах «земли» участками.'),
    ('Примеры ручного разбора и сверки: по 2 пути на каждую причину в каждой области', '162',
     'ошибок авторазбора нет; 9 шаблонов с неверной причиной',
     'Исправлены: объект найден, а ниже непонятное имя — теперь C (литеру выбрать), а не D (объект не найден); '
     '«2 корпус шк» давало литеру «ШК»; «Лит А1й этаж» — литера А и этаж, а не этаж без литеры; «По литерам» '
     'читалось как литера; папка литеры с номером («Литера А - 1корп 7-эт») считалась вложенным объектом; '
     '«Постановление 2», «Тех паспорт нежил. здания» — документы; «Школьная 89» — улица; «Лид Ж» — литера Ж; '
     'длинные списки литер «Д, Е, Ж, З, Л».'),
    ('Авторазбор (A, B), контрольная выборка после всех исправлений: по 10 путей на область', '70',
     '1 ошибка, 1 неточность',
     'Исправлены: «92Б - подвал» — номер дома читался как код заказчика и литера Б; «лит c» латиницей — '
     'литера приводится к кириллице.'),
]
LIMITS = [
    'Сопоставление с базой не проверялось: базы в разборе нет. Проверено только, верно ли из пути извлечены '
    'адрес, литера и категория. Найдётся ли объект в базе по этому адресу — отдельная проверка.',
    'Адрес в папке записан свободным текстом («г.Манас, ул.Пушкина №15 - Школа №3 =++(+)»). Код заказчика есть '
    'в 97–100 % путей, полный ЕНИ — в 0–2 %; сопоставление надёжнее вести по паре «код заказчика + адрес».',
    'Литера берётся из имени папки. Есть ли такая литера у объекта в базе, не проверялось: «Литер Й», '
    '«Литер НЗС», «лит 26» извлекаются как написаны.',
    'Категория без явного слова в имени — «Прочее». Снимки внутри папки не просматривались: папка «Литер А» '
    'может содержать и фасад, и помещения.',
]
ASSUMPTIONS = [
    '«Территория», «двор», «участок», «забор», «газон» — объект имущества «Земельный участок», категория '
    '«Земельный участок».',
    'Этаж («1 этаж», «этаж2», «подвал») — категория «Внутр. помещения» литеры, номер этажа — подпись снимка.',
    'Фото литеры без категории — категория «Прочее».',
    'Объект, названный только по населённому пункту («с. Кара-Бак Сш Нарматов»), разбирается автоматически с '
    'обязательной сверкой найденного объекта (уровни A*, B*).',
    'Объекты в группах «Земли», «Несельхоз участки» не считаются участками автоматически: там лежат и '
    'спортзалы, и стадионы с этажами, и базары.',
]


SHORT = {'Баткенская область': 'Баткен', 'Джалал-Абадская область': 'Джалал-Абад', 'Иссык-Кульская область': 'Иссык-Куль',
         'Нарынская область': 'Нарын', 'Ошская область': 'Ош', 'Таласская область': 'Талас', 'Чуйская область': 'Чуй'}


def pct(n, d):
    return '%d %%' % round(100 * n / d) if d else '—'


def h(doc, text, level):
    doc.add_heading(text, level=level)


def para(doc, text, bold_lead=None):
    q = doc.add_paragraph()
    if bold_lead:
        q.add_run(bold_lead).bold = True
    q.add_run(text)
    return q


def bullets(doc, items):
    for it in items:
        doc.add_paragraph(it, style='List Bullet')


def table(doc, head, rows, widths=None, num=()):
    t = doc.add_table(rows=1, cols=len(head))
    t.style = 'Light Grid Accent 1'
    for i, name in enumerate(head):
        t.rows[0].cells[i].text = name
    for r in rows:
        cells = t.add_row().cells
        for i, v in enumerate(r):
            cells[i].text = str(v if v is not None else '')
            if i in num:
                cells[i].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for row in t.rows:
        for i, c in enumerate(row.cells):
            for q in c.paragraphs:
                for run in q.runs:
                    run.font.size = Pt(9)
            if widths:
                c.width = Cm(widths[i])
    doc.add_paragraph()
    return t


def short_path(p):
    return BS.join([p['dist']] + p['segs'])


def examples(xs, k=2, seed=0):
    rnd = random.Random(seed)
    xs = sorted(xs, key=short_path)
    return [short_path(p) for p in rnd.sample(xs, min(k, len(xs)))]


def stats(xs):
    re_ = [p for p in xs if p['tier'] in RE_TIERS]
    t = C.Counter(p['tier'] for p in xs)
    objs = C.defaultdict(list)
    for p in re_:
        objs[object_key(p)].append(p)
    man_objs = sum(1 for v in objs.values() if any(manual_reason(p) for p in v))
    # руками только «фото без литеры», а своих литер в папках объекта нет: при одной литере в базе — автоматически
    soft = {'C-HEAP', 'C-CAT', 'C-FLOOR'}
    one_lit = sum(1 for v in objs.values() if any(manual_reason(p) for p in v)
                  and {manual_reason(p) for p in v} - {None} <= soft and not any(p['lits'] for p in v))
    return dict(all=len(xs), re=len(re_), t=t, auto=sum(t[k] for k in AUTO), objs=len(objs), man_objs=man_objs,
                check=sum(1 for p in re_ if check_reasons(p)), one_lit=one_lit)


def build():
    P = parse_all()
    for p in P:
        p['tpl'] = template(p, above=False)
    oblasts = sorted({p['obl'] for p in P})
    by = {o: [p for p in P if p['obl'] == o] for o in oblasts}
    S = {o: stats(by[o]) for o in oblasts}
    tot = stats(P)
    oblasts.sort(key=lambda o: -S[o]['auto'] / S[o]['re'])

    doc = Document()
    doc.styles['Normal'].font.size = Pt(10)
    from docx.enum.section import WD_ORIENT
    for s in doc.sections:
        s.orientation = WD_ORIENT.LANDSCAPE
        s.page_width, s.page_height = s.page_height, s.page_width
        s.left_margin = s.right_margin = Cm(1.5)

    h(doc, 'Раскладка фото регионалов по областям', 0)
    para(doc, 'Перечень папок с фото из «9. Осмотрщики (Регионалы)»: %d папок, файлов в перечне нет. Идеал '
              'раскладки — адрес → литера → категория. Категории — справочник макета: Фасад, Внутр. помещения, '
              'Кровля, Конструкции, Земельный участок, Прочее.' % tot['all'])

    # --- итог ---------------------------------------------------------------------------------
    h(doc, 'Итог по областям', 1)
    para(doc, 'Доли — от папок недвижимости. «Авто» — фото уходят в литеру без участия человека (уровни A, B). '
              '«Сверить» — из них папки, где найденный объект или литеру нужно подтвердить. «По заказчику» — '
              'адреса нет, объект определяется кодом заказчика, если у заказчика один объект. «Руками» — '
              'объект известен, литеру выбирает человек (C), или объект не определяется (D).')
    rows = []
    for o in oblasts + ['Все области']:
        s = S[o] if o in S else tot
        t = s['t']
        rows.append([SHORT.get(o, o), s['re'], pct(s['auto'], s['re']), pct(s['check'], s['re']),
                     pct(t['D-ЗАК'], s['re']), pct(t['C'], s['re']), pct(t['D'], s['re']),
                     '%d из %d' % (s['man_objs'], s['objs']), s['one_lit'], t['X-ТС'], t['X-EQP'], t['X-DOC']])
    table(doc, ['Область', 'Папок недвиж.', 'Авто', 'Сверить', 'По заказ-чику', 'Руками: литера', 'Руками: объект',
                'Объектов с ручным разбором', 'Из них авто при одной литере', 'Транс-порт', 'Обору-дование',
                'Фото докум.'], rows,
          widths=[2.4, 1.4, 1.1, 1.2, 1.2, 1.3, 1.3, 2.2, 1.6, 1.2, 1.4, 1.2], num=tuple(range(1, 12)))
    para(doc, '«Из них авто при одной литере» — объекты, где руками только фото без литеры, а папок литер у '
              'объекта нет совсем. Если у такого объекта в базе одна литера, снимки уходят в неё автоматически.')
    para(doc, 'Транспорт, оборудование и снимки документов (госакт, техпаспорт, токтом) — не фото недвижимости: '
              'им нужен свой маршрут импорта, в долях выше они не участвуют.')

    # --- причины по всем областям ---------------------------------------------------------------
    h(doc, 'Что не разбирается автоматически', 1)
    rows = []
    for code, label, do in MANUAL:
        xs = [p for p in P if manual_reason(p) == code]
        objs = {object_key(p) for p in xs}
        rows.append([label, len(xs), len(objs), do] + [sum(p['obl'] == o for p in xs) for o in oblasts])
    table(doc, ['Причина', 'Папок', 'Объектов', 'Что сделать'] + [SHORT[o] for o in oblasts],
          rows, widths=[4.2, 1.2, 1.3, 4.6] + [1.1] * len(oblasts), num=(1, 2) + tuple(range(4, 4 + len(oblasts))))
    h(doc, 'Что разбирается, но требует сверки', 2)
    rows = []
    for code, label, do in CHECK:
        xs = [p for p in P if code in check_reasons(p)]
        rows.append([label, len(xs), do] + [sum(p['obl'] == o for p in xs) for o in oblasts])
    table(doc, ['Что сверить', 'Папок', 'Как'] + [SHORT[o] for o in oblasts], rows,
          widths=[4.5, 1.2, 5.6] + [1.1] * len(oblasts), num=(1,) + tuple(range(3, 3 + len(oblasts))))

    # --- точность -----------------------------------------------------------------------------
    h(doc, 'Точность разбора', 1)
    para(doc, 'Разбор проверен чтением путей: для каждой выборки путь прочитан целиком и сверен с тем, какие '
              'объект, литера и категория из него извлечены.')
    table(doc, ['Выборка', 'Путей', 'Итог', 'Что найдено'], CHECKS, widths=[5, 1.3, 2.8, 9], num=(1,))
    h(doc, 'Что разбором не проверяется', 2)
    bullets(doc, LIMITS)
    h(doc, 'Принятые допущения', 2)
    bullets(doc, ASSUMPTIONS)

    # --- области ------------------------------------------------------------------------------
    all_re = [p for p in P if p['tier'] in RE_TIERS]
    avg = {d: sum(d in p['devs'] for p in all_re) / len(all_re) for d in PROFILE}
    for i, o in enumerate(oblasts):
        s, xs = S[o], by[o]
        re_ = [p for p in xs if p['tier'] in RE_TIERS]
        t = s['t']
        h(doc, o, 1)
        para(doc, '%d папок, из них недвижимость — %d. Автоматически — %s (%d папок), сверить — %d; по коду '
                  'заказчика — %d; руками — %d (литеру выбрать — %d, объект не определяется — %d). Объектов — %d, '
                  'с ручным разбором хотя бы одной папки — %d, из них %d станут автоматическими, если у объекта в '
                  'базе одна литера.'
             % (s['all'], s['re'], pct(s['auto'], s['re']), s['auto'], s['check'], t['D-ЗАК'], t['C'] + t['D'],
                t['C'], t['D'], s['objs'], s['man_objs'], s['one_lit']))
        if t['X-ТС'] or t['X-EQP'] or t['X-DOC']:
            para(doc, 'Отдельно: транспорт — %d, оборудование — %d, снимки документов — %d.'
                 % (t['X-ТС'], t['X-EQP'], t['X-DOC']))

        h(doc, 'Как разложено', 2)
        tc = C.Counter(p['tpl'] for p in re_)
        rows = []
        for tpl, n in tc.most_common(5):
            ex = [p for p in re_ if p['tpl'] == tpl]
            rows.append([tpl, n, pct(n, len(re_)), examples(ex, 1, seed=i)[0]])
        table(doc, ['Шаблон ниже папки объекта', 'Папок', 'Доля', 'Пример'], rows, widths=[5, 1.3, 1.3, 10.4],
              num=(1, 2))
        para(doc, 'Разных шаблонов — %d.' % len(tc))
        hi = []
        for d in PROFILE:
            v = sum(d in p['devs'] for p in re_) / len(re_)
            if v >= 0.05 and v >= 1.5 * avg[d]:
                hi.append('%s — %s папок (по всем областям %s)' % (DEV[d][0], pct(v, 1), pct(avg[d], 1)))
        if hi:
            para(doc, 'Чаще, чем в среднем:')
            bullets(doc, hi)

        h(doc, 'Что не разбирается автоматически', 2)
        rows = []
        for code, label, do in MANUAL:
            ms = [p for p in xs if manual_reason(p) == code]
            if not ms:
                continue
            rows.append([label, len(ms), len({object_key(p) for p in ms}), '\n'.join(examples(ms, 2, seed=i))])
        table(doc, ['Причина', 'Папок', 'Объектов', 'Примеры путей'], rows, widths=[4.2, 1.3, 1.5, 11], num=(1, 2))

        rows = []
        for code, label, do in CHECK:
            ms = [p for p in xs if code in check_reasons(p)]
            if ms:
                rows.append([label, len(ms), '\n'.join(examples(ms, 2, seed=i))])
        if rows:
            h(doc, 'Разбирается, но требует сверки', 2)
            table(doc, ['Что сверить', 'Папок', 'Примеры путей'], rows, widths=[4.6, 1.3, 12.1], num=(1,))

        h(doc, 'Заказчики с наибольшим ручным разбором', 2)
        g = C.defaultdict(list)
        for p in re_:
            g[p.get('org') or '— без кода заказчика'].append(p)
        rows = []
        for org, ps in g.items():
            m = [p for p in ps if manual_reason(p)]
            if m:
                rows.append([org, len(ps), len(m), pct(len(m), len(ps)), len({object_key(p) for p in m})])
        rows.sort(key=lambda r: -r[2])
        table(doc, ['Заказчик', 'Папок', 'Руками', 'Доля', 'Объектов'], rows[:8], widths=[9, 1.6, 1.6, 1.6, 1.8],
              num=(1, 2, 3, 4))
    doc.save(OUT)
    return tot


if __name__ == '__main__':
    tot = build()
    print('папок: %d, недвижимость: %d, авто: %d, объектов: %d, с ручным разбором: %d'
          % (tot['all'], tot['re'], tot['auto'], tot['objs'], tot['man_objs']))
    print('собрано: %s' % OUT)
