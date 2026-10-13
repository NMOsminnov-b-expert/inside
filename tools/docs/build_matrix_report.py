# -*- coding: utf-8 -*-
"""Матрицы и сценарии по ролям — таблица для человека: docs/matricy-i-scenarii.xlsx.

Задача графа matricy-vzaimodeystviy-putey-i-scenarii-po-rolyam. Данные —
записи графа с меткой «сценарии» (сценарии ролей) и последние прогоны
матриц (tools/matrix/out/paths.json, interactions.json). Листы:

  Сценарии                 роль, код, сценарий, чем проверен, итог путей
  Матрица путей            путь × ширина · роль: ✓ / ✗
  Провалы путей            что не сошлось: было / ждали
  Находки взаимодействий   действие с флагом: экран, блок, флаги, подробности
  Кто на кого              строка — блок, которым действовали; столбец — кто
                           ответил; число ответов по всем экранам и повторам
  Действия                 все действия всех экранов: что ответило

    python tools/matrix/paths.py && python tools/matrix/interactions.py
    python tools/docs/build_matrix_report.py
"""
import ast, glob, json, os, re, sys
from collections import Counter, defaultdict

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(ROOT, 'docs', 'matricy-i-scenarii.xlsx')
MX = os.path.join(ROOT, 'tools', 'matrix', 'out')
sys.path.insert(0, os.path.join(ROOT, 'tools', 'matrix'))

HEAD_FILL = PatternFill('solid', fgColor='1F4E5F')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
CELL_FONT = Font(size=10)
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
OK_FILL = PatternFill('solid', fgColor='E6F4EA')
BAD_FILL = PatternFill('solid', fgColor='FCE8E6')
WARN_FILL = PatternFill('solid', fgColor='FFF4E0')


def sheet(wb, title, cols, height=30):
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in cols])
    for i, (_, width) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.fill, cell.font = HEAD_FILL, HEAD_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = height
    ws.freeze_panes = 'B2' if len(cols) > 8 else 'A2'
    return ws


def finish(ws):
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(ws.max_column), max(ws.max_row, 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font, cell.border = CELL_FONT, BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if cell.value == '✓':
                cell.fill = OK_FILL
            elif cell.value == '✗':
                cell.fill = BAD_FILL


def records(tag):
    out = []
    for f in glob.glob(os.path.join(ROOT, 'knowledge', '*', '*.py')):
        mod = ast.parse(open(f, encoding='utf-8').read())
        rec = {n.targets[0].id: ast.literal_eval(n.value) for n in mod.body
               if isinstance(n, ast.Assign) and isinstance(n.targets[0], ast.Name)}
        if tag in rec.get('TAGS', []):
            out.append(rec)
    return sorted(out, key=lambda r: r['ID'])


def load(name):
    p = os.path.join(MX, name)
    return json.load(open(p, encoding='utf-8')) if os.path.exists(p) else []


def build():
    from paths_osmotry import PATHS
    titles = {p['id']: p['title'] for p in PATHS}
    paths = load('paths.json')
    inter = load('interactions.json')
    by_path = defaultdict(list)
    for x in paths:
        by_path[x['id']].append(x['ok'])

    wb = Workbook()
    wb.remove(wb.active)

    ws = sheet(wb, 'Сценарии', [('Роль', 22), ('Код', 7), ('Сценарий', 70), ('Пути матрицы', 26), ('Итог путей', 14),
                                 ('Экраны матрицы взаимодействий', 34), ('Статус записи', 12)])
    for rec in records('сценарии'):
        if rec.get('KIND') != 'понятие':
            continue
        role = re.sub(r'^Сценарии роли «([^»]+)».*$', r'\1', rec['TITLE'])
        for pt in rec['POINTS']:
            m = re.match(r'^([А-ЯA-Z]\d+)\.\s*(.*)$', pt)
            if not m:
                continue
            code, text = m.groups()
            ids = re.findall(r'\b([a-z]+(?:-[a-z]+)+)\b', text.split('Проверено:')[-1]) if 'Проверено:' in text else []
            ids = [i for i in ids if i in titles]
            res = [ok for i in ids for ok in by_path.get(i, [])]
            total = '%d из %d' % (sum(res), len(res)) if res else ('не гонялся' if ids else 'пути нет')
            scr = re.findall(r'«([^»]+)»', text.split('Проверено:')[-1]) if 'Проверено:' in text else []
            ws.append([role, code, text.split(' Проверено:')[0], '\n'.join(ids), total, '\n'.join(scr), rec['STATUS']])
            c = ws.cell(row=ws.max_row, column=5)
            c.value = total
    finish(ws)
    for row in ws.iter_rows(min_row=2, min_col=5, max_col=5):
        for c in row:
            if c.value and ' из ' in c.value:
                a, b = c.value.split(' из ')
                c.fill = OK_FILL if a == b else BAD_FILL
            elif c.value:
                c.fill = WARN_FILL

    combos = sorted({(x['role'], x['width']) for x in paths}, key=lambda k: (k[0], -k[1]))
    ws = sheet(wb, 'Матрица путей', [('Путь', 60), ('Сценарий', 9)] + [('%s · %s' % (w, r), 9) for r, w in combos], height=44)
    res = {(x['id'], x['role'], x['width']): x for x in paths}
    for p in PATHS:
        row = [p['title'], p['scenario']]
        for r, w in combos:
            x = res.get((p['id'], r, w))
            row.append('' if not x else ('✓' if x['ok'] else '✗'))
        if any(row[2:]):
            ws.append(row)
    finish(ws)

    ws = sheet(wb, 'Провалы путей', [('Путь', 50), ('Ширина', 8), ('Роль', 8), ('Что не сошлось', 40), ('Было', 45), ('Ждали', 45)])
    for x in paths:
        if x['ok']:
            continue
        bad = [s for s in x['steps'] if not s['ок']]
        if x.get('error'):
            ws.append([titles.get(x['id'], x['id']), x['width'], x['role'], 'путь остановился', x['error'], ''])
        for s in bad:
            ws.append([titles.get(x['id'], x['id']), x['width'], x['role'], s['что'] + (' (мягко)' if s.get('мягко') else ''),
                       json.dumps(s['было'], ensure_ascii=False), json.dumps(s['ждали'], ensure_ascii=False)])
        for c in x.get('console', [])[:3]:
            ws.append([titles.get(x['id'], x['id']), x['width'], x['role'], 'ошибка консоли', c, ''])
    finish(ws)

    ws = sheet(wb, 'Находки взаимодействий', [('Экран', 30), ('Роль', 7), ('Действие', 9), ('Подпись', 34), ('Блок', 26),
                                               ('Флаги', 20), ('Подробности', 70), ('Повторов', 9)])
    for sc in inter:
        for a in sc['actions']:
            if not a.get('flags'):
                continue
            r = next((x for x in a['runs'] if 'hash1' in x), a['runs'][0] if a['runs'] else {})
            det = []
            for f, key in (('ПЕРЕКРЫТ', 'cover'), ('ЗА КРАЕМ', 'edge'), ('СДВИГ', 'shift'), ('ПРЫГАЛ', 'jumpy'),
                           ('МЕЛЬКНУЛ', 'flash'), ('СБРОС', 'reset')):
                if f in a['flags'] and r.get(key):
                    det.append('%s: %s' % (f.lower(), ', '.join(r[key][:5])))
            for f, key in (('КОНСОЛЬ', 'console'), ('МОДАЛКА', 'dialog')):
                if f in a['flags'] and r.get(key):
                    det.append('%s: %s' % (f.lower(), r[key][0]))
            if r.get('err'):
                det.append(r['err'])
            ws.append([sc['screen'], sc['role'], a['kind'], a['label'], a['block'], ' '.join(a['flags']), '\n'.join(det), len(a['runs'])])
    finish(ws)

    who = defaultdict(Counter)
    for sc in inter:
        for a in sc['actions']:
            for r in a['runs']:
                for k in set(r.get('appeared', []) + r.get('gone', []) + r.get('changed', [])):
                    who[a['block'] or '—'][k] += 1
    cols = [k for k, _ in Counter(k for c in who.values() for k in c.elements()).most_common(60)]
    ws = sheet(wb, 'Кто на кого', [('чем \\ кто', 30)] + [(c, 6) for c in cols], height=120)
    for i in range(2, len(cols) + 2):
        ws.cell(row=1, column=i).alignment = Alignment(text_rotation=90, vertical='bottom')
    for src in sorted(who, key=lambda s: -sum(who[s].values())):
        ws.append([src] + [who[src].get(c) or None for c in cols])
    finish(ws)

    ws = sheet(wb, 'Действия', [('Экран', 30), ('Действие', 9), ('Подпись', 34), ('Блок', 26), ('Адрес после', 30),
                                 ('Появилось', 30), ('Исчезло', 30), ('Изменилось', 30), ('Флаги', 18)])
    for sc in inter:
        for a in sc['actions']:
            r = next((x for x in a['runs'] if 'hash1' in x), None)
            if not r:
                why = 'не нажимается в матрице (%s)' % a['kind'] if a['kind'] in ('наружу', 'файл') else ('недоступно' if a['off'] else '—')
                ws.append([sc['screen'], a['kind'], a['label'], a['block'], '', why, '', '', ' '.join(a.get('flags', []))])
                continue
            ws.append([sc['screen'], a['kind'], a['label'], a['block'], r['hash1'] if r['hash1'] != r['hash0'] else '',
                       '\n'.join(r['appeared'][:6]), '\n'.join(r['gone'][:6]),
                       '\n'.join(k for k in r['changed'][:6] if k not in r['replaced']), ' '.join(a.get('flags', []))])
    finish(ws)

    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(OUT)


if __name__ == '__main__':
    sys.stdout.reconfigure(encoding='utf-8')
    build()
