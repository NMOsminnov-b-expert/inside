# -*- coding: utf-8 -*-
"""Аудит раскладки фото регионалов относительно идеала «адрес → литера → категория».

Исходные данные — перечень папок с фото (Анализ по фото регионалов/папки_с_фото.xlsx):
название папки и полный путь. Файлов в перечне нет, только папки.

Каждый уровень пути получает роль (photo_audit/classify.py): заказчик, служебная группа,
адрес, название объекта, обёртка «Фото/Осмотр», литера, этаж, категория, строение без
литеры, фото документов, транспорт, оборудование. По ролям путь разбирается
(photo_audit/parse.py): где объект, какие литеры и категории, какие отклонения от идеала,
на каком уровне папку можно импортировать автоматически.

Категории — справочник макета: Фасад, Внутр. помещения, Кровля, Конструкции,
Земельный участок, Прочее. «Территория», «двор», «участок» — Земельный участок;
этаж — внутренние помещения литеры.

    python tools/docs/build_photo_audit.py

На выходе — «Анализ по фото регионалов/аудит_раскладки_фото.xlsx» (локальный файл).
"""
import os
import re
import sys
import collections as C

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.hyperlink import Hyperlink

sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), 'photo_audit'))
from common import SRC_DIR, BS, norm  # noqa: E402
from parse import parse_all, DEV, TIER, MANUAL, CHECK, manual_reason, check_reasons, object_key  # noqa: E402
from template import template  # noqa: E402

OUT = os.path.join(SRC_DIR, 'аудит_раскладки_фото.xlsx')

TIER_ORDER = ['A', 'A*', 'B', 'B*', 'D-ЗАК', 'C', 'D', 'X-ТС', 'X-EQP', 'X-DOC']
RE_TIERS = ['A', 'A*', 'B', 'B*', 'D-ЗАК', 'C', 'D']
AUTO = {'A', 'A*', 'B', 'B*'}
TIER_DO = {
    'A':  'Фото уходят в литеру и категорию без участия человека.',
    'A*': 'Как A; объект сопоставляется по населённому пункту и названию — нужна сверка найденного объекта.',
    'B':  'Фото уходят в литеру; категория «Прочее», разбирать по категориям — уже в карточке.',
    'B*': 'Как B; объект сопоставляется по населённому пункту и названию — нужна сверка найденного объекта.',
    'D-ЗАК': 'Объект ищется по коду заказчика. Один объект у заказчика — автоматически, несколько — выбор руками.',
    'C':  'Объект определяется автоматически, литеру выбирает человек (фото всего объекта, строение без литеры, '
          'несколько литер в одной папке, «нет на чертежах»).',
    'D':  'Объект не определяется: ни адреса, ни кода заказчика, либо путь не распознан. Разбор руками.',
    'X-ТС': 'Не недвижимость: транспорт. Импорт в карточки ТС отдельным правилом.',
    'X-EQP': 'Не недвижимость: оборудование и инвентарь (в основном Бишкекская ТЭЦ). Отдельный маршрут.',
    'X-DOC': 'Сняты документы (госакт, техпаспорт, токтом). Место им — во вложениях объекта, а не в фото.',
}
DEV_FIX = {
    'SVC_ABOVE': 'Пропускаются: объект ищется по адресу, а не по глубине. Код заказчика сохраняется как ключ сверки.',
    'WRAP': 'Пропускается по словарю обёрток («фото», «осмотр», «фото с осмотра», «фото и примечание»).',
    'DUP_OBJ': 'Объектом берётся верхняя папка с адресом, повтор ниже пропускается.',
    'LIT_SPELL': 'Буква литеры извлекается из любой формы: «лит а», «литера А», «объект А», «Блок А», «А».',
    'CAT_SPELL': 'Словарь написаний (лист «Написание категорий») сводит слово к категории макета.',
    'MERGED': 'Имя разбирается на литеру и этаж/категорию.',
    'MARKS': 'Отбрасываются: «+», «(+-)», «=++», «осмотрено», «не полн».',
    'FLOOR_AS_CAT': 'Этаж → «Внутр. помещения», номер этажа сохраняется в подписи снимка.',
    'INVERTED': 'Порядок уровней не важен: литера и категория ищутся по всему пути ниже объекта.',
    'OBJ_IN_LIT': 'Название объекта в имени литеры отбрасывается, литера извлекается.',
    'PHOTO_IN_MID': 'Фото из самой папки относятся к её уровню (литера без категории или объект целиком).',
    'ADDR_PLACE': 'Объект ищется по населённому пункту и названию; найденное сверяет человек.',
    'NO_CAT': 'Категория «Прочее».',
    'NO_LIT': 'Литеру выбирает человек; если у объекта одна литера — можно ставить её автоматически.',
    'BLD_NO_LIT': 'Литеру выбирает человек: котельная, туалет, навес чаще всего без своей литеры.',
    'NODOC': 'Руками: строения нет в техпаспорте, литеры для него в базе нет.',
    'MULTI_LIT': 'Руками: фото нескольких литер вперемешку, раскладывать по снимкам.',
    'NO_ADDR': 'Объект ищется по заказчику (уровень D-ЗАК) или руками.',
    'UNREAD': 'Руками: слово не из словарей (лист «Нераспознанное»); после разбора пополнить словарь.',
    'DOC_PHOTO': 'Во вложения объекта, не в фото.',
    'VEH': 'В карточки ТС.',
    'EQP': 'В движимое имущество.',
    'NESTED_OBJ': 'Объектом берётся вложенная папка; к какой карточке ОЦ её относить — сверка по базе.',
    'LIT_AMBIG': 'Принимается как литера с цифрой (А1); если у объекта нет литеры А1 — этаж литеры А.',
    'NUMBERED': 'Номер папки ничего не говорит о содержимом: уровень пропускается.',
}
DEV_ORDER = ['NO_ADDR', 'ADDR_PLACE', 'NO_LIT', 'BLD_NO_LIT', 'NODOC', 'MULTI_LIT', 'UNREAD', 'NO_CAT', 'INVERTED',
             'OBJ_IN_LIT', 'PHOTO_IN_MID', 'NESTED_OBJ', 'LIT_AMBIG', 'NUMBERED', 'SVC_ABOVE', 'WRAP', 'DUP_OBJ', 'LIT_SPELL', 'CAT_SPELL',
             'MERGED', 'FLOOR_AS_CAT', 'MARKS', 'DOC_PHOTO', 'VEH', 'EQP']
NOT_SCORED = {'VEH', 'EQP', 'DOC_PHOTO'}

HEAD_FILL = PatternFill('solid', fgColor='1F4E5F')
TOTAL_FILL = PatternFill('solid', fgColor='E8EEF1')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
CELL_FONT = Font(size=10)
BOLD_FONT = Font(size=10, bold=True)
LINK_FONT = Font(size=10, color='0563C1', underline='single')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
PCT = '0%'
TOC = 'Оглавление'


def link(cell, target):
    cell.hyperlink = Hyperlink(ref=cell.coordinate, location="'%s'!A1" % target, display=str(cell.value))
    cell.font = LINK_FONT


def sheet(wb, title, cols, rows, pct_cols=(), num_cols=(), total=False):
    """cols — [(заголовок, ширина)]; pct_cols/num_cols — номера колонок (с 1) для формата и выравнивания."""
    ws = wb.create_sheet(title)
    ws.append([c for c, _ in cols])
    for r in rows:
        ws.append(list(r))
    for i, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    last = ws.max_row
    for row in ws.iter_rows():
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical='top', wrap_text=True,
                                    horizontal='right' if c.row > 1 and c.column in (set(pct_cols) | set(num_cols)) else None)
            c.font = HEAD_FONT if c.row == 1 else CELL_FONT
            if c.row == 1:
                c.fill = HEAD_FILL
            elif c.column in pct_cols:
                c.number_format = PCT
            elif c.column in num_cols:
                c.number_format = '#,##0' if isinstance(c.value, int) else '0.00'
            if total and c.row == last and last > 1:
                c.fill, c.font = TOTAL_FILL, BOLD_FONT
    ws.freeze_panes = 'B2' if len(cols) > 6 else 'A2'
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(len(cols)), last - (1 if total else 0))
    back = ws.cell(row=1, column=len(cols) + 1, value='← Оглавление')
    link(back, TOC)
    ws.column_dimensions[get_column_letter(len(cols) + 1)].width = 16
    return ws


def score(p):
    return sum(DEV[d][1] for d in p['devs'] if d not in NOT_SCORED)


def strict_ideal(p):
    """Ровно «адрес → литера → категория»: без обёрток, повторов, слияний и своих написаний."""
    extra = p['devs'] - {'SVC_ABOVE', 'MARKS'}
    return p['tier'] == 'A' and not extra


def share(n, d):
    return n / d if d else 0


def lit_form(seg):
    """Форма записи литеры: буква → X, цифры → 9 («лит а1» → «лит X9»)."""
    keys = ('литера', 'литеры', 'литер', 'литтер', 'лит', 'объект', 'обьект', 'объек', 'блок', 'корпус')
    out = []
    for w in re.split(r'(\s+|[,()\-])', norm(seg)):
        w = re.sub(r'\d+', '9', w)
        if w.rstrip('.') in keys:
            out.append(w)
            continue
        glued = next((k for k in keys if w.startswith(k) and re.fullmatch(r'[а-яa-e]9?', w[len(k):])), None)
        if glued:
            out.append(glued + 'X' + ('9' if w.endswith('9') else ''))
        elif re.fullmatch(r'[абвгдежзийклмнопрстуфхцчшщэюяabcde]9?', w):
            out.append('X' + ('9' if w.endswith('9') else ''))
        else:
            out.append(w)
    return ''.join(out)


def build():
    P = parse_all()
    for p in P:
        p['score'] = score(p)
        p['tpl'] = template(p, above=False)
    oblasts = sorted({p['obl'] for p in P})
    RE = [p for p in P if p['tier'] in RE_TIERS]
    by_obl = C.defaultdict(list)
    for p in P:
        by_obl[p['obl']].append(p)

    wb = Workbook()
    wb.remove(wb.active)
    toc = wb.create_sheet(TOC)
    contents = {}

    # --- уровни импорта ------------------------------------------------------------------------
    tc = C.Counter(p['tier'] for p in P)
    rows = [(t, TIER[t], TIER_DO[t], tc[t], share(tc[t], len(P))) for t in TIER_ORDER]
    rows.append(('Всего', '', '', len(P), 1.0))
    sheet(wb, 'Уровни импорта', [('Уровень', 9), ('Что это', 48), ('Что делать при импорте', 70), ('Папок', 10),
                                 ('Доля', 8)], rows, pct_cols=(5,), num_cols=(4,), total=True)
    contents['Уровни импорта'] = 'Уровни автоимпорта: что каждый значит, что делать с папкой, сколько папок'

    # --- области -------------------------------------------------------------------------------
    def obl_row(name, xs):
        re_ = [p for p in xs if p['tier'] in RE_TIERS]
        t = C.Counter(p['tier'] for p in xs)
        n = len(re_)
        devc = C.Counter(d for p in re_ for d in p['devs'] if DEV[d][1] > 0)
        top = '; '.join('%s — %.0f%%' % (DEV[d][0], 100 * share(v, n)) for d, v in devc.most_common(3))
        return [name, len(xs), n, t['X-ТС'], t['X-EQP'], t['X-DOC']] + \
            [share(t[k], n) for k in RE_TIERS] + \
            [share(sum(t[k] for k in AUTO), n), share(sum(t[k] for k in AUTO) + t['D-ЗАК'], n),
             share(t['C'] + t['D'], n), share(sum(strict_ideal(p) for p in re_), n),
             sum(p['score'] for p in re_) / n if n else 0, len({p['tpl'] for p in re_}), top]
    rows = [obl_row(o, by_obl[o]) for o in oblasts]
    rows.sort(key=lambda r: -r[13])
    rows.append(obl_row('Все области', P))
    cols = [('Область', 22), ('Папок всего', 9), ('Недвижимость', 11), ('Транспорт', 9), ('Оборудование', 11),
            ('Фото документов', 10)] + [(k, 7) for k in RE_TIERS] + \
           [('Авто (A, B)', 9), ('Авто, если у заказчика один объект', 12), ('Руками (C, D)', 9),
            ('Строго по идеалу', 9), ('Средний балл отклонения', 10), ('Разных шаблонов', 9),
            ('Главные отклонения (доля папок недвижимости)', 70)]
    sheet(wb, 'Области', cols, rows, pct_cols=tuple(range(7, 18)), num_cols=(2, 3, 4, 5, 6, 18, 19), total=True)
    contents['Области'] = ('Сводка по областям: доли уровней импорта среди папок недвижимости, строгое соответствие '
                           'идеалу, средний балл отклонения, главные отклонения')

    # --- отклонения ----------------------------------------------------------------------------
    rows = []
    for d in DEV_ORDER:
        r = [DEV[d][0], DEV[d][1], DEV_FIX[d]]
        base = P if d in NOT_SCORED else RE
        for o in oblasts:
            xs = [p for p in base if p['obl'] == o]
            r.append(share(sum(d in p['devs'] for p in xs), len(xs)))
        r.append(share(sum(d in p['devs'] for p in base), len(base)))
        r.append(sum(d in p['devs'] for p in base))
        rows.append(r)
    cols = [('Отклонение', 40), ('Вес', 5), ('Как снимается при импорте', 50)] + [(o.replace(' область', ''), 10)
                                                                                 for o in oblasts] + [('Все', 8), ('Папок', 8)]
    ws = sheet(wb, 'Отклонения', cols, rows, pct_cols=tuple(range(4, 4 + len(oblasts) + 1)),
               num_cols=(2, 5 + len(oblasts)))
    ws.cell(row=len(rows) + 3, column=1, value='Вес: 0 — снимается правилом без потерь; 1 — правилом, с риском '
            'ошибки; 3 — только руками. Доля считается от папок недвижимости, для транспорта, оборудования и фото '
            'документов — от всех папок.').font = CELL_FONT
    contents['Отклонения'] = 'Доля папок с каждым отклонением по областям, вес отклонения и как оно снимается при импорте'

    # --- районы и заказчики --------------------------------------------------------------------
    def grp_rows(key, head):
        g = C.defaultdict(list)
        for p in RE:
            g[key(p)].append(p)
        out = []
        for k, xs in g.items():
            t = C.Counter(p['tier'] for p in xs)
            n = len(xs)
            tpl = C.Counter(p['tpl'] for p in xs).most_common(1)[0]
            out.append(list(k) + [n] + [share(t[x], n) for x in RE_TIERS] +
                       [share(sum(t[x] for x in AUTO), n), share(t['C'] + t['D'], n),
                        sum(p['score'] for p in xs) / n, '%s (%.0f%%)' % (tpl[0], 100 * tpl[1] / n)])
        out.sort(key=lambda r: (r[0], -r[len(head)]))
        return out
    head = [('Область', 20), ('Район', 26)]
    rows = grp_rows(lambda p: (p['obl'], p['dist']), head)
    n0 = len(head)
    sheet(wb, 'Районы', head + [('Папок', 8)] + [(k, 7) for k in RE_TIERS] +
          [('Авто (A, B)', 9), ('Руками (C, D)', 9), ('Средний балл', 8), ('Главный шаблон ниже объекта', 50)],
          rows, pct_cols=tuple(range(n0 + 2, n0 + 11)), num_cols=(n0 + 1, n0 + 12))
    contents['Районы'] = 'То же по районам внутри области: где раскладка хуже или лучше средней по области'

    head = [('Область', 20), ('Район', 22), ('Заказчик', 44)]
    rows = grp_rows(lambda p: (p['obl'], p['dist'], p.get('org') or '— без кода заказчика'), head)
    n0 = len(head)
    sheet(wb, 'Заказчики', head + [('Папок', 8)] + [(k, 7) for k in RE_TIERS] +
          [('Авто (A, B)', 9), ('Руками (C, D)', 9), ('Средний балл', 8), ('Главный шаблон ниже объекта', 50)],
          rows, pct_cols=tuple(range(n0 + 2, n0 + 11)), num_cols=(n0 + 1, n0 + 12))
    contents['Заказчики'] = ('По заказчикам (папка с кодом вида «110.1.», «115.65_»): объём ручной работы по каждому '
                             'заказчику')

    # --- шаблоны -------------------------------------------------------------------------------
    g = C.defaultdict(list)
    for p in RE:
        g[p['tpl']].append(p)
    rows = []
    for k, xs in sorted(g.items(), key=lambda kv: -len(kv[1])):
        t = C.Counter(p['tier'] for p in xs).most_common(1)[0][0]
        ex = xs[0]
        rows.append([k, len(xs), share(len(xs), len(RE)), t, BS.join([ex['obl'], ex['dist']] + ex['segs'])] +
                    [sum(p['obl'] == o for p in xs) for o in oblasts])
    sheet(wb, 'Шаблоны', [('Шаблон ниже объекта', 44), ('Папок', 8), ('Доля', 7), ('Чаще всего уровень', 9),
                          ('Пример пути', 90)] + [(o.replace(' область', ''), 9) for o in oblasts], rows,
          pct_cols=(3,), num_cols=(2,) + tuple(range(6, 6 + len(oblasts))))
    contents['Шаблоны'] = ('Цепочки уровней от папки объекта вниз («АДРЕС / фото / литер / категория»), сколько папок '
                           'устроено так в каждой области, пример пути')

    # --- словари написания ---------------------------------------------------------------------
    litc = C.defaultdict(lambda: C.Counter())
    litex = {}
    catc = C.defaultdict(lambda: C.Counter())
    catmap = {}
    unk = C.defaultdict(lambda: C.Counter())
    unkrole = {}
    for p in RE:
        oi = p['obj_i'] if p['obj_i'] is not None else -1
        for i, (r, info) in enumerate(p['roles']):
            if i <= oi:
                continue
            seg = p['segs'][i]
            if 'LIT' in r:
                f = lit_form(seg)
                litc[f][p['obl']] += 1
                litex.setdefault(f, seg)
            elif r & {'CAT', 'FLOOR'}:
                k = norm(seg)
                catc[k][p['obl']] += 1
                cats = list(info.get('cat', []))
                if 'FLOOR' in r and 'Внутр. помещения' not in cats:
                    cats.append('Внутр. помещения (этаж)')
                catmap[k] = ', '.join(cats)
            elif r & {'UNK', 'BLD', 'NODOC', 'NUM'} and not r & {'ADDR', 'OBJ', 'WRAP'}:
                k = norm(seg)
                unk[k][p['obl']] += 1
                unkrole[k] = ('Нераспознано' if 'UNK' in r else 'Нет на чертежах' if 'NODOC' in r else
                              'Номер' if 'NUM' in r else 'Строение без литеры')
    rows = [[f, litex[f], sum(c.values())] + [c[o] for o in oblasts] for f, c in
            sorted(litc.items(), key=lambda kv: -sum(kv[1].values()))]
    sheet(wb, 'Написание литер', [('Форма (X — буква, 9 — цифры)', 34), ('Пример', 30), ('Папок', 8)] +
          [(o.replace(' область', ''), 9) for o in oblasts], rows, num_cols=tuple(range(3, 4 + len(oblasts))))
    contents['Написание литер'] = 'Как в разных областях записывают литеру: «Литер А», «лит а», «А», «литера А» и т. д.'
    rows = [[k, catmap[k], sum(c.values())] + [c[o] for o in oblasts] for k, c in
            sorted(catc.items(), key=lambda kv: -sum(kv[1].values()))]
    sheet(wb, 'Написание категорий', [('Имя папки', 34), ('Категория макета', 26), ('Папок', 8)] +
          [(o.replace(' область', ''), 9) for o in oblasts], rows, num_cols=tuple(range(3, 4 + len(oblasts))))
    contents['Написание категорий'] = ('Словарь для импорта: имя папки → категория макета, сколько раз встречается '
                                       'в каждой области')
    rows = [[k, unkrole[k], sum(c.values())] + [c[o] for o in oblasts] for k, c in
            sorted(unk.items(), key=lambda kv: (kv[0] != '', -sum(kv[1].values())))]
    sheet(wb, 'Нераспознанное', [('Имя папки ниже объекта', 40), ('Что это', 20), ('Папок', 8)] +
          [(o.replace(' область', ''), 9) for o in oblasts], rows, num_cols=tuple(range(3, 4 + len(oblasts))))
    contents['Нераспознанное'] = ('Имена ниже объекта, которые не литера и не категория: строения без литеры, '
                                  '«нет на чертежах», номера, незнакомые слова — материал для пополнения словарей')

    # --- разбор руками и сверка ----------------------------------------------------------------
    man = {k: (label, do) for k, label, do in MANUAL}
    chk = {k: (label, do) for k, label, do in CHECK}
    rows = []
    for p in P:
        r = manual_reason(p)
        codes = [(r, man[r])] if r else [(c, chk[c]) for c in check_reasons(p)]
        for code, (label, do) in codes:
            rows.append([p['obl'], p['dist'], p.get('org') or '', BS.join(object_key(p)[2:]), BS.join(p['segs']),
                         p['tier'], 'Руками' if r else 'Сверить', label, do])
    rows.sort(key=lambda r: (r[0], r[6], r[7], r[3]))
    sheet(wb, 'Разбор руками', [('Область', 16), ('Район', 18), ('Заказчик', 28), ('Объект (путь до папки объекта)', 50),
                                ('Путь ниже района', 70), ('Уровень', 7), ('Что нужно', 9), ('Причина', 40),
                                ('Что сделать', 50)], rows)
    contents['Разбор руками'] = ('Папки, которые не разбираются автоматически (с причиной), и авторазобранные папки, '
                                 'у которых нужно сверить найденный объект или литеру')

    # --- все папки -----------------------------------------------------------------------------
    rows = []
    for p in P:
        obj = p['segs'][p['obj_i']] if p['obj_i'] is not None else ''
        addr = {'full': 'улица и номер', 'place': 'населённый пункт', 'name': 'только название'}.get(p['addr_q'], '')
        devs = '; '.join(DEV[d][0] for d in DEV_ORDER if d in p['devs'])
        rows.append([p['obl'], p['dist'], BS.join(p['segs']), p['segs'][-1] if p['segs'] else '', obj, addr,
                     p.get('eni') or '', p.get('org') or '', ', '.join(p['lits']), ', '.join(p['cats']),
                     p['floor'] or '', p['tier'], p['score'], p['tpl'], devs])
    sheet(wb, 'Все папки', [('Область', 16), ('Район', 18), ('Путь ниже района', 70), ('Папка', 24),
                            ('Папка объекта', 36), ('Адрес', 14), ('ЕНИ', 16), ('Заказчик', 28), ('Литеры', 10),
                            ('Категория', 16), ('Этаж', 8), ('Уровень', 7), ('Балл', 6), ('Шаблон ниже объекта', 34),
                            ('Отклонения', 70)], rows, num_cols=(13,))
    contents['Все папки'] = 'Каждая папка перечня с разбором: объект, литеры, категория, уровень импорта, отклонения'

    # --- оглавление ----------------------------------------------------------------------------
    toc.append(['Лист', 'Что на листе', 'Строк'])
    for name in wb.sheetnames:
        if name == TOC:
            continue
        toc.append([name, contents[name], wb[name].max_row - 1])
        link(toc.cell(row=toc.max_row, column=1), name)
    for i, w in enumerate((24, 100, 8), start=1):
        toc.column_dimensions[get_column_letter(i)].width = w
    for row in toc.iter_rows():
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical='top', wrap_text=True)
            if c.row == 1:
                c.font, c.fill = HEAD_FONT, HEAD_FILL
            elif c.column != 1:
                c.font = CELL_FONT
    toc.freeze_panes = 'A2'
    wb.save(OUT)
    return P, tc


if __name__ == '__main__':
    P, tc = build()
    print('папок: %d' % len(P))
    for t in TIER_ORDER:
        print('  %-6s %6d' % (t, tc[t]))
    print('собрано: %s' % OUT)
