# -*- coding: utf-8 -*-
"""Сверка макета с исходными таблицами оценщиков — таблицей расхождений.

Источники в корне проекта:
  «Группы движкимого имущества (параметры).xlsx» — классы, подгруппы, типы и
  параметры движимого имущества; «Спецтехника» из него живёт не в механизмах,
  а в типах ТС (решение пользователя 22.09.2026).
  «Классы литер-помещений_Никите (сравн(капит и сост-е)).xlsx» — характеристики
  капитальности гражданских и производственных помещений, их коэффициенты,
  карманы классов и шкала состояния.

Макет берётся не глазами, а из его же справочников (tools/docs/app_data_dump.mjs):
классификатор механизмов, поля механизмов, типы и поля ТС, перечни карточки
литеры.

    python tools/docs/build_sverka_s_tablicami.py

На выходе — docs/sverka-s-tablicami.xlsx:

    Итого                сводка: сколько расхождений в каждом разделе
    Движимое — состав    классы, подгруппы и типы — чего нет в макете и чего
                         нет в таблице
    Движимое — параметры параметр таблицы → поля макета; «не покрыт» значит,
                         что параметру не соответствует ни одно поле
    Поля сверх таблицы   поля макета, которых в таблице нет
    Спецтехника          типы и параметры класса «Спецтехника» → типы и поля ТС
    Классы литер         характеристики капитальности → поля карточки литеры
"""
import io
import json
import os
import re
import subprocess
import sys

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOCS = os.path.join(ROOT, 'docs')
OUT = os.path.join(DOCS, 'sverka-s-tablicami.xlsx')
SRC_MOVABLE = os.path.join(ROOT, 'Группы движкимого имущества (параметры).xlsx')
SRC_CLASSES = os.path.join(ROOT, 'Классы литер-помещений_Никите (сравн(капит и сост-е)).xlsx')

HEAD_FILL = PatternFill('solid', fgColor='1F4E5F')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
CELL_FONT = Font(size=10)
BAD_FILL = PatternFill('solid', fgColor='FCE4E4')
WARN_FILL = PatternFill('solid', fgColor='FFF4E0')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WORD = re.compile(r'[А-Яа-яЁёA-Za-z]{3,}')
STOP = {'для', 'или', 'при', 'под', 'над', 'без', 'если', 'как', 'что', 'это', 'том',
        'его', 'все', 'него', 'them', 'and'}
# Обобщённые параметры таблицы: одно название покрывает несколько полей макета.
GENERAL = ('основные технические', 'габаритные размеры', 'характеристики')


# Тот же разбор текста, что у сборщика классификатора: он правит опечатки,
# пробелы у косой черты и заглавную букву. Без этого сверка ловила бы не
# расхождения, а разницу в написании.
sys.path.insert(0, os.path.join(ROOT, 'tools', 'data'))
import build_mech_classifier as BUILD  # noqa: E402

clean = BUILD.clean
cap = BUILD.cap
param = BUILD.param


def stem_list(text):
    return [w.lower().replace('ё', 'е')[:6] for w in WORD.findall(text) if w.lower() not in STOP]


def stems(text):
    return set(stem_list(text))


def head_stem(text):
    st = stem_list(text)
    return st[0] if st else ''


def covers(par, f):
    """Покрывает ли поле карточки параметр таблицы.

    Совпадение по двум общим словам, по вхождению подписи поля целиком
    («Комплектность» ↔ «Комплектность (наличие насадок…)») или по первому слову
    («Наработка, моточасы» ↔ «Наработка / остаточный ресурс»)."""
    ps = stems(par)
    fs = stems(' '.join([f['label'], ' '.join(f.get('options', [])), ' '.join(f.get('units', []))]))
    lab = stems(f['label'])
    return (len(ps & fs) >= 2 or (len(ps) <= 2 and bool(ps & fs))
            or (bool(lab) and lab <= ps)
            or (head_stem(par) != '' and head_stem(f['label']) == head_stem(par)))


def app_data():
    out = subprocess.run([node(), os.path.join(ROOT, 'tools', 'docs', 'app_data_dump.mjs')],
                         cwd=ROOT, capture_output=True)
    if out.returncode:
        sys.exit('не удалось снять справочники макета:\n' + out.stderr.decode('utf-8', 'replace'))
    return json.loads(out.stdout.decode('utf-8'))


def node():
    return 'node.exe' if os.name == 'nt' else 'node'


# --- исходные таблицы ---------------------------------------------------------
def movable_table():
    """Классификатор движимого: {класс: {подгруппа: {'types': [...], 'params': [...]}}}."""
    ws = openpyxl.load_workbook(SRC_MOVABLE, data_only=True)['Классификатор']
    rows = list(ws.iter_rows(values_only=True))
    base = [clean(v) for v in rows[1][4:6]]

    out, cls, sub = {}, '', ''
    for r in rows[2:]:
        r = list(r) + [None] * 14
        cls = cap(clean(r[1])) or cls
        sub = clean(r[2]) or sub
        typ = clean(r[3])
        typ = cap(BUILD.TYPOS.get(typ, typ))
        if not cls or not sub:
            continue
        s = out.setdefault(cls, {}).setdefault(sub, {'types': [], 'params': list(base)})
        if typ and typ not in s['types']:
            s['types'].append(typ)
        for i in (6, 7, 8, 9, 10):
            v = param(r[i])
            if v and v not in s['params']:
                s['params'].append(v)
    return out


def classes_table():
    """Лист dict: {(тип, характеристика): [(значение, коэффициент)]}."""
    ws = openpyxl.load_workbook(SRC_CLASSES, data_only=True)['dict']
    out = {}
    for r in range(2, ws.max_row + 1):
        char, kind = clean(ws.cell(r, 1).value), clean(ws.cell(r, 2).value)
        val, k = clean(ws.cell(r, 3).value), ws.cell(r, 4).value
        if not char or not kind:
            continue
        out.setdefault((kind, char), []).append((val, k))
    return out


# --- сверка движимого ---------------------------------------------------------
# Класс таблицы, который в механизмы не переносится: это транспорт.
SPEC = 'Спецтехника'

# Общие поля карточки механизма (oi/mech/view.js): они стоят над параметрами
# подгруппы и в справочник полей не входят, поэтому при сверке их надо знать
# отдельно — иначе базовые параметры таблицы выглядели бы непокрытыми.
COMMON_FIELDS = [
    'Наименование', 'Инвентарный номер', 'Год ввода в эксплуатацию',
    'Количество, шт.', 'Балансовая стоимость, сом', 'Страна происхождения',
    'Комментарий', 'Свои поля',
]


def compare_movable(table, app):
    """Состав: расхождения по классам, подгруппам и типам."""
    rows, stat = [], {'нет в макете': 0, 'нет в таблице': 0}
    app_by = {c['name']: c for c in app['mech']}

    for cls, subs in table.items():
        if cls == SPEC:
            continue
        c = app_by.get(cls)
        if not c:
            rows.append([cls, '', '', 'нет в макете', 'класс есть в таблице, в классификаторе его нет'])
            stat['нет в макете'] += 1
            continue
        app_subs = {s['name']: s for s in c['subgroups']}
        for sub, d in subs.items():
            s = app_subs.get(sub)
            if not s:
                rows.append([cls, sub, '', 'нет в макете', 'подгруппа есть в таблице'])
                stat['нет в макете'] += 1
                continue
            for t in d['types']:
                if t not in s['types']:
                    rows.append([cls, sub, t, 'нет в макете', 'тип есть в таблице'])
                    stat['нет в макете'] += 1
            for t in s['types']:
                if t not in d['types']:
                    rows.append([cls, sub, t, 'нет в таблице', 'тип заведён в макете'])
                    stat['нет в таблице'] += 1
        for sub in app_subs:
            if sub not in subs:
                rows.append([cls, sub, '', 'нет в таблице', 'подгруппа заведена в макете'])
                stat['нет в таблице'] += 1

    for c in app['mech']:
        if c['name'] in table or c['name'] == SPEC:
            continue
        why = 'взят с листа «Схема»' if c['fromSchema'] else 'заведён в макете'
        rows.append([c['name'], '', '', 'нет в таблице', why])
        stat['нет в таблице'] += 1
    return rows, stat


def compare_params(table, app):
    """Параметры подгрупп: покрыт ли параметр таблицы полями макета."""
    covered, extra = [], []
    app_by = {c['name']: c for c in app['mech']}

    for cls, subs in table.items():
        if cls == SPEC:
            continue
        c = app_by.get(cls)
        if not c:
            continue
        app_subs = {s['name']: s for s in c['subgroups']}
        for sub, d in subs.items():
            s = app_subs.get(sub)
            if not s:
                continue
            fields = [{'key': 'общее:' + x, 'label': x, 'options': [], 'units': []}
                      for x in COMMON_FIELDS] + list(s['fields'])
            for bt in s['byType']:
                fields += bt['add']
            used = set()
            for p in d['params']:
                hit = []
                for f in fields:
                    if covers(p, f):
                        hit.append(f['label'])
                        used.add(f['key'])
                general = any(g in p.lower() for g in GENERAL)
                if hit:
                    covered.append([cls, sub, p, 'покрыт', '; '.join(sorted(set(hit))[:6])])
                elif general:
                    covered.append([cls, sub, p, 'обобщённый', 'раскрыт полями подгруппы'])
                else:
                    covered.append([cls, sub, p, 'НЕ ПОКРЫТ', ''])
            for f in fields:
                if f['key'].startswith('общее:'):
                    continue
                if f['key'] not in used:
                    extra.append([cls, sub, f['label'], f['key']])
    return covered, extra


# Опознавательный блок карточки ТС (vehicle/view.js): стоит над параметрами
# типа, в справочник полей не входит.
VEHICLE_COMMON = ['Тип ТС', 'Марка и модель', 'Государственный номер', 'VIN',
                  'Год выпуска', 'Цвет', 'Страна-изготовитель']


def compare_spec(table, app):
    """Спецтехника таблицы → типы и поля ТС."""
    rows = []
    subs = table.get(SPEC, {})
    app_by = {v['type']: v for v in app['vehicle']}
    for sub, d in subs.items():
        v = app_by.get(sub)
        if not v:
            rows.append([sub, '', 'нет в макете', 'подгруппа таблицы не заведена типом ТС'])
            continue
        kinds = []
        for f in v['passport']:
            if f['key'] == 'specKind':
                kinds = f['options']
        for t in d['types']:
            rows.append([sub, t, 'есть' if t in kinds else 'НЕТ В МАКЕТЕ',
                         'вид техники' if t in kinds else ''])
        for t in kinds:
            if t not in d['types']:
                rows.append([sub, t, 'нет в таблице', 'заведён в макете'])
        fields = ([{'label': x, 'options': [], 'units': []} for x in VEHICLE_COMMON]
                  + v['passport'] + v['inspect'])
        for p in d['params']:
            hit = [f['label'] for f in fields if covers(p, f)]
            rows.append([sub, 'параметр: ' + p, 'покрыт' if hit else 'НЕ ПОКРЫТ', '; '.join(hit[:4])])
    return rows


# --- сверка классов литер -----------------------------------------------------
# Сопоставление характеристик таблицы с тем, что есть в карточке литеры.
# Заполняется руками: в таблице это характеристики расчёта капитальности, в
# карточке — поля осмотра и техпаспорта, автоматически они не сходятся.
LETTER_MAP = [
    ('гражд', 'Высота', 'Высота по внешним / внутренним замерам, м (число)',
     'число вместо четырёх диапазонов таблицы; коэффициент по диапазону не считается'),
    ('гражд', 'Архитектура', '—', 'поля нет'),
    ('гражд', 'Конструкция', 'Конструктив и износ (материалы по элементам)',
     'сводной оценки «некапитальная / смешанная / капитальная» нет'),
    ('гражд', 'Планировка', '—', 'поля нет'),
    ('гражд', 'Уровень отделки', 'Отделка (материал в таблице конструктива)',
     'уровня из четырёх ступеней нет'),
    ('гражд', 'Инженерное оснащение', 'Отопление, коммуникации участка',
     'шкалы из трёх ступеней нет'),
    ('произв', 'Высота', 'Высота, м (ТП) в доп. параметрах', 'число вместо диапазонов'),
    ('произв', 'Наличие / Возможность кран-балки', 'Наличие/возможность кран-балки', 'сверяется значениями'),
    ('произв', 'Конструкция', 'Конструктив (доп. параметры производственного)', 'сверяется значениями'),
    ('произв', 'Полы', 'Полы (несущая способность)', 'сверяется значениями'),
    ('произв', 'Инженерное оснащение', '—', 'поля нет'),
]
# Какие перечни макета сверяются со значениями таблицы.
LETTER_DICTS = {
    ('произв', 'Конструкция'): 'PROD_FRAME',
    ('произв', 'Полы'): 'PROD_FLOORS',
    ('произв', 'Наличие / Возможность кран-балки'): 'CRANE_BEAM',
}


def compare_letters(dic, app):
    rows = []
    for kind, char, where, note in LETTER_MAP:
        vals = [v for v, _ in dic.get((kind, char), [])]
        key = LETTER_DICTS.get((kind, char))
        app_vals = app['civil'].get(key, []) if key else []
        if where == '—':
            status = 'ПОЛЯ НЕТ'
            diff = 'значения таблицы: ' + '; '.join(vals)
        elif key:
            miss = [v for v in vals if v not in app_vals]
            extra = [v for v in app_vals if v not in vals]
            status = 'значения расходятся' if (miss or extra) else 'совпадает'
            diff = ' | '.join(filter(None, [
                ('нет в макете: ' + '; '.join(miss)) if miss else '',
                ('нет в таблице: ' + '; '.join(extra)) if extra else '']))
        else:
            status = 'заполняется иначе'
            diff = 'значения таблицы: ' + '; '.join(vals)
        rows.append([kind, char, where, status, note, diff])

    rows.append(['оба', 'Класс помещения', 'Класс ОИ — выбирается вручную', 'СЧИТАЕТСЯ В ТАБЛИЦЕ',
                 'в таблице класс вычисляется из капитальности (карманы 0,33 / 0,55 / 0,77)',
                 'капитальности и коэффициентов в макете нет'])
    rows.append(['оба', 'Капитальность', '—', 'ПОЛЯ НЕТ',
                 'произведение коэффициентов характеристик', ''])
    rows.append(['оба', 'Физический износ, %', 'Износ по элементам: «Умеренный» / «Значительный»',
                 'ЕДИНИЦЫ РАЗНЫЕ', 'в таблице процент из затратного подхода', ''])
    rows.append(['оба', 'Состояние', 'Внешнее / внутреннее / итоговое, пять ступеней',
                 'ШКАЛЫ РАЗНЫЕ',
                 'в таблице восемь ступеней из процента износа, на листе «Шкалы» — пять диапазонов, '
                 'на «Выжимке» — девять словесных градаций', ''])
    return rows


# --- вывод --------------------------------------------------------------------
def sheet(wb, title, cols):
    ws = wb.create_sheet(title)
    ws.append([c for c, _ in cols])
    for i, (_, w) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    return ws


def finish(ws, bad_col=None, bad_words=()):
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical='top', wrap_text=True)
            c.font = HEAD_FONT if c.row == 1 else CELL_FONT
            if c.row == 1:
                c.fill = HEAD_FILL
        if bad_col and row[0].row > 1:
            v = str(row[bad_col - 1].value or '')
            if any(w in v for w in bad_words):
                for c in row:
                    c.fill = BAD_FILL
    ws.auto_filter.ref = ws.dimensions


def build():
    table = movable_table()
    dic = classes_table()
    app = app_data()

    comp, stat = compare_movable(table, app)
    covered, extra = compare_params(table, app)
    spec = compare_spec(table, app)
    letters = compare_letters(dic, app)

    wb = Workbook()
    wb.remove(wb.active)

    it = sheet(wb, 'Итого', [('Раздел', 46), ('Расхождений', 14), ('Пояснение', 70)])
    not_covered = sum(1 for r in covered if r[3] == 'НЕ ПОКРЫТ')
    spec_bad = sum(1 for r in spec if 'НЕ' in str(r[2]))
    letters_bad = sum(1 for r in letters if r[3].isupper() or 'расходятся' in r[3])
    it.append(['Движимое — состав классификатора', len(comp),
               'классы, подгруппы и типы, которые есть только с одной стороны'])
    it.append(['Движимое — параметры таблицы без полей', not_covered,
               'параметр таблицы, которому не соответствует ни одно поле карточки'])
    it.append(['Движимое — поля сверх таблицы', len(extra),
               'поля карточки, которых в таблице нет (раскрытие обобщённых параметров и предложения)'])
    it.append(['Спецтехника → ТС', spec_bad, 'типы и параметры спецтехники без своего места в карточке ТС'])
    it.append(['Классы литер', letters_bad,
               'характеристики капитальности, которых нет в карточке или которые считаются иначе'])
    finish(it)

    ws = sheet(wb, 'Движимое — состав', [('Класс', 40), ('Подгруппа', 40), ('Тип', 44),
                                        ('Где', 16), ('Пояснение', 40)])
    for r in comp:
        ws.append(r)
    finish(ws, 4, ('нет в макете',))

    ws = sheet(wb, 'Движимое — параметры', [('Класс', 32), ('Подгруппа', 34),
                                           ('Параметр таблицы', 60), ('Статус', 16), ('Поля карточки', 60)])
    for r in covered:
        ws.append(r)
    finish(ws, 4, ('НЕ ПОКРЫТ',))

    ws = sheet(wb, 'Поля сверх таблицы', [('Класс', 32), ('Подгруппа', 34), ('Поле', 48), ('Ключ', 20)])
    for r in extra:
        ws.append(r)
    finish(ws)

    ws = sheet(wb, 'Спецтехника', [('Подгруппа таблицы / тип ТС', 40), ('Тип или параметр', 62),
                                   ('Статус', 18), ('Поля карточки ТС', 50)])
    for r in spec:
        ws.append(r)
    finish(ws, 3, ('НЕТ', 'НЕ ПОКРЫТ'))

    ws = sheet(wb, 'Классы литер', [('Тип', 10), ('Характеристика таблицы', 30),
                                    ('Где в карточке', 46), ('Статус', 22),
                                    ('Пояснение', 52), ('Разница значений', 60)])
    for r in letters:
        ws.append(r)
    finish(ws, 4, ('НЕТ', 'РАЗНЫЕ', 'СЧИТАЕТСЯ', 'расходятся'))

    if not os.path.isdir(DOCS):
        os.makedirs(DOCS)
    wb.save(OUT)
    return {'состав': len(comp), 'без полей': not_covered, 'сверх таблицы': len(extra),
            'спецтехника': spec_bad, 'литеры': letters_bad}


if __name__ == '__main__':
    res = build()
    print('docs/sverka-s-tablicami.xlsx собран')
    for k, v in res.items():
        print('  %-16s %d' % (k, v))
