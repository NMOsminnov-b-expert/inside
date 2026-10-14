# -*- coding: utf-8 -*-
"""Аудит экрана выездного осмотрщика по типам ОЦ и видам ОИ — таблица xlsx.

Задача пользователя 14.10.2026: «Нужен аудит, что и как в интерфейсе
осмотрщика. Как ведёт себя интерфейс с разными ОЦ ОИ. Что по ТС?».

Сценарий проходит в браузере каждую пару «тип ОЦ × вид ОИ» из данных макета
(первую попавшуюся запись с таким ОИ) по экранам осмотрщика — задача,
объект, осмотр ОИ, документы, фото — и записывает, что показано: заголовок,
ЕНИ, карта, «Участок по Кадастру», сводка, разделы формы осмотра, число
полей, документы, категории снимков, ошибки консоли, прокрутка вбок.
Замечания ставит по правилам ниже (RULES) — по фактам экрана, без догадок.

    python tools/docs/build_audit_osmotrschika.py   # → docs/audit-ekrana-osmotrschika.xlsx
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
OUT = os.path.join(ROOT, 'docs', 'audit-ekrana-osmotrschika.xlsx')
sys.path.insert(0, os.path.join(ROOT, 'tools', 'checks'))

KIND = {'building': 'Литера (строение)', 'land': 'Земельный участок', 'apartment': 'Квартира', 'vehicle': 'ТС',
        'mech': 'Механизмы и оборудование', 'movable': 'Движимое (прежний вид)', 'aux': 'Вспомогательная постройка'}
BUILDING_FORM = ('Чек-лист коммуникаций', 'Конструктивные элементы')

PAIRS_JS = """async () => { const r = await import('./app/kernel/registry.js'); const out = []; const types = {};
  for (const t of r.sortedTypes()) {
    const rows = t.records.queryRecords({ filter: {}, offset: 0, limit: 100000 }).rows;
    types[t.manifest.id] = { label: t.manifest.label || t.manifest.title, n: rows.length };
    const seen = new Set();
    for (const s of rows) { const rec = t.records.loadRecord(s.id); if (!rec) continue;
      for (const o of rec.oi || []) { if (seen.has(o.card)) continue; seen.add(o.card);
        out.push({ type: t.manifest.id, typeLabel: types[t.manifest.id].label, oc: s.id, oi: o.id, card: o.card, eni: rec.eni || '', hasGeo: !!(rec.geo && rec.geo.extent), gps: rec.gps || '' }); } }
  } return { pairs: out, types }; }"""

SCREEN_JS = """() => { const q = (s) => document.querySelector(s); const qa = (s) => [...document.querySelectorAll(s)];
  const txt = (s) => (q(s) ? q(s).textContent.trim().replace(/\\s+/g, ' ') : '');
  return {
    title: txt('.ins-head-title'), eni: txt('.ins-head-eni b'), sub: txt('.ins-head-sub'), strip: txt('.ins-strip b'),
    map: !!q('[data-ins-map]'), cadp: !!q('.ins-cadp'), cadpText: txt('.ins-cadp'),
    rows: qa('.ins-body .ins-row-l').map((e) => e.textContent.trim()).slice(0, 30),
    blocks: qa('.ins-body .ins-blk-h, .ins-body .ins-acc-t b').map((e) => e.textContent.trim()),
    sections: qa('.ins-acc-t b').map((e) => e.textContent.trim()),
    fields: qa('[data-fs-field]').length, req: qa('.ins-fs-req').length,
    secnav: qa('[data-goto-sec]').map((e) => e.textContent.trim().replace(/\\s+/g, ' ')),
    oiItems: qa('.ins-oi-h, .ins-oi-t, .ins-oi-name').map((e) => e.textContent.trim().replace(/\\s+/g, ' ')).slice(0, 12),
    docs: qa('.ins-doc, .ins-doc-i, [data-doc-open]').length, photoCats: qa('[data-cat], .ins-cat').map((e) => e.textContent.trim()).slice(0, 12),
    empty: qa('.ins-empty').map((e) => e.textContent.trim()),
    pageX: document.documentElement.scrollWidth > innerWidth + 1,
  }; }"""


def collect():
    import harness
    from playwright.sync_api import sync_playwright
    srv = harness.serve()
    base = 'http://127.0.0.1:%d/app.html' % harness.PORT
    res = []
    try:
        with sync_playwright() as p:
            b = p.chromium.launch()
            pg = b.new_page(viewport={'width': 390, 'height': 844})
            errs = []
            pg.on('pageerror', lambda e: errs.append(str(e)[:160]))
            pg.on('console', lambda m: m.type == 'error' and errs.append(m.text[:160]))
            pg.goto(base + '#/insp')
            pg.wait_for_selector('.ins-body')
            meta = pg.evaluate(PAIRS_JS)
            for pr in meta['pairs']:
                task = '#/insp/%s/%s' % (pr['type'], pr['oc'])
                screens = [('Задача', task), ('Объект', task + '/object'), ('Осмотр ОИ', task + '/object/' + pr['oi']),
                           ('Документы', task + '/docs'), ('Фото', task + '/photo')]
                for name, route in screens:
                    errs.clear()
                    pg.goto(base + route)
                    try:
                        pg.wait_for_selector('.ins-body', timeout=8000)
                    except Exception:
                        pass
                    pg.wait_for_timeout(700)
                    try:
                        info = pg.evaluate(SCREEN_JS)
                    except Exception as e:
                        info = {'error': str(e)[:120]}
                    res.append({**pr, 'screen': name, 'route': route, 'info': info, 'errors': list(errs)})
            b.close()
            return meta, res
    finally:
        srv.terminate()


# --- правила замечаний: (условие, текст, важность) --------------------------

def findings(r):
    i = r['info']
    out = []
    if 'error' in i:
        return [('Экран не открылся: ' + i['error'], 'высокая')]
    if r['errors']:
        out.append(('Ошибки в консоли: ' + '; '.join(r['errors'][:2]), 'высокая'))
    if i.get('pageX'):
        out.append(('Страница прокручивается вбок на 390 px', 'средняя'))
    card = r['card']
    movable = card in ('vehicle', 'mech', 'movable')
    if r['screen'] == 'Осмотр ОИ':
        secs = i.get('sections', [])
        # Чужие разделы: чек-лист коммуникаций — только у строения; конструктив
        # — у строения и квартиры (он есть в карточке квартиры).
        own = {'building': BUILDING_FORM, 'apartment': ('Конструктивные элементы',)}.get(card, ())
        alien = [x for x in secs if x in BUILDING_FORM and x not in own]
        if alien:
            out.append(('Разделы формы строения на «%s»: %s — поля не про этот вид ОИ' % (KIND.get(card, card), ', '.join(alien)), 'высокая'))
        if not secs and not i.get('fields'):
            if card == 'vehicle':
                out.append(('Формы осмотра ТС нет — ждёт решения пользователя («ТС пока что ждём», 14.10.2026)', 'низкая'))
            else:
                out.append(('Полей осмотра нет — осмотреть этот ОИ нечем, кроме снимков', 'высокая'))
        if not i.get('title'):
            out.append(('Нет заголовка экрана', 'средняя'))
    if r['screen'] == 'Задача':
        if movable and i.get('cadp') and r['type'] in ('vehicle', 'mechanisms'):
            out.append(('«Участок по Кадастру» у движимого ОЦ — участка у него нет', 'средняя'))
        if not i.get('map'):
            out.append(('Карты нет: у объекта нет координат ни в карточке, ни с Кадастра', 'низкая'))
        if not i.get('eni'):
            out.append(('В шапке нет ЕНИ', 'средняя'))
    if r['screen'] == 'Документы' and not i.get('docs') and not any('не назначены' in e for e in i.get('empty', [])):
        out.append(('Документов не показано, и нет пояснения почему', 'средняя'))
    # Адрес — в сведениях об объекте, а не в шапке (решение пользователя 14.10.2026).
    if i.get('sub') and (',' in i['sub'] or 'ул.' in i['sub']) and len(i['sub']) > 20:
        out.append(('Адрес в шапке под ЕНИ («%s…») — по решению адрес стоит в сведениях об объекте' % i['sub'][:30], 'средняя'))
    if r['screen'] == 'Осмотр ОИ' and movable and i.get('eni') and r['eni'] and i['eni'].replace('-', '') in r['eni'].replace('-', ''):
        out.append(('В шапке ЕНИ здания, а не номер «%s» (госномер, VIN, инвентарный номер)' % KIND.get(card, card), 'средняя'))
    return out


def build(meta, res):
    wb = Workbook()
    hdr_fill = PatternFill('solid', fgColor='DCE6F1')
    sev_fill = {'высокая': PatternFill('solid', fgColor='F8D7D3'), 'средняя': PatternFill('solid', fgColor='FCEBC8'),
                'низкая': PatternFill('solid', fgColor='EEF2F6')}

    def sheet(ws, head, rows, widths):
        ws.append(head)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = hdr_fill
            c.alignment = Alignment(wrap_text=True, vertical='top')
        for row in rows:
            ws.append(row)
        for col, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(col)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical='top')
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions

    # Замечания
    ws = wb.active
    ws.title = 'Замечания'
    rows = []
    for r in res:
        for text, sev in findings(r):
            rows.append([r['typeLabel'], KIND.get(r['card'], r['card']), r['screen'], text, sev, r['route']])
    for k, v in meta['types'].items():
        if not v['n']:
            rows.append([v['label'], '—', 'Все', 'Записей этого типа ОЦ в данных нет — назначить и осмотреть его нельзя, экран осмотрщика не проверить', 'высокая', ''])
    order = {'высокая': 0, 'средняя': 1, 'низкая': 2}
    rows.sort(key=lambda x: (order[x[4]], x[0], x[1], x[2]))
    sheet(ws, ['Тип ОЦ', 'Вид ОИ', 'Экран', 'Замечание', 'Важность', 'Адрес экрана'], rows, [22, 22, 12, 70, 11, 46])
    for row in ws.iter_rows(min_row=2):
        row[4].fill = sev_fill.get(row[4].value, PatternFill())

    # Что показано
    ws2 = wb.create_sheet('Что показано')
    rows2 = []
    for r in res:
        i = r['info']
        if 'error' in i:
            rows2.append([r['typeLabel'], KIND.get(r['card'], r['card']), r['screen'], '', '', '', '', '', '', i['error']])
            continue
        what = []
        if r['screen'] == 'Задача':
            what.append('карта: %s; участок по Кадастру: %s' % ('да' if i['map'] else 'нет', 'да' if i['cadp'] else 'нет'))
            what.append('сводка: ' + ', '.join(i['rows'][:8]))
        elif r['screen'] == 'Объект':
            what.append('блоки: ' + ', '.join(i['blocks']))
        elif r['screen'] == 'Осмотр ОИ':
            what.append('разделы: ' + (', '.join(i['sections']) or '—'))
            what.append('строка разделов: ' + (', '.join(i['secnav']) or '—'))
        elif r['screen'] == 'Документы':
            what.append('документов: %d' % i['docs'])
        elif r['screen'] == 'Фото':
            what.append('блоки: ' + ', '.join(i['blocks']))
        rows2.append([r['typeLabel'], KIND.get(r['card'], r['card']), r['screen'], i.get('title', ''), i.get('eni', ''),
                      i.get('fields', 0), i.get('req', 0), '; '.join(what), '; '.join(i.get('empty', [])[:2]), '; '.join(r['errors'][:2])])
    sheet(ws2, ['Тип ОЦ', 'Вид ОИ', 'Экран', 'Заголовок', 'ЕНИ в шапке', 'Полей', 'Обязательных', 'Что на экране', 'Пустые места', 'Ошибки консоли'],
          rows2, [20, 20, 12, 30, 20, 8, 12, 70, 40, 30])

    # Типы ОЦ
    ws3 = wb.create_sheet('Типы ОЦ в данных')
    kinds_by_type = {}
    for pr in meta['pairs']:
        kinds_by_type.setdefault(pr['type'], []).append(KIND.get(pr['card'], pr['card']))
    sheet(ws3, ['Тип ОЦ', 'Записей в данных', 'Виды ОИ в записях', 'Замечание'],
          [[v['label'], v['n'], ', '.join(kinds_by_type.get(k, [])) or '—',
            'Записей нет: назначить и осмотреть объект этого типа в макете нельзя, экран осмотрщика для него не проверить' if not v['n'] else '']
           for k, v in meta['types'].items()], [28, 16, 50, 60])

    wb.save(OUT)
    return rows


def main():
    meta, res = collect()
    rows = build(meta, res)
    print('пар ОЦ × ОИ: %d, экранов: %d, замечаний: %d (высоких %d) → %s' % (
        len(meta['pairs']), len(res), len(rows), sum(1 for r in rows if r[4] == 'высокая'), os.path.relpath(OUT, ROOT)))
    for r in rows:
        print(' ·', r[4], '|', r[0], '|', r[1], '|', r[2], '|', r[3][:120])


if __name__ == '__main__':
    main()
