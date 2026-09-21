# -*- coding: utf-8 -*-
"""Журнал изменений макета — docs/reestr-izmeneniy.xlsx.

    python tools/docs/build_change_log.py

Строки журнала пишутся руками в change_log_data.py (что изменилось и почему —
это знает человек, а не история git). Остальное снимается с git при каждой
сборке, поэтому не устаревает: дата, автор, в каких ветках правка, влита ли
она в refactor и main, какие модули затронуты.

Листы:
  «Изменения»            — весь журнал;
  «Только в гражданском» — содержательные правки, которых пока нет в других
                           типах ОЦ: по ним актуализируют остальные карточки;
  «Переименования»       — было → стало, для сверки подписей;
  «Не разобрано»         — коммиты с начала журнала, которых в нём нет.
                           Лист не пустой — журнал отстал.
"""
import io
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(os.path.dirname(HERE))
OUT = os.path.join(ROOT, 'docs', 'reestr-izmeneniy.xlsx')
DATA = 'tools/docs/change_log_data.py'

sys.path.insert(0, HERE)
import change_log_data as D  # noqa: E402

HEAD_FILL = PatternFill('solid', fgColor='1F4E5F')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
CELL_FONT = Font(size=10)
MONO = Font(size=10, name='Consolas')
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WARN_FILL = PatternFill('solid', fgColor='FFF4E0')

MODULES = [
    ('app/modules/civil/', 'Гражданское'),
    ('app/modules/apartment/', 'Квартира'),
    ('app/modules/residential-house/', 'Жилое здание'),
    ('app/modules/production/', 'Производственное'),
    ('app/modules/land-plot/', 'Земельный участок'),
    ('app/modules/vehicle/', 'ТС'),
    ('app/kernel/', 'Ядро'),
    ('app/pages/', 'Страницы'),
    ('app/shell/', 'Оболочка'),
    ('tools/checks/', 'Проверки'),
    ('docs/', 'Документы'),
    ('.claude/knowledge-graph/', 'Граф знаний'),
]


def git(*args):
    out = subprocess.run(['git', *args], cwd=ROOT, stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    return out.stdout.decode('utf-8', 'replace').strip() if out.returncode == 0 else ''


def ref_exists(name):
    return bool(git('rev-parse', '--verify', '--quiet', name))


def target(name):
    """Ветка для признака «влито»: своя, иначе удалённая."""
    for ref in (name, 'origin/' + name):
        if ref_exists(ref):
            return ref
    return None


class Commit:
    def __init__(self, line):
        self.short, self.full, self.date, self.author, self.subject = line.split('|', 4)

    def __repr__(self):
        return self.short


def commits_since(start):
    raw = git('log', '--all', '--since=%sT00:00' % start, '--date=short',
              '--format=%h|%H|%ad|%an|%s')
    return [Commit(line) for line in raw.splitlines() if line]


_cache = {}


def info(h):
    """Сведения о коммите: дата, автор, модули, ветки, влито ли."""
    if h in _cache:
        return _cache[h]
    line = git('log', '-1', '--date=short', '--format=%h|%H|%ad|%an|%s', h)
    if not line:
        _cache[h] = None
        return None
    c = Commit(line)
    files = git('diff-tree', '--no-commit-id', '--name-only', '-r', '-m', '--first-parent', c.full).splitlines()
    c.modules = ', '.join(label for prefix, label in MODULES if any(f.startswith(prefix) for f in files))
    heads = [b.strip().lstrip('* ').replace('remotes/', '') for b in
             git('branch', '-a', '--contains', c.full).splitlines()]
    heads = sorted({b for b in heads if b and '->' not in b and not b.startswith('origin/HEAD')})
    c.branches = ', '.join(heads)
    for name in ('refactor', 'main'):
        ref = target(name)
        ok = ref and subprocess.run(['git', 'merge-base', '--is-ancestor', c.full, ref], cwd=ROOT,
                                    stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL).returncode == 0
        setattr(c, 'in_' + name, 'да' if ok else 'нет')
    _cache[h] = c
    return c


def commit_of_row(text):
    """Коммит, в котором строка появилась в данных журнала."""
    found = git('log', '--reverse', '--format=%h', '-S', text, '--', DATA).splitlines()
    return found[0] if found else ''


def resolve(c, text):
    h = c or commit_of_row(text)
    return info(h) if h else None


def sheet(wb, title, cols):
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in cols])
    for i, (_, width) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 30
    ws.freeze_panes = 'A2'
    return ws


def finish(ws, mono_cols=()):
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(ws.max_column), max(ws.max_row, 2))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = MONO if cell.column in mono_cols else CELL_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=True)


def ru_date(iso):
    return '.'.join(reversed(iso.split('-'))) if iso else ''


def build():
    wb = Workbook()
    wb.remove(wb.active)

    rows = []
    for i, r in enumerate(D.CHANGES):
        c = resolve(r['c'], r['what'])
        rows.append((c.date if c else '9999', i, r, c))
    rows.sort(key=lambda x: (x[0], x[1]))

    # --- Изменения ------------------------------------------------------------
    ws = sheet(wb, 'Изменения', [
        ('Дата', 11), ('Раздел', 22), ('Вид', 14), ('Что изменилось', 70), ('Охват', 18),
        ('Повод', 36), ('Модули в коде', 22), ('Коммит', 10), ('Автор', 18), ('Ветки', 30),
        ('В refactor', 10), ('В main', 9),
    ])
    for _, _, r, c in rows:
        ws.append([
            ru_date(c.date) if c else 'не закоммичено', r['where'], r['kind'], r['what'],
            D.SCOPES.get(r['scope'], r['scope']), r['why'],
            c.modules if c else '', c.short if c else '', c.author if c else '',
            c.branches if c else '', c.in_refactor if c else '', c.in_main if c else '',
        ])
    finish(ws, mono_cols=(8,))

    # --- Только в гражданском ------------------------------------------------
    ws = sheet(wb, 'Только в гражданском', [
        ('Дата', 11), ('Раздел', 22), ('Вид', 14), ('Что изменилось', 80), ('Повод', 36), ('Коммит', 10),
    ])
    for _, _, r, c in rows:
        if r['scope'] == 'civil' and r['kind'] in D.CONTENT_KINDS:
            ws.append([ru_date(c.date) if c else 'не закоммичено', r['where'], r['kind'], r['what'],
                       r['why'], c.short if c else ''])
    finish(ws, mono_cols=(6,))

    # --- Переименования -------------------------------------------------------
    ws = sheet(wb, 'Переименования', [
        ('Дата', 11), ('Где', 34), ('Было', 36), ('Стало', 36), ('Коммит', 10),
    ])
    ren = []
    for h, where, was, now in D.RENAMES:
        c = info(h) if h else None
        ren.append((c.date if c else '9999', ru_date(c.date) if c else '', where, was, now, c.short if c else ''))
    for r in sorted(ren, key=lambda x: x[0]):
        ws.append(list(r[1:]))
    finish(ws, mono_cols=(5,))

    # --- Не разобрано ---------------------------------------------------------
    covered = {c.full for _, _, _, c in rows if c}
    missing = []
    for c in commits_since(D.START):
        if c.full in covered:
            continue
        # Правка, у которой строка журнала заведена в том же коммите.
        if DATA in git('diff-tree', '--no-commit-id', '--name-only', '-r', c.full).splitlines():
            continue
        missing.append(info(c.short))
    ws = sheet(wb, 'Не разобрано', [
        ('Дата', 11), ('Коммит', 10), ('Автор', 18), ('Описание коммита', 70), ('Ветки', 36),
    ])
    for c in sorted(missing, key=lambda x: x.date):
        ws.append([ru_date(c.date), c.short, c.author, c.subject, c.branches])
    if not missing:
        ws.append(['', '', '', 'Все коммиты с %s внесены в журнал' % ru_date(D.START), ''])
    finish(ws, mono_cols=(2,))
    if missing:
        wb['Не разобрано'].sheet_properties.tabColor = 'E07B00'
        for row in wb['Не разобрано'].iter_rows(min_row=2):
            for cell in row:
                cell.fill = WARN_FILL

    wb.save(OUT)
    out = io.open(sys.stdout.fileno(), 'w', encoding='utf-8', closefd=False)
    out.write('журнал: %s\n' % os.path.relpath(OUT, ROOT))
    out.write('строк: %d, только в гражданском: %d, переименований: %d, не разобрано: %d\n' % (
        len(rows), wb['Только в гражданском'].max_row - 1, len(ren), len(missing)))
    for c in missing:
        out.write('  не разобран: %s %s %s\n' % (c.short, c.date, c.subject))


if __name__ == '__main__':
    build()
