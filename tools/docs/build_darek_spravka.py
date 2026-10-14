# -*- coding: utf-8 -*-
"""Документация запросов darek.kg (карта ГИСАР и адресный регистр) — docx.

Просьба пользователя 14.10.2026: «Соберешь информацию? Как получать такие
данные. Так же в формате docx». Знания — задача графа
izuchit-darek-kg-adresnyy-reestr; соседняя документация по cadastre.kg —
tools/docs/build_kadastr_spravka.py.

Сборщик делает немного лёгких запросов (около десяти, с паузой): состав слоёв
и полей берётся из описания картосервера, примеры — настоящие ответы. Слой
целиком в примере ограничен тремя объектами: сервер государственный, и
выкачивать слой при каждой пересборке незачем.

    python tools/docs/build_darek_spravka.py      # docs/darek-api.docx

На сервер уходят только ЕНИ и точка демо-записи и служебные запросы. Личные
данные и закрытое входом не запрашиваются.
"""
import http.cookiejar
import json
import os
import re
import sys
import time
import urllib.parse
import urllib.request
import xml.etree.ElementTree as ET

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, 'docs', 'darek-api.docx')
QGIS = 'http://map.darek.kg/qgis/qgis_mapserv.fcgi.exe'
PROJECT = 'C:/OSGeo4W64/projects/ГИСАР/ГИСАР.qgz'
ENI = '1030700400754'
LON, LAT = 74.62421, 42.86995          # ул. Суюмбаева А., д. 33, Бишкек
BISHKEK_ATE = '7402'

_jar = http.cookiejar.CookieJar()
_op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(_jar))


def call(url, data=None, method='GET', ref='http://darek.kg/Map'):
    body = urllib.parse.urlencode(data).encode() if (data is not None and method == 'POST') else None
    if data is not None and method == 'GET':
        url += ('&' if '?' in url else '?') + urllib.parse.urlencode(data)
    hdr = {'User-Agent': 'Mozilla/5.0', 'Referer': ref, 'X-Requested-With': 'XMLHttpRequest'}
    if body:
        hdr['Content-Type'] = 'application/x-www-form-urlencoded'
    time.sleep(1)                                  # бережно к государственному серверу
    try:
        r = _op.open(urllib.request.Request(url, data=body, headers=hdr, method=method), timeout=90)
        raw = r.read()
        return r.status, r.headers.get('Content-Type') or '', raw
    except urllib.error.HTTPError as e:
        return e.code, '', e.read()
    except Exception as e:
        return 0, '', str(e).encode()


def text(raw):
    return raw.decode('utf-8', 'replace') if isinstance(raw, bytes) else raw


def plain(html):
    t = re.sub(r'<script.*?</script>|<style.*?</style>', ' ', html, flags=re.S)
    return re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', t)).strip()


def excerpt(t, n=1400):
    t = t.strip()
    if t[:1] in '{[':
        try:
            def cut(o):
                if isinstance(o, dict):
                    return {k: ('…координаты сокращены…' if k == 'coordinates' else cut(v)) for k, v in o.items()}
                if isinstance(o, list):
                    return [cut(x) for x in o[:3]] + (['…ещё %d' % (len(o) - 3)] if len(o) > 3 else [])
                return o
            t = json.dumps(cut(json.loads(t)), ensure_ascii=False, indent=1)
        except ValueError:
            pass
    elif t.startswith('<?xml'):
        t = re.sub(r'\s+', ' ', t)
    elif t.startswith('<') or '<html' in t[:400].lower():
        t = '[HTML, текст без разметки] ' + plain(t)
    return t[:n] + ('…' if len(t) > n else '')


def curl(method, url, data=None):
    q = "curl -s -H 'Referer: http://darek.kg/Map'"
    if method == 'GET':
        return "%s '%s%s'" % (q, url, (('&' if '?' in url else '?') + urllib.parse.urlencode(data)) if data else '')
    return "%s -X POST '%s' %s" % (q, url, ' '.join("--data-urlencode '%s=%s'" % (k, v) for k, v in data.items()))


def qp(**kw):
    return dict({'map': PROJECT}, **kw)


# --- слои: из описания проекта ------------------------------------------------
def project_layers():
    code, ct, raw = call(QGIS, qp(SERVICE='WMS', REQUEST='GetProjectSettings'))
    ns = {'w': 'http://www.opengis.net/wms'}
    rows = []
    try:
        root = ET.fromstring(raw)
    except ET.ParseError:
        return rows, code
    for L in root.iter('{http://www.opengis.net/wms}Layer'):
        name = L.find('w:Name', ns)
        if name is None or name.text in ('GISAR',):
            continue
        geom = L.get('geometryType')
        attrs = []
        for a in L.iter():
            if a.tag.endswith('}Attribute'):
                attrs.append((a.get('name'), a.get('alias') or '', a.get('type') or ''))
        group = geom is None
        rows.append(dict(name=name.text, geom=geom or 'группа', attrs=attrs, group=group))
    return rows, code


GEOM = {'Polygon': 'многоугольник', 'MultiPolygon': 'многоугольники', 'MultiSurface': 'многоугольники', 'MultiLineString': 'линии',
        'LineStringZ': 'линии', 'группа': 'группа слоёв'}


GENERAL = [
    ('Что это', 'darek.kg — «Адресный регистр Кыргызской Республики» ГУ «Кадастр». Две системы: карта ГИСАР (http://darek.kg/Map) с '
                'картосервером QGIS и адресный регистр АРКР (http://address.darek.kg). Работает по http (без https).'),
    ('Картосервер', '%s, всегда с параметром map=%s. Стандарт WMS 1.3.0 (GetCapabilities, GetProjectSettings, GetMap, '
                    'GetFeatureInfo); WFS слоёв не отдаёт.' % (QGIS, PROJECT)),
    ('Сессия и заголовки', 'Картосервер отвечает без сессии. Запросы карты (/Map/…) — с заголовком Referer: http://darek.kg/Map; '
                           'X-Requested-With: XMLHttpRequest — как у страницы.'),
    ('Координаты', 'В WMS 1.3.0 при CRS=EPSG:4326 рамка BBOX пишется широтой вперёд: «мин. широта, мин. долгота, макс. широта, макс. '
                   'долгота». При CRS=EPSG:3857 — метры, x вперёд. Геометрия в ответе GetFeatureInfo — в той же системе, что CRS запроса.'),
    ('Названия слоёв', 'В параметрах LAYERS, QUERY_LAYERS, FILTER — русские названия слоёв как есть («Границы районов»); в адресе — '
                       'в кодировке URL.'),
    ('Кодировка', 'UTF-8.'),
    ('Ошибки', 'Слой закрыт для внешних запросов — 403 Forbidden («Граница Кыргызстана», «Озёра»); условие FILTER не подошло — пустой '
               'список features; закрыто входом — страница авторизации вместо ответа.'),
    ('Нагрузка', 'Сервер государственный: слой целиком — только когда нужен, с паузами между запросами. Массовое получение — по '
                 'официальному доступу (решение пользователя 14.10.2026).'),
    ('Не делаем', 'Личные данные и закрытое входом не запрашиваются.'),
]


def sections(layers):
    lat, lon, d = LAT, LON, 0.0015
    bbox_pt = '%f,%f,%f,%f' % (lat - d, lon - d, lat + d, lon + d)
    point_layers = 'Участки ЕНИ,Здания ЕНИ,Улицы,Квартал (насел. пункт),Границы айылных аймаков'
    return [
        dict(name='Какие слои есть и какие у них поля', method='GET', url=QGIS,
             data=qp(SERVICE='WMS', REQUEST='GetProjectSettings'),
             purpose='Описание проекта карты: все слои, их геометрия, поля и подписи полей, какие слои можно запрашивать. Полный '
                     'перечень — раздел «Слои и поля».',
             params=[('map', 'да', 'путь проекта', PROJECT), ('SERVICE', 'да', 'WMS', 'WMS'),
                     ('REQUEST', 'да', 'GetProjectSettings (подробно) или GetCapabilities (кратко)', 'GetProjectSettings')],
             fmt='XML', fields=[('Layer/Name', 'название слоя — его пишут в LAYERS и FILTER'), ('Layer@geometryType', 'вид геометрии'),
                                ('Layer@queryable', '1 — слой отвечает на GetFeatureInfo'),
                                ('Attributes/Attribute@name, @alias', 'поле и его подпись на карте')],
             errors='—', note='Ответ большой (около 80 КБ): брать один раз и хранить.'),
        dict(name='Сведения в точке: что лежит в этом месте', method='GET', url=QGIS,
             data=qp(SERVICE='WMS', VERSION='1.3.0', REQUEST='GetFeatureInfo', LAYERS=point_layers, QUERY_LAYERS=point_layers,
                     CRS='EPSG:4326', BBOX=bbox_pt, WIDTH=101, HEIGHT=101, I=50, J=50, INFO_FORMAT='application/json',
                     FEATURE_COUNT=20, WITH_GEOMETRY='TRUE', FI_POLYGON_TOLERANCE=2, FI_LINE_TOLERANCE=8),
             purpose='Все объекты указанных слоёв в точке — со сведениями и геометрией. Точка — центр картинки WIDTH×HEIGHT, '
                     'описанной рамкой BBOX; I, J — пиксель точки.',
             params=[('LAYERS, QUERY_LAYERS', 'да', 'названия слоёв через запятую', point_layers),
                     ('CRS', 'да', 'EPSG:4326 или EPSG:3857', 'EPSG:4326'),
                     ('BBOX', 'да', 'рамка; в EPSG:4326 — широта вперёд', bbox_pt),
                     ('WIDTH, HEIGHT, I, J', 'да', 'размер «картинки» и пиксель точки; точка — центр', '101, 101, 50, 50'),
                     ('INFO_FORMAT', 'да', 'application/json — GeoJSON; text/xml — XML', 'application/json'),
                     ('WITH_GEOMETRY', 'нет', 'TRUE — с геометрией', 'TRUE'),
                     ('FEATURE_COUNT', 'нет', 'сколько объектов вернуть на слой', '20'),
                     ('FI_POLYGON_TOLERANCE, FI_LINE_TOLERANCE', 'нет', 'допуск попадания в пикселях — для линий (улиц) нужен больше',
                      '2, 8')],
             fmt='JSON (GeoJSON)', fields=[('features[].id', '«Слой.номер объекта»'), ('features[].properties', 'сведения — подписи полей '
                                           'как на карте (раздел «Слои и поля»)'), ('features[].geometry', 'геометрия в системе CRS')],
             errors='Пусто — в точке нет объектов этих слоёв.', note='Пример ниже — ул. Суюмбаева А., д. 33.'),
        dict(name='Слой целиком (по условию, без точки)', method='GET', url=QGIS,
             data=qp(SERVICE='WMS', VERSION='1.3.0', REQUEST='GetFeatureInfo', LAYERS='Границы районов', QUERY_LAYERS='Границы районов',
                     CRS='EPSG:4326', BBOX='39.1,69.2,43.3,80.3', WIDTH=400, HEIGHT=400, INFO_FORMAT='application/json',
                     FEATURE_COUNT=3, WITH_GEOMETRY='TRUE', FILTER='Границы районов:"dname_r" IS NOT NULL'),
             purpose='Все объекты слоя, подходящие под условие FILTER, — без точки. Проверено 14.10.2026: «Границы районов» целиком — '
                     '46 районов с границами (1,8 МБ, 2 с). В примере FEATURE_COUNT=3, чтобы не выкачивать слой.',
             params=[('FILTER', 'да', '«Слой:условие» — выражение QGIS по полям слоя (имена полей — без подписей)',
                      'Границы районов:"dname_r" IS NOT NULL'),
                     ('FEATURE_COUNT', 'да', 'предел числа объектов; для всего слоя — с запасом', '3'),
                     ('BBOX', 'да', 'рамка всей страны', '39.1,69.2,43.3,80.3'),
                     ('остальные', 'да', 'как в «Сведениях в точке», без I и J', '')],
             fmt='JSON (GeoJSON)', fields=[('features[]', 'объекты слоя, как в «Сведениях в точке»')],
             errors='Условие не подошло — пустой список (так было у «Границ областей» и «Границ айыл аймаков» — подходящее условие не '
                    'найдено). Закрытый слой — 403.',
             note='Массово — по официальному доступу.'),
        dict(name='Картинка слоя', method='GET', url=QGIS,
             data=qp(SERVICE='WMS', VERSION='1.3.0', REQUEST='GetMap', LAYERS='Улицы,Участки ЕНИ', STYLES='', CRS='EPSG:4326',
                     BBOX=bbox_pt, WIDTH=256, HEIGHT=256, FORMAT='image/png', TRANSPARENT='TRUE'),
             purpose='Растровая картинка слоёв в рамке — для подложки карты.',
             params=[('LAYERS', 'да', 'названия слоёв', 'Улицы,Участки ЕНИ'), ('BBOX, CRS, WIDTH, HEIGHT', 'да', 'рамка и размер', ''),
                     ('FORMAT', 'да', 'image/png', 'image/png'), ('TRANSPARENT', 'нет', 'TRUE — прозрачный фон', 'TRUE')],
             fmt='картинка PNG', fields=[], errors='—', note='Есть и выгрузка в DXF (FORMAT=application/dxf) — кнопка на карте.'),
        dict(name='Справочник территорий (АТЕ): область → район → айыл аймак → село', method='POST', url='http://darek.kg/Map/ateChildren',
             data={'ate_id': BISHKEK_ATE},
             purpose='Дочерние административно-территориальные единицы с кодами СОАТЕ. Коды областей (ate_id): Баткенская 5971, '
                     'г. Бишкек 7402, Джалал-Абадская 5219, Иссык-Кульская 4949, Нарынская 5764, г. Ош 7412, Ошская 6217, Таласская 6802, '
                     'Чуйская 6940.',
             params=[('ate_id', 'да', 'номер родительской единицы', BISHKEK_ATE)],
             fmt='JSON', fields=[('child[].id', 'номер единицы — передавать дальше как ate_id'), ('child[].value', 'название по-русски'),
                                 ('child[].valueKg', 'название по-киргизски'), ('child[].type', 'уровень (3 — район)'),
                                 ('child[].code', 'код СОАТЕ, 14 цифр'), ('ateH.type', 'уровень родителя'),
                                 ('header', 'сведения органа самоуправления (контакты) — в примере пусто')],
             errors='—', note=''),
        dict(name='Улицы населённого пункта и их прежние названия (адресный регистр)', method='GET',
             url='http://address.darek.kg/search/street/%s' % BISHKEK_ATE, data={'streetType': 0, 'streetName': 'Суюмбаева', 'filter': 'sort'},
             purpose='Поиск улиц в адресном регистре: действующее и прежние названия на русском и киргизском, населённый пункт, ID '
                     'улицы. Номер территории — в пути адреса (ate_id из справочника территорий).',
             params=[('<ate_id> в пути', 'да', 'номер территории', BISHKEK_ATE), ('streetName', 'нет', 'часть названия; пробелы — «_»',
                                                                                  'Суюмбаева'),
                     ('streetType', 'нет', 'тип улицы: 0 — любой, 6 — улица, 1 — переулок, 2 — проспект, 4 — микрорайон и др.', '0'),
                     ('filter', 'нет', 'sort — по названию', 'sort')],
             fmt='HTML', fields=[('таблица', 'столбцы: № п/п, ID, действующее название (кырг., рус.), предыдущие названия (кырг., рус.), '
                                           'населённый пункт')],
             errors='—', note='Ссылка «Сообщить об улице» на сайте — форма отправки, её не используем.'),
        dict(name='Квартиры дома по ЕНИ', method='GET', url='http://darek.kg/Map/GetFlats', data={'propcode': ENI},
             purpose='Список квартир с ЕНИ и адресами для участка или здания.',
             params=[('propcode', 'да', 'ЕНИ, 13 цифр', ENI)], fmt='JSON',
             fields=[('list[].propaddress_id', 'внутренний номер адреса'), ('list[].propcode', 'ЕНИ квартиры'),
                     ('list[].address', 'адрес квартиры')],
             errors='Квартир нет — пустой list.', note='У демо-участка квартир нет — список пуст.'),
    ]


CLOSED = [
    ('Сведения о недвижимости по ЕНИ', 'http://darek.kg/Map/GetPropertyInfo?propcode=…', 'нужен вход — открывается страница авторизации'),
    ('Раздел «Недвижимость» адресного регистра', 'http://address.darek.kg/estate/ — поиск по ЕНИ, коду, адресу',
     'нужен вход («Вы не авторизованы»)'),
    ('Слои «Граница Кыргызстана», «Озёра»', 'GetFeatureInfo', '403 Forbidden — закрыты для внешних запросов'),
    ('Поиск ЕНИ на карте', 'http://darek.kg/Map/SearchPropCodes', 'отвечает {success: true} без данных — формат запроса не разобран'),
    ('Справочник территорий адресного регистра', 'http://address.darek.kg/ajax/ateChildren', 'ошибка 500 — формат запроса не разобран; '
                                                                                              'тот же справочник отдаёт карта (/Map/ateChildren)'),
    ('Поиск по карте', '/wsgi/search.wsgi?query=…, /wsgi/getSearchGeom.wsgi', 'в настройках страницы выключены'),
]


def safe_out(path):
    lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path)[2:])
    if os.path.exists(lock) or os.path.exists(os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))):
        raise SystemExit('документ открыт — закройте и запустите снова')
    if os.path.exists(path):
        from docx import Document
        by = Document(path).core_properties.last_modified_by
        if by and by != 'Сборщик макета':
            raise SystemExit('документ правил %s — не перезаписываю' % by)
    return path


def build(path, layers, recs):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(10.5)

    def shade(cell, color):
        pr = cell._tc.get_or_add_tcPr()
        sh = OxmlElement('w:shd')
        sh.set(qn('w:val'), 'clear')
        sh.set(qn('w:color'), 'auto')
        sh.set(qn('w:fill'), color)
        pr.append(sh)

    def table(cols, rows, widths, mono=()):
        t = doc.add_table(rows=1, cols=len(cols))
        t.style = 'Table Grid'
        t.alignment = WD_TABLE_ALIGNMENT.LEFT
        t.autofit = False
        lay = OxmlElement('w:tblLayout')
        lay.set(qn('w:type'), 'fixed')
        t._tbl.tblPr.append(lay)
        for i, c in enumerate(cols):
            cell = t.rows[0].cells[i]
            run = cell.paragraphs[0].add_run(c)
            run.bold = True
            run.font.size = Pt(9.5)
            shade(cell, 'DCE9F7')
        for row in rows:
            cells = t.add_row().cells
            for i, v in enumerate(row):
                run = cells[i].paragraphs[0].add_run(str(v if v is not None else ''))
                run.font.size = Pt(9)
                if i in mono:
                    run.font.name = 'Consolas'
        for i, w in enumerate(widths):
            t.columns[i].width = Cm(w)
            for cell in t.columns[i].cells:
                cell.width = Cm(w)
        doc.add_paragraph()

    def code(s):
        run = doc.add_paragraph().add_run(s)
        run.font.name = 'Consolas'
        run.font.size = Pt(8.5)
        run.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    def label(s):
        doc.add_paragraph().add_run(s).bold = True

    doc.add_heading('darek.kg (карта ГИСАР и адресный регистр): запросы к данным', 0)
    doc.add_heading('Как обращаться', 1)
    for t, x in GENERAL:
        p = doc.add_paragraph(style='List Bullet')
        p.add_run(t + '. ').bold = True
        p.add_run(x)

    doc.add_heading('Что нужно — где искать', 1)
    table(['Задача', 'Что даёт', 'Запрос', 'Сейчас'],
          [[r['s']['name'], re.split(r'(?<=\.)\s', r['s']['purpose'])[0], r['s']['method'] + ' ' + r['s']['url'].replace('http://', ''),
            r['status']] for r in recs], [4.4, 6.6, 4.6, 1.8], mono=(2,))

    for r in recs:
        s = r['s']
        doc.add_heading(s['name'], 1)
        doc.add_paragraph(s['purpose'])
        label('Запрос')
        code('%s %s' % (s['method'], s['url']))
        if s['params']:
            label('Параметры')
            table(['Параметр', 'Нужен', 'Описание', 'Пример'], [list(p) for p in s['params']], [3.6, 1.6, 6.6, 5.6], mono=(0, 3))
        label('Ответ — %s' % s['fmt'])
        if s['fields']:
            table(['Поле', 'Что значит'], [list(f) for f in s['fields']], [5.2, 12.2], mono=(0,))
        label('Пример запроса')
        code(r['curl'])
        label('Ответ сейчас (%s)' % r['status'])
        code(r['answer'])
        if s['errors'] and s['errors'] != '—':
            label('Ошибки и особенности')
            doc.add_paragraph(s['errors'])
        if s['note']:
            doc.add_paragraph(s['note'])

    doc.add_heading('Слои и поля', 1)
    doc.add_paragraph('По описанию проекта карты (GetProjectSettings). Название слоя пишется в LAYERS, QUERY_LAYERS и FILTER как есть; '
                      'в условии FILTER — имя поля (первый столбец), в ответе GetFeatureInfo поле приходит под подписью (второй столбец), '
                      'если она есть.')
    for L in layers:
        if L['group']:
            continue
        doc.add_heading('%s — %s' % (L['name'], GEOM.get(L['geom'], L['geom'])), 2)
        seen, rows = set(), []
        for n, a, t in L['attrs']:
            if n in seen:
                continue
            seen.add(n)
            rows.append([n, a, t])
        if rows:
            table(['Поле', 'Подпись на карте', 'Тип'], rows, [5.0, 9.6, 2.8], mono=(0,))

    doc.add_heading('Что закрыто или не разобрано', 1)
    table(['Что', 'Запрос', 'Почему'], [list(c) for c in CLOSED], [5.0, 7.0, 5.4], mono=(1,))
    doc.add_heading('Подписи полей Кадастра', 1)
    doc.add_paragraph('ГИСАР подписывает поля сведений участка так: naznacheni — «Форма использования» («Нежилое помещение»), '
                      'propfor_name — «Целевое назначение» («Коммерческое»), propform_name — «Форма собственности», sts — «Тип участка '
                      '(1 — сельхоз)», from_gps — «GPS-съёмка», amnestied — «Амнистирован». Те же поля отдаёт cadastre.kg '
                      '(документация docs/kadastr-api.docx).')
    doc.core_properties.author = doc.core_properties.last_modified_by = 'Сборщик макета'
    doc.save(path)


def main():
    path = safe_out(OUT)
    layers, _ = project_layers()
    recs = []
    for s in sections(layers):
        code, ct, raw = call(s['url'], s['data'], s['method'],
                             ref='http://address.darek.kg/search/' if 'address.darek' in s['url'] else 'http://darek.kg/Map')
        t = text(raw)
        if 'image' in ct:
            ans = 'картинка %s, %d байт' % (ct, len(raw))
        elif s['name'].startswith('Какие слои'):
            ans = 'XML %d КБ; слоёв: %d — %s' % (len(raw) // 1024, len([x for x in layers if not x['group']]),
                                                   ', '.join(x['name'] for x in layers if not x['group']))
        else:
            ans = excerpt(t)
        ok = code == 200 and 'Авторизация' not in t[:3000]
        status = 'отвечает' if ok else 'ошибка %s' % code
        recs.append(dict(s=s, status=status, curl=curl(s['method'], s['url'], s['data']), answer=ans))
        print('  %s: %s' % (s['name'][:55], status))
    build(path, layers, recs)
    print('готово: %s — разделов %d, слоёв %d' % (os.path.relpath(path, ROOT), len(recs), len([x for x in layers if not x['group']])))


if __name__ == '__main__':
    sys.exit(main())
