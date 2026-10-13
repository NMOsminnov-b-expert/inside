# -*- coding: utf-8 -*-
"""ТЗ «Рабочий день осмотрщика» — документ Word из записи графа.

Источник один — запись knowledge/tasks/tz_rabota_osmotrschika.py: правка
ТЗ делается в графе, документ пересобирается этой командой, копии текста
нет. Пункты записи вида «N. Заголовок — текст» становятся разделами,
«N.M. текст» — пунктами раздела.

    python tools/docs/build_tz_osmotrschika.py      # → docs/tz-rabota-osmotrschika.docx
    python tools/docs/build_tz_osmotrschika.py ux   # → docs/razbor-ui-ux-osmotry.docx (разбор UI/UX)
"""
import ast
import io
import os
import re

from docx import Document
from docx.shared import Pt, Cm

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
# Записи и их документы: python build_tz_osmotrschika.py [ux | panel] — разбор
# UI/UX раздела «Осмотры», ТЗ панели осмотров регионала.
DOCS = {
    'tz': ('tz_rabota_osmotrschika.py', 'tz-rabota-osmotrschika.docx'),
    'ux': ('razbor_ui_ux_razdela_osmotry_scenarii_putanica_navigacii_plan.py', 'razbor-ui-ux-osmotry.docx'),
    'panel': ('panel_osmotrov_tz.py', 'tz-panel-osmotrov.docx'),
}
import sys
_key = sys.argv[1] if len(sys.argv) > 1 else 'tz'
SRC = os.path.join(ROOT, 'knowledge', 'tasks', DOCS[_key][0])
OUT = os.path.join(ROOT, 'docs', DOCS[_key][1])


def record():
    tree = ast.parse(io.open(SRC, encoding='utf-8').read())
    data = {}
    for st in tree.body:
        if isinstance(st, ast.Assign) and isinstance(st.targets[0], ast.Name):
            try:
                data[st.targets[0].id] = ast.literal_eval(st.value)
            except ValueError:
                pass
    return data


def main():
    rec = record()
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    for sec in doc.sections:
        sec.left_margin = sec.right_margin = Cm(2)

    title = re.sub(r'^ТЗ:\s*', '', rec['TITLE'])
    doc.add_heading(title[0].upper() + title[1:], level=1)
    p = doc.add_paragraph(f"Дата: {'.'.join(reversed(rec['DATE'].split('-')))}. Основание: {rec['SOURCE']}.")
    p.runs[0].italic = True

    for point in rec['POINTS']:
        sub = re.match(r'^(\d+\.\d+)\.\s+(.*)$', point, re.S)
        head = re.match(r'^(\d+)\.\s+(.*)$', point, re.S)
        if sub:
            para = doc.add_paragraph()
            num = para.add_run(sub.group(1) + '. ')
            num.bold = True
            para.add_run(sub.group(2))
            para.paragraph_format.left_indent = Cm(0.6)
            para.paragraph_format.space_after = Pt(4)
        elif head:
            n, rest = head.group(1), head.group(2)
            name, _, text = rest.partition(' — ')
            doc.add_heading(f'{n}. {name.rstrip(".")}', level=2)
            if text:
                doc.add_paragraph(text[0].upper() + text[1:])
        else:
            doc.add_paragraph(point)

    doc.save(OUT)
    print(OUT)


if __name__ == '__main__':
    main()
