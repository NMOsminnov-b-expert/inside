# -*- coding: utf-8 -*-
"""Документация запросов cadastre.kg для программиста — в Excel.

Просьбы пользователя 14.10.2026: «Выдай мне все, что можно выкопать с
кадастра. И где это искать тоже. С форматами запросов», «подробное описание
куда и как стучаться за данными», «руководство программиста должно быть…
АПИ с описанием запроса и ответа… реальная документация», «Мне удобнее
будет вообще excel». Знания — запись графа
cadastre-kg-otkrytye-zaprosy-portala-chto-mozhno-poluchit.

Сборщик опрашивает портал вживую: каждый открытый запрос выполняется, поля
ответа и примеры берутся из настоящих ответов, состав слоёв — из описания
картосервера. Поэтому документация пересобирается одной командой и не
отстаёт от портала:

    python tools/docs/build_kadastr_spravka.py      # docs/kadastr-api.xlsx

На портал уходят только ЕНИ демо-записи, адрес демо-записи и служебные
запросы (решения пользователя 13.10.2026: ЕНИ и адреса отправлять можно).
Закрытые входом запросы и поиск по физическим лицам не выполняются — только
описаны.
"""
import http.cookiejar
import json
import math
import os
import re
import sys
import urllib.parse
import urllib.request

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, 'docs', 'kadastr-api.xlsx')
BASE = 'https://cadastre.kg/svc-portal/'
WFS = 'map/proxy.do?http://localhost/o2map/services/wfs'
WMS = 'map/proxy.do?http://localhost/o2map/services/wms'
ENI = '1030700400754'          # демо-запись: нежилое, Бишкек, ул. Суюмбаева А., д. 33
STREET_WORD = 'Суюмбаева'
HOUSE = '33'

_jar = http.cookiejar.CookieJar()
_op = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(_jar))
HDR = {'User-Agent': 'Mozilla/5.0', 'Referer': BASE + 'map/main.do', 'X-Requested-With': 'XMLHttpRequest'}


def call(path, data=None, method='GET'):
    """(код, текст) без исключений — документация пишет и неудачи."""
    url = BASE + path
    body = None
    if data is not None and method == 'GET':
        url += ('&' if '?' in url else '?') + urllib.parse.urlencode(data)
    elif data is not None:
        body = urllib.parse.urlencode(data).encode()
    hdr = dict(HDR, **({'Content-Type': 'application/x-www-form-urlencoded'} if body else {}))
    try:
        r = _op.open(urllib.request.Request(url, data=body, headers=hdr, method=method), timeout=45)
        return r.status, r.read().decode('utf-8', 'replace')
    except urllib.error.HTTPError as e:
        return e.code, e.read().decode('utf-8', 'replace')
    except Exception as e:
        return 0, str(e)


def dashed(c):
    return '%s-%s-%s-%s-%s' % (c[0], c[1:3], c[3:5], c[5:9], c[9:13])


def eni_filter(c):
    return ('<Filter><And><PropertyIsEqualTo matchCase="false"><PropertyName>PROPCODE</PropertyName>'
            '<Literal>%s</Literal></PropertyIsEqualTo></And></Filter>' % dashed(c))


def merc(lon, lat):
    R = 6378137.0
    return R * math.radians(lon), R * math.log(math.tan(math.pi / 4 + math.radians(lat) / 2))


def curl(method, path, data=None):
    url = BASE + path
    q = "curl -s -b jar.txt -c jar.txt -H 'Referer: %smap/main.do'" % BASE
    if method == 'GET':
        return "%s '%s%s'" % (q, url, (('&' if '?' in url else '?') + urllib.parse.urlencode(data)) if data else '')
    return "%s -X POST '%s' %s" % (q, url, ' '.join("--data-urlencode '%s=%s'" % (k, v) for k, v in data.items()))


def plain(html):
    t = re.sub(r'<script.*?</script>|<style.*?</style>', ' ', html, flags=re.S)
    return re.sub(r'\s+', ' ', re.sub(r'<[^>]+>', ' ', t)).strip()


def excerpt(text, n=1500):
    t = text.strip()
    if t.startswith('{'):
        try:
            def cut(o):
                if isinstance(o, dict):
                    return {k: ('…координаты сокращены…' if k in ('coordinates', 'wkt') else cut(v)) for k, v in o.items()}
                if isinstance(o, list):
                    return [cut(x) for x in o[:3]] + (['…ещё %d' % (len(o) - 3)] if len(o) > 3 else [])
                return o
            t = json.dumps(cut(json.loads(t)), ensure_ascii=False, indent=1)
        except ValueError:
            pass
    elif t.startswith('<?xml'):
        t = re.sub(r'\s+', ' ', t)
    elif t.startswith('<') or '<html' in t[:300].lower():
        t = '[HTML, текст без разметки] ' + plain(t)
    return t[:n] + ('…' if len(t) > n else '')


def flatten(o, path='', out=None):
    """Поля JSON-ответа: [(путь, тип, пример)]; у массива берётся первый элемент."""
    out = [] if out is None else out
    if isinstance(o, dict):
        for k, v in o.items():
            flatten(v, '%s.%s' % (path, k) if path else k, out)
    elif isinstance(o, list):
        out.append((path, 'массив', '%d эл.' % len(o)))
        if o and not path.endswith('coordinates'):
            flatten(o[0], path + '[]', out)
    else:
        typ = {bool: 'логическое', int: 'число', float: 'число', str: 'строка', type(None): 'null'}.get(type(o), type(o).__name__)
        out.append((path, typ, '' if o is None else str(o)[:80]))
    return out


# Описание полей ответов — по последней части пути; неясное так и помечено.
FIELD = {
    'searchAddrList': 'список элементов уровня справочника', 'addrCd': 'код элемента справочника адресов (32 знака) — передавать на следующий уровень',
    'addrName': 'название элемента (область, район, айыл аймак, село, улица)', 'type': 'служебное, обычно null', 'addr': 'адрес одной строкой',
    'searchAddrDetail': 'сведения об адресе участка', 'propCode': 'ЕНИ, 13 цифр', 'lndPclId': 'внутренний код участка (32 знака)',
    'id': 'внутренний код объекта', 'oblast': 'код области', 'region': 'код района', 'ao': 'код айыл аймака', 'village': 'код села',
    'street': 'код улицы', 'buildingNumb': 'номер дома', 'flatNumb': 'номер квартиры', 'uchNo': 'номер участка', 'mainRegion': 'код района',
    'searchMode': 'служебное: режим поиска', 'locale': 'язык ответа', 'pageInfo': 'служебное: постраничность', 'page': 'страница',
    'pageListSize': 'записей на странице', 'totalRows': 'всего записей', 'viewType': 'служебное', 'className': 'служебное',
    'hasValidateError': 'служебное', 'searchCondition': 'служебное: условие поиска', 'searchKeyword': 'служебное', 'mode': 'служебное',
    'extent': 'рамка', 'minx': 'наименьшая x (в C1 — долгота)', 'miny': 'наименьшая y (в C1 — широта)', 'maxx': 'наибольшая x',
    'maxy': 'наибольшая y', 'feature': 'фигура', 'wkt': 'геометрия в WKT (местная система координат — лист «Общие сведения»)',
    'features': 'объекты слоя', 'geometry': 'геометрия объекта', 'coordinates': 'координаты (долгота, широта при srsname=EPSG:4326)',
    'properties': 'сведения объекта — поля слоя (лист «Слои картосервера»)', 'totalFeatures': 'всего объектов', 'numberMatched': 'найдено',
    'numberReturned': 'возвращено', 'crs': 'система координат ответа', 'nameRu': 'название по-русски', 'nameKg': 'название по-киргизски',
    'comment': 'комментарий', 'userLroCode': 'служебное', 'fromIndex': 'служебное', 'toIndex': 'служебное', 'autoIndex': 'служебное',
    'maxPaginationSize': 'служебное', 'name': 'название', 'properties.name': 'название системы координат',
    'terEd': 'код территориальной единицы', 'terOb': 'код территориального объекта', 'regionCode': 'код района',
    'termDate': 'дата прекращения — смысл не уточнён', 'removeDate': 'дата снятия — смысл не уточнён', 'restricted': 'отметка ограничения — смысл не уточнён',
    'updated': 'служебное',
    # Сведения о владельцах и правах — без входа пустые (null); личные данные, не используем.
    'owners': 'владельцы — без входа пусто; не используем', 'rightTbl': 'права — без входа пусто', 'ownerName': 'владелец — без входа пусто; не используем',
    'ownerId': 'владелец — без входа пусто; не используем', 'tin': 'ИНН владельца — без входа пусто; не используем',
    'birthday': 'личные данные — без входа пусто; не используем', 'idDocIssue': 'личные данные — без входа пусто; не используем',
    'idDocNo': 'личные данные — без входа пусто; не используем', 'telNo': 'личные данные — без входа пусто; не используем',
}
LAYER_FIELD = {
    'PROPCODE': 'ЕНИ с дефисами (1-03-07-0040-0754)', 'DOMCODE': 'номер дома',
    'USLCODE': 'код — смысл не ясен (у участка в Канте 42214000, в Бишкеке пусто)', 'PRIMECHANI': 'примечание',
    'THE_GEOM': 'контур — мультиполигон', 'ADDRESS': 'адрес коротко', 'FULL_ADDRESS': 'адрес полностью', 'FULL_ADDRE': 'адрес полностью',
    'NAZNACHENI': 'назначение по Кадастру', 'DATEINS': 'дата внесения записи', 'DATEUPD': 'дата изменения записи',
    'PROPFOR': 'вид использования, код', 'PROPFOR_NAME': 'вид использования', 'PROPFOR_NA': 'вид использования',
    'PROPFORM_NAME': 'форма собственности', 'PROPFORM_N': 'форма собственности',
    'OID_LRO': 'служебный код регистрационного органа', 'CODE_LRO': 'код регистрационного органа (совпадает с началом ЕНИ: 103, 704)',
    'AMNESTIED': 'отметка амнистии (по названию поля — узаконен по амнистии)', 'FROM_GPS': 'координаты сняты по GPS',
    'STS': 'состояние — смысл не ясен', 'OID': 'служебный номер объекта', 'GID': 'служебный номер объекта', 'OID_OLD': 'прежний служебный номер',
    'BUILD_TYPE': 'тип строения',
}


def P(name, req, typ, desc, ex):
    return (name, req, typ, desc, ex)


def endpoints(ctx):
    """Описание запросов. ctx — коды, полученные по ходу (область, улица, рамка)."""
    bish, st, sid, bbox = ctx.get('bish', ''), ctx.get('street', ''), ctx.get('sid', ''), ctx.get('bbox', '')
    return [
        dict(id='E1', group='Служебные', name='Сессия и список областей', method='GET', path='map/main.do', data=None,
             purpose='Страница карты портала. Выдаёт куку сессии (JSESSIONID) — её передавать во всех запросах; в разметке — '
                     'список областей с кодами для справочника адресов (<select id="oblast">: value — код, текст — название).',
             params=[], fmt='HTML', html='Области: <select id="oblast"> → <option value="код">название</option>.',
             errors='—', used='kernel/cadastreGeo.js — oblastList'),
        dict(id='A1', group='Справочник адресов', name='Уровень справочника адресов', method='GET', path='map/searchAddr.do',
             data={'mode': 'street', 'oblast': bish, 'region': '', 'ao': '', 'village': bish, 'street': ''},
             purpose='Справочник адресов каскадом: область → район → айыл аймак → село → улица. Каждый вызов отдаёт список следующего '
                     'уровня; код выбранного элемента (addrCd) передаётся в следующий вызов.',
             params=[P('mode', 'да', 'строка', 'какой уровень вернуть: region | ao | village | street', 'street'),
                     P('oblast', 'да', 'код 32 знака', 'код области (из E1)', bish or '<код>'),
                     P('region', 'для ao и ниже', 'код', 'код района (из mode=region)', ''),
                     P('ao', 'для village', 'код', 'код айыл аймака (из mode=ao)', ''),
                     P('village', 'для street', 'код', 'код села; у Бишкека и Оша — код самой «области»: улицы идут сразу после неё',
                       bish or '<код>'),
                     P('street', 'нет', 'код', 'не используется', '')],
             fmt='JSON', errors='Неверные коды — пустой список.', used='kernel/cadastreGeo.js — eniByAddress (каскад)'),
        dict(id='B1', group='Поиск', name='ЕНИ по адресу / адрес по ЕНИ (список)', method='POST', path='map/searchAddrResult.do',
             data={'searchCondition': 'ADDR', 'inputPropCode': '', 'propCode': '', 'oblast': bish, 'region': '', 'ao': '', 'village': bish,
                   'street': st, 'buildingNumb': HOUSE, 'flatNumb': '', 'uchNo': '', 'page': 1, 'pageListSize': 12},
             purpose='Поиск объектов недвижимости: по адресу (коды из A1 и номер дома) или по ЕНИ. Возвращает список ЕНИ с адресами; '
                     'участок и квартиры дома — отдельными строками.',
             params=[P('searchCondition', 'да', 'строка', 'ADDR — по адресу; ENI — по коду', 'ADDR'),
                     P('propCode', 'для ENI', 'строка', 'ЕНИ, 13 цифр', ENI),
                     P('oblast, region, ao, village, street', 'для ADDR', 'коды', 'коды из A1 (у Бишкека и Оша village = код области)', ''),
                     P('buildingNumb', 'нет', 'строка', 'номер дома', HOUSE), P('flatNumb', 'нет', 'строка', 'номер квартиры', ''),
                     P('uchNo', 'нет', 'строка', 'номер участка', ''), P('inputPropCode', 'нет', 'строка', 'пусто', ''),
                     P('page', 'да', 'число', 'страница', '1'), P('pageListSize', 'да', 'число', 'записей на странице', '12')],
             fmt='HTML', html='Каждая строка списка: атрибут data-gab-item-key — ЕНИ (13 цифр, иногда с хвостом); адрес — в <dd>.',
             errors='Ничего не найдено — список пуст.', used='tools/cadastre/proxy.py — /eni; kernel/cadastreGeo.js — eniByAddress'),
        dict(id='B2', group='Поиск', name='Адрес и внутренние коды по ЕНИ', method='GET', path='map/searchAddrDetail.do', data={'propCode': ENI},
             purpose='Адрес участка одной строкой с районом и внутренние коды портала: участка, адреса, района, улицы.',
             params=[P('propCode', 'да', 'строка', 'ЕНИ, 13 цифр без дефисов', ENI)], fmt='JSON',
             errors='Нет такого ЕНИ — searchAddrDetail пустой или null.', used='не используется'),
        dict(id='B3', group='Поиск', name='Адресный реестр «Дарек»', method='POST', path='map/darekSearchAddrResult.do',
             data={'searchCondition': 'ENI', 'propCode': ENI, 'page': 1, 'pageListSize': 10},
             purpose='Поиск в адресном реестре «Дарек» — по ЕНИ или по адресу. Адрес в написании реестра.',
             params=[P('searchCondition', 'да', 'строка', 'ENI или адресный режим', 'ENI'), P('propCode', 'для ENI', 'строка', 'ЕНИ, 13 цифр', ENI),
                     P('oblast, region, ao, village, terEd, terOb, street', 'для адреса', 'коды', 'коды справочника; terEd, terOb — '
                       'территориальная единица и объект', ''),
                     P('buildingNumb, flatNumb, uchNo', 'нет', 'строка', 'дом, квартира, участок', ''),
                     P('page, pageListSize', 'да', 'число', 'постраничность', '1, 10')],
             fmt='HTML', html='Как в B1: data-gab-item-key — ЕНИ, адрес — в <dd>.', errors='—', used='не используется'),
        dict(id='B4', group='Поиск', name='Поиск улиц по слову', method='POST', path='map/totalSearchStreetResult.do',
             data={'page': 1, 'pageListSize': 10, 'searchCondition': 'TOTAL', 'searchKeyword': STREET_WORD},
             purpose='Улицы, в названии которых есть слово. id улицы — для линии улицы (C2).',
             params=[P('searchKeyword', 'да', 'строка', 'слово из названия улицы', STREET_WORD), P('searchCondition', 'да', 'строка', 'TOTAL', 'TOTAL'),
                     P('page, pageListSize', 'да', 'число', 'постраничность', '1, 10')],
             fmt='HTML', html='Строка улицы: data-gab-item-key — id улицы (32 знака), текст — название.', errors='—', used='не используется'),
        dict(id='B5', group='Поиск', name='Полнотекстовый поиск адресов', method='POST', path='map/totalSearchAddrResult.do',
             data={'page': 1, 'pageListSize': 10, 'searchCondition': 'TOTAL', 'searchKeyword': STREET_WORD},
             purpose='Поиск адресов по слову. С простым словом отвечает ошибкой: страница портала шлёт параметр qry — запрос в формате их '
                     'поискового движка; формат не разобран.',
             params=[P('searchKeyword', 'да', 'строка', 'слово', STREET_WORD), P('searchCondition', 'да', 'строка', 'TOTAL', 'TOTAL'),
                     P('qry', 'вместо слова', 'строка', 'запрос поискового движка — формат не разобран', ''),
                     P('page, pageListSize', 'да', 'число', 'постраничность', '1, 10')],
             fmt='HTML', html='—', errors='«Произошла ошибка во время выполнения программы»', used='не используется'),
        dict(id='C1', group='Геометрия', name='Рамка участка по ЕНИ', method='GET', path='map/getFeatureExtent.do', data={'propCode': ENI},
             purpose='Рамка (наименьшие и наибольшие долгота и широта) участка.',
             params=[P('propCode', 'да', 'строка', 'ЕНИ, 13 цифр без дефисов', ENI)], fmt='JSON',
             errors='Нет участка — {"extent":null}.', used='kernel/cadastreGeo.js — geoByEni'),
        dict(id='C2', group='Геометрия', name='Линия улицы', method='GET', path='map/getStreetFeatureWkt.do', data={'id': sid},
             purpose='Геометрия улицы по её id (из B4).', params=[P('id', 'да', 'строка 32 знака', 'id улицы из B4', sid)], fmt='JSON',
             errors='—', used='не используется'),
        dict(id='C3', group='Геометрия', name='Области без названий', method='GET', path='map/feature.do', data={'geoId': 1},
             purpose='Многоугольники по номеру: 1–4 — области 8–14 км в Бишкеке (вероятно, районы города), 5 и 6 — 37–40 км к западу и '
                     'востоку (вероятно, соседние районы Чуйской области); с 7 — пусто. Названий нет.',
             params=[P('geoId', 'да', 'число', 'номер фигуры, 1–6', '1')], fmt='JSON', errors='Номер без фигуры — пустой ответ.',
             used='не используется'),
        dict(id='D1', group='Картосервер (WFS)', name='Описание сервиса', method='POST', path=WFS,
             data={'service': 'WFS', 'version': '1.1.0', 'request': 'GetCapabilities'},
             purpose='Список слоёв картосервера: ENI_AREA, ENI_WGS84, ENI_FACT_VALID (участки), ENI_BUILDINGS (строения).',
             params=[P('service', 'да', 'строка', 'WFS', 'WFS'), P('version', 'да', 'строка', '1.1.0', '1.1.0'),
                     P('request', 'да', 'строка', 'GetCapabilities', 'GetCapabilities')], fmt='XML',
             html='Имена слоёв — в <Name>.', errors='—', used='—'),
        dict(id='D2', group='Картосервер (WFS)', name='Состав полей слоя', method='POST', path=WFS,
             data={'service': 'WFS', 'version': '1.1.0', 'request': 'DescribeFeatureType', 'typename': 'ENI_AREA'},
             purpose='Поля слоя и их типы (XSD). Состав всех слоёв — лист «Слои картосервера».',
             params=[P('request', 'да', 'строка', 'DescribeFeatureType', 'DescribeFeatureType'),
                     P('typename', 'да', 'строка', 'слой', 'ENI_AREA'), P('service, version', 'да', 'строка', 'WFS, 1.1.0', '')], fmt='XML',
             html='Поля — элементы <element name="…" type="…">.', errors='Нет слоя — исключение «is not exist».', used='—'),
        dict(id='D3', group='Картосервер (WFS)', name='Участок по ЕНИ: все сведения и контур', method='POST', path=WFS,
             data={'service': 'WFS', 'version': '1.1.0', 'request': 'GetFeature', 'typename': 'ENI_AREA', 'srsname': 'EPSG:4326',
                   'outputFormat': 'application/json', 'filter': eni_filter(ENI)},
             purpose='Главный запрос: сведения участка (назначение, вид использования, форма собственности, адрес, даты и др.) и контур.',
             params=[P('service, version, request', 'да', 'строка', 'WFS, 1.1.0, GetFeature', ''),
                     P('typename', 'да', 'строка', 'ENI_AREA | ENI_WGS84 | ENI_FACT_VALID | ENI_BUILDINGS', 'ENI_AREA'),
                     P('filter', 'да (или bbox)', 'XML (OGC Filter)', 'отбор по ЕНИ; ЕНИ — с дефисами', eni_filter(ENI)),
                     P('propertyname', 'нет', 'строка', 'поля через запятую, иначе все', 'PROPCODE,NAZNACHENI,THE_GEOM'),
                     P('srsname', 'нет', 'строка', 'система координат ответа; EPSG:4326 — долгота/широта', 'EPSG:4326'),
                     P('outputFormat', 'нет', 'строка', 'application/json — GeoJSON; без него — GML (XML)', 'application/json'),
                     P('maxFeatures', 'нет', 'число', 'предел числа объектов', '')], fmt='JSON (GeoJSON)',
             errors='Нет участка — пустой features.', used='kernel/cadastreGeo.js — контур, назначение (экран осмотрщика)'),
        dict(id='D4', group='Картосервер (WFS)', name='Участки или строения в рамке', method='POST', path=WFS,
             data={'service': 'WFS', 'version': '1.1.0', 'request': 'GetFeature', 'typename': 'ENI_WGS84', 'srsname': 'EPSG:4326',
                   'outputFormat': 'application/json', 'propertyname': 'PROPCODE,ADDRESS,NAZNACHENI,THE_GEOM', 'maxFeatures': '5', 'bbox': bbox},
             purpose='Все объекты слоя в прямоугольнике — соседние участки, строения на участке.',
             params=[P('bbox', 'да (или filter)', 'строка', 'x1,y1,x2,y2,EPSG:3857 — метры', bbox),
                     P('typename', 'да', 'строка', 'ENI_WGS84 — участки; ENI_BUILDINGS — строения', 'ENI_WGS84'),
                     P('propertyname, srsname, outputFormat, maxFeatures', 'нет', '', 'как в D3', '')], fmt='JSON (GeoJSON)',
             errors='—', used='kernel/cadastreLayersGL.js — участки и строения вокруг на картах'),
        # Параметры — после «?» в самом адресе за прокси: с «&» прокси портала отвечает 500.
        dict(id='D5', group='Картосервер (WMS)', name='Слои картинкой', method='GET', path=WMS + '?SERVICE=WMS&REQUEST=GetCapabilities',
             data=None,
             purpose='Те же четыре слоя растровой картинкой. GetMap: LAYERS, BBOX, WIDTH, HEIGHT, SRS, FORMAT.',
             params=[P('SERVICE', 'да', 'строка', 'WMS', 'WMS'), P('REQUEST', 'да', 'строка', 'GetCapabilities | GetMap', 'GetCapabilities'),
                     P('LAYERS, BBOX, WIDTH, HEIGHT, SRS, FORMAT', 'для GetMap', '', 'параметры картинки', '')], fmt='XML / картинка',
             html='—', errors='—', used='не используется'),
        dict(id='E2', group='Служебные', name='Настройки карты портала', method='GET', path='map/mapConfig.do', data=None,
             purpose='JavaScript-настройки карты: подложки (OSM, спутник Google), слои ENI_AREA и ENI_BUILDINGS; группы легенды «Границы», '
                     '«Гидрография», «Дороги» — только названия и стили, слоёв за ними нет (лист «Чего нет»).',
             params=[], fmt='JavaScript', html='—', errors='—', used='—'),
    ]


GENERAL = [
    ('Адрес', 'https://cadastre.kg/svc-portal/ — все пути запросов от этого адреса.'),
    ('Сессия', 'Первым — GET map/main.do (E1): портал ставит куку JSESSIONID. Её передавать во всех запросах.'),
    ('Заголовки', 'Referer: https://cadastre.kg/svc-portal/map/main.do. Для POST — Content-Type: application/x-www-form-urlencoded. '
                  'X-Requested-With: XMLHttpRequest — как у страницы портала.'),
    ('Кодировка', 'UTF-8 во всех ответах.'),
    ('ЕНИ', '13 цифр без дефисов (1030700400754) — в map/*.do; с дефисами (1-03-07-0040-0754) — в отборе картосервера (поле PROPCODE).'),
    ('Координаты', 'C1 и GeoJSON картосервера при srsname=EPSG:4326 — долгота/широта. Отбор bbox картосервера — EPSG:3857 (метры). WKT из '
                   'C2 и C3 — в местной системе: по замеру 14.10.2026 совпадает с «Kyrg-06, зона 3» (центральный меридиан 74°31′ в. д., '
                   'ложное восточное смещение 3 300 000 м; ул. Суюмбаева в Бишкеке — x ≈ 3 308 480, y ≈ 4 765 050). Номер системы '
                   'портал не называет — перед расчётами сверить.'),
    ('Картосервер', 'O2Map за прокси портала: map/proxy.do?http://localhost/o2map/services/wfs (и /wms, /wmts). Слоёв четыре: ENI_AREA, '
                    'ENI_WGS84, ENI_FACT_VALID — участки; ENI_BUILDINGS — строения.'),
    ('Ошибки', 'Закрытое входом — ответ 200 с текстом «You do not have access privileges»; нет участка — {"extent":null} или пустой '
               'features; нет слоя — исключение «is not exist»; сбой поиска — HTML «Произошла ошибка во время выполнения программы».'),
    ('Из браузера', 'Портал не разрешает запросы с других сайтов (нет Access-Control-Allow-Origin) — нужен посредник. В макете: Live '
                    'Server проксирует /cadastre → https://cadastre.kg/svc-portal/map/ (.vscode/settings.json); запасной — '
                    'python tools/cadastre/proxy.py (http://127.0.0.1:8787/portal/<путь>).'),
    ('Порядок', 'Участок по ЕНИ: E1 → C1 (рамка) → D3 (сведения и контур). ЕНИ по адресу: E1 (код области) → A1 (каскад до улицы) → '
                'B1 (searchCondition=ADDR). Адрес по ЕНИ: B2 или B1 (searchCondition=ENI). Соседи на карте: D4.'),
    ('Статус', 'Внутренние запросы общедоступной карты, а не публичный API: без описания и обещаний, портал может поменять их без '
               'предупреждения. Решение пользователя 14.10.2026: если наработки пойдут в дело — запросить официальный доступ.'),
    ('Не делаем', 'Поиск по физическим лицам и владельцы — личные данные, не используются; закрытое входом не обходится; капча не обходится.'),
]

CLOSED = [
    ('Поиск по юридическим лицам — участки организации', 'map/searchCorpOwnerResult.do', 'POST',
     'companyName — наименование; companyTypeId — организационно-правовая форма (коды из списка на странице карты); tin — ИНН; docNo; '
     'page; pageListSize=5', 'нужен вход: без него «нет прав доступа»', 'полезно (пользователь 14.10.2026) — нужен вход на портал или официальный доступ'),
    ('Техническая инвентаризация участка', 'map/getInventoryForm.do', 'POST', 'propCode', 'нужен вход', 'не нужно — инвентаризация предоставлена'),
    ('Техническая инвентаризация строения', 'map/getInventoryBuildForm.do', 'GET', 'propCode, lndPclId, unitNumb', 'нужен вход', 'не нужно'),
    ('Владельцы участка', 'map/getInventoryOwnerDetail.do', 'POST', 'lndPclId, propCode, addrKey, addr', 'нужен вход; личные данные', 'не нужно'),
    ('Поиск по физическим лицам', 'map/searchIndvOwnerResult.do', 'POST', 'фамилия, имя, отчество, ИНН, документ', 'личные данные', 'не используется'),
    ('Сведения об участке для вошедших', 'aproc/parcel/searchAddressList.do', 'POST', 'parcelCode, searchCondition=PARCEL, page, pageListSize',
     'только вошедшим', 'не используется'),
    ('Сделки с муниципальной недвижимостью', 'app/marketReport.do', 'POST (форма)',
     'startDate, endDate, kuplya / arenda / polzovanie, oblast, region, ao, village, terOb, captcha',
     'открыто человеку, каждый запрос — с капчей', 'автоматически нельзя; при нужде — официальный доступ'),
    ('Выписки и отчёты', 'app/civilPage.do; app/reports.do', '—', '—', 'платные услуги портала', 'не изучались'),
]

ABSENT = [
    ('Границы (областей, районов, сёл)', 'нет слоя на картосервере; в C3 — шесть фигур без названий вокруг Бишкека', 'OpenStreetMap'),
    ('Дороги, железные дороги', 'нет слоя (ROADS, ROADM10, ROADM2, RAILW — «is not exist»); в коде страницы — только стили', 'OpenStreetMap'),
    ('Гидрография', 'нет слоя (HYDRO, RIVER, LAKE — «is not exist»)', 'OpenStreetMap'),
    ('Кадастровые зоны, кварталы', 'нет слоя (ZONE, BLOCK, QUARTER — «is not exist»)', '—'),
    ('Улицы слоем', 'нет слоя; линия отдельной улицы — C2', 'OpenStreetMap'),
    ('Площадь участка, правообладатели, обременения, кадастровая стоимость', 'нет в полях слоёв', 'официальный доступ'),
]


def run():
    """Выполнить запросы, собрать строки листов."""
    _op.open(urllib.request.Request(BASE + 'map/main.do', headers=HDR), timeout=30).read()
    ctx = {}
    code, html = call('map/main.do')
    m = re.search(r'<select[^>]*id="oblast"[\s\S]*?</select>', html)
    obl = re.findall(r'value="([0-9A-F]{32})">\s*([^<]+?)\s*<', m.group(0) if m else '')
    ctx['bish'] = next((c for c, n in obl if 'Бишкек' in n), '')
    c2, js = call('map/searchAddr.do', {'mode': 'street', 'oblast': ctx['bish'], 'region': '', 'ao': '', 'village': ctx['bish'], 'street': ''})
    try:
        ctx['street'] = next((x['addrCd'] for x in json.loads(js).get('searchAddrList', []) if STREET_WORD in x['addrName']), '')
    except ValueError:
        pass
    c3, sh = call('map/totalSearchStreetResult.do', {'page': 1, 'pageListSize': 10, 'searchCondition': 'TOTAL', 'searchKeyword': STREET_WORD}, 'POST')
    ms = re.search(r'data-gab-item-key="([^"]+)"', sh)
    ctx['sid'] = ms.group(1) if ms else ''
    c4, ext = call('map/getFeatureExtent.do', {'propCode': ENI})
    try:
        e = json.loads(ext)['extent']
        x1, y1 = merc(e['minx'], e['miny'])
        x2, y2 = merc(e['maxx'], e['maxy'])
        ctx['bbox'] = '%d,%d,%d,%d,EPSG:3857' % (x1 - 60, y1 - 60, x2 + 60, y2 + 60)
    except (ValueError, KeyError, TypeError):
        ctx['bbox'] = ''

    overview, params, answers, examples = [], [], [], []
    for ep in endpoints(ctx):
        code, txt = call(ep['path'], ep['data'], ep['method'])
        txt_ex = ('Области (код — название): ' + '; '.join('%s — %s' % (c, n) for c, n in obl)) if ep['id'] == 'E1' else excerpt(txt)
        ok = code == 200 and 'access privileges' not in txt and 'Произошла ошибка' not in txt and 'is not exist' not in txt
        overview.append([ep['id'], ep['group'], ep['name'], ep['purpose'], ep['method'], ep['path'], 'открыто', ep['used'],
                         ('отвечает (%d)' % code) if ok else ('ошибка (%d): %s' % (code, plain(txt)[:90]))])
        for pr in ep['params']:
            params.append([ep['id'], ep['name']] + list(pr))
        parsed = False
        if txt.strip().startswith('{'):
            try:
                seen = set()
                for path, typ, ex in flatten(json.loads(txt)):
                    key = path.replace('[]', '')
                    if key in seen:
                        continue
                    seen.add(key)
                    last = path.split('.')[-1].replace('[]', '')
                    desc = LAYER_FIELD.get(last) if '.properties.' in path else None
                    answers.append([ep['id'], ep['name'], path, typ, desc or FIELD.get(last, ''), ex])
                parsed = True
            except ValueError:
                pass
        if not parsed:
            answers.append([ep['id'], ep['name'], '(%s)' % ep['fmt'], ep['fmt'], ep.get('html', ''), ''])
        examples.append([ep['id'], ep['name'], curl(ep['method'], ep['path'], ep['data']), txt_ex, ep['errors']])
        print('  %s %s: %s' % (ep['id'], ep['name'], overview[-1][-1][:70]))
    return overview, params, answers, examples


def layers():
    rows, samples = [], {}
    for lay in ('ENI_AREA', 'ENI_WGS84', 'ENI_FACT_VALID', 'ENI_BUILDINGS'):
        st, txt = call(WFS, {'service': 'WFS', 'version': '1.1.0', 'request': 'GetFeature', 'typename': lay, 'srsname': 'EPSG:4326',
                             'outputFormat': 'application/json', 'maxFeatures': '1', 'filter': eni_filter(ENI)}, 'POST')
        try:
            for f in (json.loads(txt).get('features') or [])[:1]:
                samples.update({(lay, k): v for k, v in f['properties'].items()})
                samples[(lay, 'THE_GEOM')] = 'мультиполигон'
        except ValueError:
            pass
        st, d = call(WFS, {'service': 'WFS', 'version': '1.1.0', 'request': 'DescribeFeatureType', 'typename': lay}, 'POST')
        for name, typ in re.findall(r'<(?:xsd:|xs:)?element[^>]*name="([^"]+)"[^>]*type="([^"]+)"', d):
            if name == lay:
                continue
            v = samples.get((lay, name), '')
            rows.append([lay, name, typ.split(':')[-1], LAYER_FIELD.get(name, ''), '' if v in (None, -9999) else str(v)[:80]])
    return rows


def safe_out(path):
    """Правило пользователя 14.10.2026: открытый или правленный человеком файл не перезаписывается."""
    lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))
    if os.path.exists(lock):
        raise SystemExit('файл открыт — закройте и запустите снова')
    if os.path.exists(path):
        from openpyxl import load_workbook
        by = load_workbook(path, read_only=True).properties.lastModifiedBy
        if by and by != 'Сборщик макета':
            raise SystemExit('файл правил %s — не перезаписываю' % by)
    return path


def build(path, overview, params, answers, examples, lay_rows):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head = PatternFill('solid', fgColor='DCE9F7')
    bad = PatternFill('solid', fgColor='FDECEC')
    wrap = Alignment(vertical='top', wrap_text=True)
    mono = Font(name='Consolas', size=9)

    def sheet(title, cols, rows, widths, mono_cols=(), mark=None, freeze='A2'):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for i in range(1, len(cols) + 1):
            c = ws.cell(row=1, column=i)
            c.fill = head
            c.font = Font(bold=True, color='1F3A5F')
            c.alignment = Alignment(vertical='center', wrap_text=True)
        for r in rows:
            ws.append(r)
            if mark and mark(r):
                for i in range(1, len(cols) + 1):
                    ws.cell(row=ws.max_row, column=i).fill = bad
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
                if c.column in mono_cols:
                    c.font = mono
        ws.freeze_panes = freeze
        ws.auto_filter.ref = ws.dimensions

    sheet('Общие сведения', ['Тема', 'Как устроено'], [list(x) for x in GENERAL], [16, 120])
    sheet('Запросы', ['Код', 'Группа', 'Запрос', 'Назначение', 'Метод', 'Путь (от https://cadastre.kg/svc-portal/)', 'Доступ', 'В макете',
                      'Проверено сейчас'], overview, [6, 18, 28, 60, 8, 40, 10, 34, 30], mono_cols=(6,),
          mark=lambda r: str(r[-1]).startswith('ошибка'), freeze='C2')
    sheet('Параметры', ['Код', 'Запрос', 'Параметр', 'Обязателен', 'Тип', 'Описание', 'Пример'], params, [6, 28, 26, 14, 14, 56, 44],
          mono_cols=(3, 7), freeze='C2')
    sheet('Ответы', ['Код', 'Запрос', 'Поле (путь в ответе)', 'Тип', 'Описание', 'Пример значения'], answers, [6, 28, 40, 12, 56, 40],
          mono_cols=(3, 6), freeze='C2')
    sheet('Примеры', ['Код', 'Запрос', 'Запрос командой (curl; кука — из E1)', 'Ответ сейчас (геометрия сокращена)', 'Ошибки и особенности'],
          examples, [6, 26, 70, 90, 36], mono_cols=(3, 4), freeze='C2')
    sheet('Слои картосервера', ['Слой', 'Поле', 'Тип', 'Описание', 'Пример (ЕНИ %s)' % dashed(ENI)], lay_rows, [18, 18, 14, 56, 40],
          mono_cols=(2,))
    sheet('Закрытые', ['Что', 'Путь', 'Метод', 'Параметры', 'Почему закрыто', 'Для нас'], [list(x) for x in CLOSED],
          [36, 34, 10, 56, 30, 40], mono_cols=(2,))
    sheet('Чего нет', ['Данные', 'Что на портале', 'Где взять'], [list(x) for x in ABSENT], [40, 70, 24])
    wb.properties.creator = wb.properties.lastModifiedBy = 'Сборщик макета'
    wb.save(path)


def main():
    path = safe_out(OUT)
    overview, params, answers, examples = run()
    lay_rows = layers()
    build(path, overview, params, answers, examples, lay_rows)
    print('готово: %s — запросов %d, параметров %d, полей ответов %d, полей слоёв %d'
          % (os.path.relpath(path, ROOT), len(overview), len(params), len(answers), len(lay_rows)))


if __name__ == '__main__':
    sys.exit(main())
