# -*- coding: utf-8 -*-
"""Предложения по расширению шаблонов машин ТС — книга на согласование.

Просьба пользователя 06.10.2026: «Давай потом повторно расширим шаблоны. Найди
что есть что и предложи. Даже экзотические варианты». Найденное (поиск общими
словами о технике, у каждого предложения — источник) разложено по листам:
новые модули, новые сочетания «модуль + носитель», составные наборы, обиходные
названия, сомнительное. Столбец «Замечание» — проверка Claude по модели
карточки (одна карточка — одна машина и т.п.), «Решение» — пустой, для
пользователя. В карточку предложения не вносятся до решения.

Коды носителей — как в tools/data/ts_templates.py (CARRIER).

    python tools/docs/build_predlozheniya_shablonov_ts.py   # docs/predlozheniya-shablonov-ts.xlsx
"""
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
OUT = os.path.join(ROOT, 'docs', 'predlozheniya-shablonov-ts.xlsx')
sys.path.insert(0, os.path.join(ROOT, 'tools', 'data'))
from ts_templates import CARRIER  # noqa: E402

KK = 'https://www.kolesa.ru/article/podyomnye-zapravochnye-i-ne-tolko-kakim-byl-sovetskij-spetstransport-dlya-aeroportov'
OKPD = 'https://ppt.ru/classifier/okpd/29-10-59'
NICK = 'https://www.kolesa.ru/article/zapor-utyug-i-skotovoz-narodnye-prozvishha-zaporozhtsa-uaza-i-drugih'
TWO = 'Две машины — две карточки: в макете одна карточка — одна машина. Раскладывается на два шаблона.'

MODULES = [
    # группа, модуль, носители, пример, источник, замечание
    ('Аэродромные (новая)', 'Трап самоходный пассажирский', 'ГА', 'СПТ-20А на УАЗ-451Д, СПТ-21 на УАЗ-452Д', KK, ''),
    ('Аэродромные (новая)', 'Багажный транспортёр (перегружатель)', 'ГА', 'АТ-4М на ГАЗ-69, АТ-6 на УАЗ-452Д', KK, ''),
    ('Аэродромные (новая)', 'Автолифт бортового питания', 'ГА', 'АЛ-2 на ГАЗ-53А, АЛ-3А на ЗИЛ-130Г', KK, ''),
    ('Аэродромные (новая)', 'Грузовой автоподъёмник аэродромный', 'ГА', 'АПК-12 на ГАЗ-53А, АПК-9 на ЗИЛ-130Г', KK, ''),
    ('Аэродромные (новая)', 'Аэродромный пусковой агрегат (АПА)', 'ГА', 'АПА-50М на ЗИЛ-131', KK, ''),
    ('Аэродромные (новая)', 'Подогреватель аэродромный (моторный, салонный)', 'ГА', 'УМП-350 на ЗИЛ-131', KK, ''),
    ('Аэродромные (новая)', 'Противообледенительная машина (деайсер)', 'ГА', '—',
     'https://www.techinsider.ru/vehicles/11629-bortsy-so-ldom/', ''),
    ('Аэродромные или коммунальные', 'Ветровая (газоструйная) машина', 'ГА', 'обдув аэродромов и дорог горячим газом',
     'https://bigenc.ru/wiki/Ветровая_машина', ''),
    ('Аэродромные или коммунальные', 'Аэродромная уборочная машина (отвал, щётка, вакуум, обдув)', 'ГА',
     'УДК-АЭРО на КамАЗ-5399', 'https://www1.ru/news/2025/05/27/novaia-aerodromnaia-udk-aero-na-baze-kamaz-5399-uborka-moika-i-reagenty-vse-v-odnoi-masine.html',
     'Можно составным из имеющихся: отвал снежный + щётка + подметально-уборочная + ветровая машина.'),
    ('Коммунальные', 'Мойка мусорных контейнеров', 'ГА, ПР', 'Isuzu Grafter (Bin Boutique)',
     'https://isuzutruck.co.uk/news/articles/bin-boutique-cleans-up-with-new-isuzu-grafters/', ''),
    ('Коммунальные или дорожные', 'Мойка ограждений, знаков и тоннелей (щётка на стреле)', 'ГА, МКШ',
     'Schmidt RPS-H; ОПС-20.10', 'https://www.aebi-schmidt.com/_PDFs/released/web/PFS/665_PFS_rps-h_schmidt_navesnoe-oborudovanie_a4_m_ru_web.pdf', ''),
    ('Строительные и дорожные', 'Машина прикрытия с ударогасителем (TMA)', 'ГА', '—',
     'https://www.vicroads.vic.gov.au/-/media/files/documents/utilities/about-vr/ohs/guidelines-for-the-use-of-truck-mounted-attenuators0714web.ashx',
     'Источник зарубежный; в СНГ встречается на дорожных работах.'),
    ('Служебные и аварийные', 'Осветительная мачта передвижная (с генератором)', 'ПР', 'Mobilight',
     'https://ltcompany.com/en/series/mobilight', ''),
    ('Зимний спорт (новая)', 'Машина подготовки лыжных трасс и катков (снегоуплотнитель, ледозаливщик)', 'ТР',
     'BELARUS МСУ-622, МЛ-428, КЛ-418', 'https://belarus-tractor.com/en/production/tekhnika-dlya-zimnikh-vidov-sporta/',
     'МЛ-428 и КЛ-418 источники называют по-разному (подготовка трасс или каток) — сведены в один модуль.'),
    ('Культура (новая)', 'Автоклуб (передвижной клуб со сценой)', 'ГА', 'ГАЗон Next; ГАЗ-53А',
     'https://rg.ru/2019/08/07/reg-sibfo/v-omskoj-oblasti-poiavilsia-peredvizhnoj-kulturnyj-centr.html', ''),
    ('Культура (новая)', 'Библиобус (автобиблиотека)', 'АВ, ГА', '—',
     'https://riamo.ru/news/politika/vlasti-podmoskovya-namereny-udvoit-chislo-bibliobusov-do-20-shtuk-xl/', ''),
    ('Культура (новая)', 'Мобильная сцена', 'ПП, ПР, ГА', '—', 'https://patents.google.com/patent/US9200462', ''),
    ('Реклама (новая)', 'Рекламный автомобиль со светодиодным экраном', 'ГА, ПР', '—',
     'https://patents.google.com/patent/MD150Z/ru', ''),
    ('Быт и питание (новая)', 'Передвижная столовая (автокухня)', 'ГА', 'ТБМ на КамАЗ-43118',
     'https://gird.ru/cat/gruzopassazhirskie-avtomobili/transportno-bytovye-mashiny-tbm/transportno-bytovye-mashiny-kamaz/peredvizhnaja-stolovaja-tbm-na-shassi-kamaz-43118/pdf', ''),
    ('Быт и питание (новая)', 'Полевая кухня прицепная', 'ПР', 'КП-130, ПК-60', 'https://ru.ruwiki.ru/wiki/КП-130', ''),
    ('Быт и питание (новая)', 'Мобильная баня, комплекс санобработки', 'ГА, ПР', 'баня на ГАЗ; КПССО на КамАЗ (МЧС)',
     'https://csoor.organizations.mchs.gov.ru/export/pdf/News/4236801', ''),
    ('Быт и питание (новая)', 'Мобильный туалетный модуль', 'ПР, ГА', '—',
     'https://rentnational.com/portable-toilets/Towable-Unit', ''),
    ('Медицинские', 'Передвижной рентген-кабинет (флюорограф, маммограф)', 'ГА, ПП, АВ', 'на КамАЗ',
     'https://www.amic.ru/news/9-noveyshih-peredvizhnyh-flyuorografov-peredany-bolnicam-kraya-v-ramkah-nacproekta-i-zdorove-i-foto-52741',
     'Можно оставить видом «Передвижного медицинского кабинета», а не отдельным модулем.'),
    ('Медицинские', 'Передвижной барокомплекс', 'ГА', 'на двух КамАЗах (МЧС Камчатки)',
     'https://www1.ru/news/2025/04/09/mcs-kamcatki-polucilo-unikalnyi-peredviznoi-barokompleks-na-baze-dvux-kamazov.html', ''),
    ('Служебные и аварийные', 'Мобильная станция водоочистки', 'ГА, ПР', '—',
     'https://csoor.organizations.mchs.gov.ru/export/pdf/News/4240006', ''),
    ('Служебные и аварийные', 'Передвижная котельная (ПКУ)', 'ПР, ГА', 'ПКУ',
     'https://productcenter.ru/products/catalog-pieriedvizhnyie-kotielnyie-3066', ''),
    ('Служебные и аварийные', 'Автозак (перевозка лиц под стражей)', 'ГА', 'КамАЗ-4308, ГАЗель',
     'https://www.zr.ru/content/articles/902582-avtozaki-ot-kletki-do-biotual/', ''),
    ('Ветеринарные (новая)', 'Машина отлова безнадзорных животных', 'ГА', 'ветслужба Алматы',
     'https://bizmedia.kz/2024-01-11-obnovilsya-avtopark-veterinarnoj-sluzhby-almaty/', 'Источник косвенный.'),
    ('Ветеринарные (новая)', 'Дезинфекционная установка (ДУК)', 'ГА, ЛГ', 'ДУК-1 на ГАЗ-3309',
     'https://www.vidal.ru/veterinar/novosti/8903', ''),
    ('Строительные и дорожные', 'Автобетоновоз (самосвальный, не миксер)', 'ГА', '—', OKPD,
     'В ОКПД2 (29.10.59.113) «автобетоновоз» может означать и миксер; если так называют миксер — модуль не нужен.'),
    ('Цистерны', 'Золовоз', 'ГА, ПП', '—', OKPD, ''),
    ('Грузовые кузова', 'Опоровоз (опоры ЛЭП)', 'ГА', 'КамАЗ-43118 с КМУ ИМ-180',
     'https://gird.ru/cat/spectekhnika/bortovye-avtomobili/bortovye-otraslevye-specializirovannye-avtomobili/kamaz-43118-oporovoz-s-kmu-im-180/pdf', ''),
    ('Цистерны', 'Живорыбная цистерна', 'ГА, ПР', '—', 'https://srac.tamu.edu/fact-sheets/serve/74', ''),
    ('Цистерны', 'Криогенная цистерна (жидкий кислород, азот)', 'ГА, ПП', '—',
     'https://engjournal.bmstu.ru/articles/1684/1684.pdf', ''),
    ('Грузовые кузова', 'Гидроборт (подъёмный задний борт)', 'ГА, ПП, ПР', 'Dhollandia',
     'https://pdf.directindustry.com/pdf/dhollandia/tail-lifts-vans-dh-vz/22685-118372.html',
     'Чаще дополнение к фургону или бортовому, чем самостоятельный модуль.'),
    ('Грузовые кузова', 'Кузов-фургон КУНГ', 'ГА, ПР, ВЗ', 'ЗИЛ-131, Урал, КамАЗ', 'https://en.wikipedia.org/wiki/KUNG', ''),
    ('Сельскохозяйственные', 'Опрыскиватель (навесной, прицепной)', 'ТР, ТПР', 'ОП-2000',
     'https://productcenter.ru/products/86446/ops-2000-2',
     'Прицепной опрыскиватель уже есть видом «Прицепной машины»; навесной на тракторе — новый.'),
    ('Сельскохозяйственные', 'Разбрасыватель минеральных удобрений', 'ТР, ТПР', 'РМГ-4',
     'https://rep.bntu.by/handle/data/81280?show=full', 'Прицепной уже есть видом «Прицепной машины»; навесной — новый.'),
    ('Сельскохозяйственные', 'Передвижной пчелопавильон (кочевая пасека)', 'ПР, ПП', '—',
     'https://booksite.ru/localtxt/pch/elo/ovo/dst/pchelovodstvo/11.htm', ''),
    ('Учебные (новая)', 'Учебный автомобиль (дублирующие педали)', 'ЛГ, ГА, АВ', 'Lada Granta учебная',
     'https://pddmaster.ru/documents/gost-r-55887-2013-avtomobilnye-transportnye-sredstva-uchebnye-avtomobili-tehnicheskie-trebovaniya-i-metody-ispytanii/4-obschie-tehnicheskie-trebovaniya-k-dubliruyuschim-organam-upravleniya', ''),
    ('Пассажирские (новая)', 'Школьный автобус (перевозка детей)', 'АВ, ГА', 'по ГОСТ 33552-2015',
     'https://www.gostinfo.ru/News/Details/485', ''),
    ('Пожарные', 'Пожарный дымоудаления (АД)', 'АВ, ГА', 'АД 90/22 на ПАЗ-3205',
     'https://mchs.gov.ru/ministerstvo/o-ministerstve/tehnika/pozharnaya-tehnika', ''),
    ('Грузовые кузова', 'Прицеп для перевозки лодок', 'ПР', '—',
     'https://productcenter.ru/products/catalog-lodochnyie-avtopritsiepy-3740', ''),
    ('Ходовая (новая)', 'Комбинированный (рельсовый) ход', 'ГА, ЭКС, МКШ, ТР', 'Liebherr A 922 Rail; Multicar Zweiweg',
     'https://www.liebherr.com/ru-by/p/643370-4964278',
     'Это свойство ходовой, а не надстройка: может быть уместнее признаком в поле «Ходовая».'),
]

PAIRS = [
    ('Автобетоносмеситель (миксер)', 'ТР', 'самозагружаемый миксер от ВОМ трактора (MAMMUT)',
     'https://www.mammut.at/en/pan-mixers/pan-mixer-for-tractors-mix-concrete-yourself-at-low-cost/', ''),
    ('Мусоровоз с задней загрузкой', 'ТПР', 'тракторный прицеп с прессом', 'https://dl.lib.uom.lk/items/ed77a1de-1988-44c1-8d50-e1776de0df27',
     'Источник зарубежный (Шри-Ланка).'),
    ('Каналопромывочная машина', 'ТПР', 'прицепная цистерна к трактору с насосами от ВОМ',
     'https://www.agriexpo.online/agricultural-manufacturer/vacuum-truck-151.html', ''),
    ('Пескоразбрасыватель', 'ЛГ', 'в кузов или на задний борт пикапа (SaltDogg)',
     'https://www.buyersproducts.com/product/SaltDogg-TGS07-11-Cubic-Foot-Tailgate-Spreader-1813', ''),
    ('Шнекороторный снегоочиститель', 'МП', 'снегометатель на мини-погрузчик',
     'https://www.totallandscapecare.com/business/article/15037696/paladins-blower-attachment-turns-skid-steers-into-mini-snowplows', ''),
    ('Шнекороторный снегоочиститель', 'ЭП', 'снегометатель на экскаватор-погрузчик', 'https://skidpro.com/2012/11/backhoe-snow-blower', ''),
    ('Щётка коммунальная', 'ЭП', 'щётка с бункером', 'https://www.directindustry.com/prod/cm-srl/product-88995-2732965.html', ''),
    ('Щётка коммунальная', 'ТП', 'JCB Sweeper Collector',
     'https://www.norlift.adpearance.com/catalog/attachments/recycling-and-other/sweeper-collector', ''),
    ('Ковш основной', 'ТП', 'перевалочный ковш', 'https://gap-group.co.uk/catalogue/telehandler-attachments-rehandling-bucket', ''),
    ('Косилка навесная', 'ЭКС', 'мульчирующая косилка на стреле (откосы, кюветы)',
     'https://landscapearchitect.com/landscape-articles/next-evolution-of-excavator-flail-mowers', ''),
    ('Косилка навесная', 'МЭКС', 'Bobcat FMR 40"',
     'https://www.constructionequipment.com/equipment-attachments/boom-mounted-tools-accessories/product/10751073/bobcat-company-bobcat-40-inch-fmr-flail-mower-for-bobcat-e-series-excavators', ''),
    ('Автодом', 'ЛГ', 'жилой модуль в кузов пикапа', 'https://campaddict.com/truck-bed-camper', ''),
    ('Скорая помощь класса B (экстренная)', 'ВЗ', 'гусеничная «неотложка»', 'https://www.zr.ru/content/articles/14690-brat_miloserdija/', ''),
    ('Кран-манипулятор (КМУ)', 'ВЗ', 'снегоболотоход ТМ-140 с КМУ до 5 т', 'https://auto.ru/mag/article/russnowmobiles/',
     'Источник — пересказ, страница целиком не прочитана.'),
    ('Автовышка', 'МКШ', 'Multicar с площадкой', 'https://www.truckscout24.com/tsp/ts-221-67-887', ''),
    ('Подметально-уборочная машина', 'АВ', 'МКУ-4 на ПАЗ-672 (аэродромная)', KK, ''),
    ('Кормораздатчик на шасси', 'ТПР', 'прицеп-раздатчик ПРКТ-10', 'https://www.olx.kz/d/obyavlenie/pritsep-razdatchik-kormov-prkt-10-IDoKk9G.html',
     'Прицепной раздатчик уже есть видом «Прицепной машины» (кормораздатчик прицепной) — проверить, не дубль.'),
    ('Пожарный вспомогательный (аварийно-спасательный, штабной, рукавный)', 'АВ', 'дымоудаление АД 90/22 на ПАЗ-3205',
     'https://mchs.gov.ru/ministerstvo/o-ministerstve/tehnika/pozharnaya-tehnika', 'Носитель АВ уже есть — только уточнение.'),
]

COMPOSITE = [
    ('Самосвал с КМУ', 'ГА', 'Самосвальный кузов + Кран-манипулятор (КМУ)', 'https://productcenter.ru/products/catalog-kranovyie-manipuliatory-4003', ''),
    ('Седельное ТС с КМУ', 'СТ', 'Кран-манипулятор (КМУ)', 'там же', 'Трал — отдельной карточкой полуприцепа. ' + TWO),
    ('Ломовоз', 'ГА', 'Самосвальный кузов + Кран-манипулятор (КМУ) + Грейфер', 'https://www.drom.ru/info/misc/98836.html',
     'Грейфер в справочнике — навесное для экскаватора и погрузчика; на КМУ он — рабочий орган крана.'),
    ('Опоровоз с КМУ', 'ГА', 'Бортовая платформа + Кран-манипулятор (КМУ)', 'см. «Опоровоз» на листе модулей',
     'Роспуск — отдельной карточкой прицепа. ' + TWO),
    ('Автоцистерна с лестницей (АЦЛ)', 'ГА', 'Пожарная автоцистерна (АЦ) + Пожарная автолестница (АЛ)',
     'https://mchs.gov.ru/ministerstvo/o-ministerstve/tehnika/pozharnaya-tehnika/avtocisterna-pozharnaya-s-lestnicey-acl-3-40-4-24-43118', ''),
    ('ПАРМ с КМУ', 'ГА', 'Передвижная автомастерская + Кран-манипулятор (КМУ)',
     'https://gird.ru/cat/spectekhnika/peredvizhnye-avtoremontnye-masterskie/parm-s-kmu/parm-kamaz-43118-s-kmu-rk-17001/pdf', ''),
    ('Миксер с бетононасосом', 'ГА', 'Автобетоносмеситель (миксер) + Автобетононасос',
     'https://www.directindustry.com/prod/putzmeister/product-21069-1886674.html', ''),
    ('Эвакуатор: платформа и подхват', 'ГА', 'Эвакуатор с платформой + Эвакуатор с частичной погрузкой',
     'https://helpix.ru/news/202311/131700-kakie_byvajut_evakuatory/index.html', ''),
    ('Аэродромная уборочно-продувочная', 'ГА', 'Отвал снежный + Щётка коммунальная + Подметально-уборочная машина + Ветровая машина (новая)',
     'https://www1.ru/news/2025/05/27/novaia-aerodromnaia-udk-aero-na-baze-kamaz-5399-uborka-moika-i-reagenty-vse-v-odnoi-masine.html',
     'Зависит от модуля «Ветровая машина» (лист модулей).'),
    ('Фургон или бортовой с гидробортом', 'ГА, ПП, ПР', 'Промтоварный/Изотермический фургон/Рефрижератор + Гидроборт (новый)',
     'https://pdf.directindustry.com/pdf/dhollandia/tail-lifts-vans-dh-vz/22685-118372.html', 'Зависит от модуля «Гидроборт».'),
    ('Передвижной ФАП с флюорографом', 'ГА, АВ', 'Передвижной медицинский кабинет + Передвижной рентген-кабинет (новый)',
     'https://riamo.ru/news/zdravoohranenie/esche-5-peredvizhnyh-fapov-s-flyuorografami-postupili-v-bolnitsy-podmoskovya/', ''),
    ('Автоклуб-трансформер', 'ГА', 'Автоклуб (новый) + Мобильная сцена (новая) + Электростанция (дизель-генератор)',
     'https://rg.ru/2019/08/07/reg-sibfo/v-omskoj-oblasti-poiavilsia-peredvizhnoj-kulturnyj-centr.html', ''),
    ('Снегоболотоход с КМУ', 'ВЗ', 'Бортовая платформа + Кран-манипулятор (КМУ)', 'https://auto.ru/mag/article/russnowmobiles/',
     'Источник — пересказ.'),
    ('Сцепка: самосвал + самосвальный прицеп', 'ГА + ПР', 'Самосвальный кузов на ГА и на ПР', 'https://lalafo.kg/alekseevka/ads/pricep-kamaz-samosval-id-76572463',
     TWO + ' Оба шаблона уже есть.'),
]

WORDS = [
    ('буханка, батон, горбушка, булочница', 'УАЗ-452 (фургон, вахтовка, база модулей)', NICK, 'Название модели — уместно в поиске шаблонов на грузовом ТС с фургоном.'),
    ('таблетка, санитарка', 'Скорая помощь класса A (УАЗ-452)', NICK, '«Санитарка» уже есть.'),
    ('бобик, козёл, коробок', 'УАЗ-469 (легковой), оперативный (патрульный)', NICK, ''),
    ('патрик', 'УАЗ Патриот (легковой)', NICK, 'Название модели, не вида — в шаблоны не обязательно.'),
    ('рафик', 'РАФ-977 (микроавтобус; скорая и катафалк на его базе)', NICK, ''),
    ('гармошка, колбаса', 'сочленённый автобус', NICK, ''),
    ('кальмар', 'К-700 «Кировец» (трактор)', NICK, ''),
    ('зилок, крокодил', 'ЗИЛ-130, ЗИЛ-164/166 (грузовое ТС)', NICK, ''),
    ('шишига', 'ГАЗ-66 (база кунга, вахтовки, мастерской)', 'https://www.zr.ru/content/articles/938024-pochemu-gaz-66-prozvali-shishiga/', ''),
    ('табуретка', 'ГАЗ-53', 'https://sport24.ru/auto/article-teper-ponyatno-pochemu-legendarnyy-gruzovik-sssr-gaz-53-nazyvali-tak-stranno-chto-yeshche-za-taburetka', ''),
    ('колхозник', 'сельхозсамосвал КамАЗ-55102 / ЗИЛ-ММЗ-554', 'https://auto.drom.ru/spec/kamaz/55102/', ''),
    ('сцепка', 'самосвал с самосвальным прицепом', 'https://lalafo.kg/alekseevka/ads/pricep-kamaz-samosval-id-76572463', 'Две карточки.'),
    ('хова, хово, мини-хова', 'самосвал или грузовик Howo; малый китайский самосвал', 'https://lalafo.kg/ky/dmitrievka/transport/q-тяга-камаз', '«Хово» уже есть у самосвала.'),
    ('портер', 'малотоннажный бортовой', 'https://lalafo.kg/kyrgyzstan/transport/q-авто-портер', 'Уже есть.'),
    ('манипулятор', 'бортовой с КМУ', 'https://lalafo.kg/user/2440251', 'Уже есть.'),
    ('шаланда', 'длинный бортовой полуприцеп', 'https://dic.academic.ru/dic.nsf/dic_synonims/197560/', ''),
    ('буратино', 'лесовоз', 'https://trans.info/ru/kak-voditeli-dalnoboyshhiki-obshhayutsya-mezhdu-soboy-dva-slova-o-ten-kodah-i-voditelskom-zhargone-197532', ''),
    ('кунг', 'кузов-фургон КУНГ (новый модуль)', 'https://en.wikipedia.org/wiki/KUNG', ''),
    ('вахтовка', 'Вахтовый автобус', 'https://www.zr.ru/content/articles/854581-avtobus-vaxtovka-rodstvennik-kunga/', 'Уже есть.'),
    ('деайсер, диайсер', 'противообледенительная машина (новый модуль)', 'https://www.techinsider.ru/vehicles/11629-bortsy-so-ldom/', ''),
    ('поливалка, мусорка, пожарка, аварийка, кдмка', 'поливомоечная, мусоровоз, пожарная, аварийная, КДМ', 'общеупотребительные',
     'Источника нет; «поливалка», «мусорка», «пожарка» уже есть.'),
]

DOUBT = [
    ('«Пухто» / «пухтовоз» (бункер, бункеровоз)', 'Надёжного словаря не нашлось; слово, вероятно, петербургское.'),
    ('«Лодочка» (бункер-накопитель бункеровоза)', 'Источника не нашлось.'),
    ('Кыргызские названия: «тез жардам» (скорая), «өрт өчүрүүчү» (пожарная), «таштанды ташуучу» (мусоровоз), «суу ташуучу» (водовоз)',
     'Подтверждения не нашлось; если описи бывают на кыргызском — уточнить у пользователя.'),
    ('Отвал снежный на легковом (УАЗ, Нива, пикап), на экскаваторе-погрузчике, на телескопическом погрузчике',
     'Встречается широко, прямого подтверждения не нашлось.'),
    ('Вахтовый автобус на полуприцепе (пассажирский полуприцеп)', 'Только определение полуприцепа, без модели.'),
    ('Автотопливомаслозаправщик (АТМЗ) — топливозаправщик + маслозаправщик', 'Термин известный, поиск дал только АТЗ.'),
    ('Мусоровоз с КМУ для заглублённых контейнеров; мультилифт с КМУ', 'Источники общие, без модели.'),
    ('Передвижная хлебопекарня', 'Есть военные полевые пекарни; страница ГАЗа не прочитана.'),
    ('Передвижная трансформаторная подстанция на прицепе; подвижный узел связи МЧС на КамАЗе',
     'Реальны, но источники — о стационарных КТП или мимоходом; узел связи близок к «Штабному автомобилю».'),
    ('Измельчитель веток (щепорез) на прицепе', 'Нашлись только садовые модели.'),
    ('Ассенизаторская бочка, навешенная на трактор', 'Подтверждения нет; прицепная (тракторный прицеп) уже есть.'),
    ('Мусоровоз — пресс на навеске трактора', 'Только патенты и проекты; тракторный прицеп предложен на листе сочетаний.'),
]


def sheet(wb, title, head, rows, widths):
    ws = wb.create_sheet(title)
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True)
        c.fill = PatternFill('solid', fgColor='DCEBF8')
        c.alignment = Alignment(wrap_text=True, vertical='center')
    for r in rows:
        ws.append(list(r))
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical='top')
    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions


def build():
    wb = Workbook()
    wb.remove(wb.active)
    sheet(wb, 'Новые модули', ['Группа', 'Модуль', 'Носители', 'Пример', 'Источник', 'Замечание', 'Решение'],
          [r + ('',) for r in MODULES], [24, 40, 14, 30, 40, 40, 20])
    sheet(wb, 'Новые сочетания', ['Модуль (как в справочнике)', 'Добавить носитель', 'Пример', 'Источник', 'Замечание', 'Решение'],
          [r + ('',) for r in PAIRS], [40, 14, 36, 40, 40, 20])
    sheet(wb, 'Составные', ['Название', 'Носители', 'Модули', 'Источник', 'Замечание', 'Решение'],
          [r + ('',) for r in COMPOSITE], [32, 12, 50, 40, 44, 20])
    sheet(wb, 'Обиходные названия', ['Слова (для поиска)', 'К чему относится', 'Источник', 'Замечание', 'Решение'],
          [r + ('',) for r in WORDS], [30, 40, 40, 40, 20])
    sheet(wb, 'Сомнительное', ['Что', 'Почему сомнительно', 'Решение'], [r + ('',) for r in DOUBT], [60, 60, 20])
    sheet(wb, 'Коды носителей', ['Код', 'Носитель'], [(k, v[3]) for k, v in CARRIER.items()], [10, 30])
    wb.save(OUT)
    return len(MODULES), len(PAIRS), len(COMPOSITE), len(WORDS), len(DOUBT)


if __name__ == '__main__':
    print('docs/predlozheniya-shablonov-ts.xlsx собран: модулей %d, сочетаний %d, составных %d, слов %d, сомнительного %d' % build())
