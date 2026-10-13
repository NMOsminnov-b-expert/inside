# -*- coding: utf-8 -*-
"""Чтение перечня папок с фото и нормализация имён."""
import os
import re
import collections as C

from openpyxl import load_workbook

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__)))))
SRC_DIR = os.path.join(ROOT, 'Анализ по фото регионалов')
SRC = os.path.join(SRC_DIR, 'папки_с_фото.xlsx')
BS = chr(92)
# общий корень всех путей в перечне; ниже него — область, район и дальше папки регионала
PREF = BS * 2 + '?' + BS + 'Y:' + BS + '1. ЦОД' + BS + 'Информация от админа' + BS + '9. Осмотрщики (Регионалы)' + BS


def load():
    """Пути папок из перечня: список уровней, начиная с области."""
    ws = load_workbook(SRC, read_only=True).worksheets[0]
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True)):
        if i == 0 or not row or not row[1]:
            continue
        path = row[1]
        assert path.startswith(PREF), 'путь вне папки регионалов: %s' % path
        out.append(path[len(PREF):].split(BS))
    return out


def norm(s):
    s = s.strip().lower().replace('ё', 'е')
    s = re.sub(r'\s*\+\s*$', '', s)
    return re.sub(r'\s+', ' ', s).strip()
