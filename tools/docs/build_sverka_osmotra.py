# -*- coding: utf-8 -*-
"""Сверка полей «Осмотров» с карточками объектов оценки.

Запрос пользователя 14.10.2026: «Сверь пожалуйста поля Осмотров с
карточками ОЦ». Форма осмотра у каждого вида ОИ своя (app/pages/inspector/
form.js, sections), а значения «по документам» рядом с полем осмотрщик
видит из карточки. Сверка показывает, где осмотр и карточка расходятся:
подпись поля, перечень значений, то, откуда берётся «по документам».

Обе стороны снимаются с живого макета, поэтому сверка пересобирается
одной командой:

    python tools/docs/build_sverka_osmotra.py   # docs/sverka-osmotra-i-kartochek.xlsx

* осмотр — разделы формы прямо из кода экрана (sections) по типу ОЦ и ОИ
  сводных записей «-all»: подпись, перечень, обязательность и значение
  «по документам» для этого ОИ;
* карточка — поля и перечни с экрана карточки того же ОИ (разбор экрана —
  из сборщика справочника полей, build_reference.py), с текущим значением.

Поле осмотра ищется в карточке по подписи: точно, затем по основам слов.
Найденное «похоже» — проверить глазами: подписи разные.
"""
import io
import os
import re
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
sys.path.insert(0, os.path.join(ROOT, 'tools', 'checks'))
sys.path.insert(0, os.path.join(ROOT, 'tools', 'docs'))
OUT = os.path.join(ROOT, 'docs', 'sverka-osmotra-i-kartochek.xlsx')

MODULES = [
    ('apartment', 'ap', 'Жилое здание (квартира)'),
    ('residential-house', 'rh', 'Жилое здание (дом)'),
    ('civil', 'cv', 'Нежилое здание'),
    ('land-plot', 'lp', 'Земельный участок'),
]

# Разбор карточки — тот же, что у справочника полей, плюс текущее значение
# поля: по нему сверяется «по документам» в осмотре.
from build_reference import PROBE  # noqa: E402
PROBE_V = PROBE.replace(
    "        значения: el.tagName === 'SELECT'",
    "        значение: el.tagName === 'SELECT' ? ((el.selectedOptions[0] || {}).textContent || '').trim() : (el.type === 'checkbox' ? (el.checked ? 'да' : 'нет') : (el.value || '').trim()),\n"
    "        значения: el.tagName === 'SELECT'", 1).replace(
    "        значения: [...el.querySelectorAll('.ms-opt')]",
    "        значение: txt(el.querySelector('.ms-control, .ms-val, .ms-text') || el).replace(/\\s*▾\\s*$/, ''),\n"
    "        значения: [...el.querySelectorAll('.ms-opt')]", 1)
PROBE_V = PROBE_V.replace(
    "        ключ: attr,\n        вид: kind,",
    "        ключ: attr,\n        атрибут: attr.startsWith('#') ? attr : attr + '=' + (el.getAttribute(attr) || ''),\n        вид: kind,", 1)
assert PROBE_V.count('значение:') == 2 and 'атрибут:' in PROBE_V, 'разбор карточки поменялся — поправить вставки'

# Поле записи ОИ → атрибут поля карточки, которое его показывает.
def card_attr(path):
    head, _, tail = (path or '').partition('.')
    return {'areas': 'data-area=', 'heights': 'data-height=', 'struct': 'data-struct-field='}.get(head, '') + tail if tail else ''

# Форма осмотра — из кода экрана.
INSPECT = r"""async (pairs) => {
  const F = await import('./app/pages/inspector/form.js');
  const R = await import('./app/kernel/registry.js');
  // Запись — как её читает экран осмотрщика (loadTask: вид литеры по правилу
  // карточки), а не напрямую.
  const T = await import('./app/pages/inspector/tasks.js');
  const str = (v) => (v == null ? '' : String(v));
  const out = [];
  for (const [typeId, ocId] of pairs) {
    const t = R.getType(typeId);
    const rec = t ? T.loadTask(typeId, ocId) : null;
    if (!rec) continue;
    for (const oi of rec.oi || []) {
      const secs = F.sections(typeId, oi);
      out.push({
        typeId, ocId, oiId: oi.id, card: oi.card || 'building', kind: oi.kind || '', name: oi.name || oi.letter || '',
        secs: secs.map((s) => ({
          title: s.title,
          fields: s.fields.map((f) => ({
            key: f.key, label: f.label, opts: f.opts || null, multi: !!f.multi, unit: f.unit || '',
            num: !!f.num, req: !!f.req, long: !!f.long, cond: !!f.visible,
            // Поле с условием показа (вариант по виду строения) — видно ли
            // оно этому ОИ при пустом осмотре: вид берётся из записи ОИ.
            shown: typeof f.visible === 'function' ? !!f.visible({}) : true,
            doc: typeof f.doc === 'function' ? str(f.doc(oi)) : (f.tpFrom ? str(F.tpValue(oi, f.tpFrom)) : null),
            docPath: typeof f.doc === 'function' ? str(f.doc.path) : (f.tpFrom ? ({ tp: 'areas.tp', pud: 'areas.pud', build: 'areas.build', heightExt: 'heights.ext', heightInt: 'heights.int' })[f.tpFrom] || f.tpFrom : ''),
          })),
        })),
      });
    }
  }
  return out;
}"""

# Переключатели сегментами карточки: подпись — aria-label группы, перечень —
# надписи кнопок, значение — нажатая.
SEGMENTS = r"""() => [...document.querySelectorAll('.oi-stack [role=radiogroup][aria-label]')].map((g) => {
  const card = g.closest('.card');
  const head = card && card.querySelector('h3');
  const btns = [...g.querySelectorAll('button, [role=radio]')];
  const on = btns.find((x) => x.getAttribute('aria-checked') === 'true' || x.getAttribute('aria-pressed') === 'true' || x.classList.contains('on'));
  return { блок: head ? head.textContent.replace(/\s+/g, ' ').trim() : '', подпись: g.getAttribute('aria-label'),
           значения: btns.map((x) => x.textContent.replace(/\s+/g, ' ').trim()).filter(Boolean),
           значение: on ? on.textContent.replace(/\s+/g, ' ').trim() : '' };
})"""

STOP = {'наличие', 'по', 'и', 'или', 'в', 'на', 'с', 'из', 'для', 'более', 'свой', 'вариант', 'м', 'м²', 'сом', 'шт', '1'}


def norm(s):
    # Звёздочка обязательного поля — не часть подписи.
    s = (s or '').lower().replace('ё', 'е').replace('*', ' ')
    s = re.sub(r'\(.*?\)', ' ', s)
    s = re.sub(r',\s*(м²|м|%|сом.*|шт\.?)\s*$', ' ', s)
    return re.sub(r'[^0-9a-zа-я]+', ' ', s).strip()


def stems(s):
    return {w[:5] for w in norm(s).split() if w not in STOP and len(w) > 1}


# Одна и та же величина, подписанная в осмотре и карточке по-разному, —
# пара задана явно («другая подпись»). None — у поля пары в карточке нет,
# похожее по словам сопоставлять нельзя (наличие отопления — не вид
# отопления).
ALIASES = {
    # Подписи осмотра сведены с карточкой 15.10.2026 (эталон — карточка):
    # пар с другой подписью не осталось.
}


def match(label, fields, section=''):
    """Поле карточки к полю осмотра: (поле, 'точно'|'другая подпись'|'похоже'|'')."""
    # Износ в карточке — колонка «Износ» таблицы конструктива: строка
    # (Кровля, Отделка, Утепление) у неё общая с материалом.
    # С 15.10.2026 износ в осмотре — поле «<элемент> — износ» в разделе
    # «Конструктив и износ», как строка таблицы карточки.
    if section == 'Износ' or label.endswith(' — износ'):
        for f in fields:
            if f['подпись'] == 'Износ':
                return f, 'колонка «Износ»'
    if label in ALIASES:
        for want in ALIASES[label] or []:
            for f in fields:
                if f['подпись'] == want:
                    return f, 'другая подпись'
        return None, ''
    n = norm(label)
    # Одна подпись в разных блоках («Полы» — материал в «Конструктив и
    # износ» и признак в «Тип и класс капитальности»): сначала поле блока с
    # тем же названием, что у раздела осмотра.
    for f in fields:
        if norm(f['подпись']) == n and norm(f.get('блок', '')) == norm(section):
            return f, 'точно'
    for f in fields:
        if norm(f['подпись']) == n:
            return f, 'точно'
    a = stems(label)
    best, score = None, 0.0
    for f in fields:
        b = stems(f['подпись'])
        if not a or not b:
            continue
        j = len(a & b) / len(a | b)
        if j > score:
            best, score = f, j
    if best and score >= 0.5:
        return best, 'похоже'
    return None, ''


def vals(lst):
    return [v.strip() for v in (lst or []) if v and v.strip() and v.strip() not in ('—', '-')
            and not re.match(r'^(не выбрано|не заполнено|выберите.*)$', v.strip(), re.I)]


EMPTY = re.compile(r'^(не выбрано|не заполнено|—|-)?$', re.I)


def same_value(a, b):
    a = '' if EMPTY.match((a or '').strip()) else a
    b = '' if EMPTY.match((b or '').strip()) else b
    if norm(a) == norm(b):
        return True
    try:
        num = lambda v: float(re.sub(r'[\s \xa0]', '', str(v)).replace(',', '.'))  # noqa: E731
        return abs(num(a) - num(b)) < 1e-6
    except ValueError:
        return False


def collect():
    import harness
    from playwright.sync_api import sync_playwright
    srv = harness.serve()
    base = 'http://127.0.0.1:%d/app.html' % harness.PORT
    data = []
    try:
        with sync_playwright() as p:
            b = p.chromium.launch()
            pg = b.new_page(viewport={'width': 1600, 'height': 1200})
            pg.goto(base + '#/')
            pg.wait_for_selector('.reg-thead', timeout=30000)
            pairs = [[mod, 'oc-%s-all' % tag] for mod, tag, _ in MODULES]
            insp = pg.evaluate(INSPECT, pairs)
            for it in insp:
                title = next(t for m, _, t in MODULES if m == it['typeId'])
                url = '%s#/oc/%s/%s/oi/%s' % (base, it['typeId'], it['ocId'], it['oiId'])
                pg.goto(url)
                try:
                    pg.wait_for_selector('.oi-stack .card', timeout=8000)
                except Exception:
                    blocks = []
                else:
                    pg.wait_for_timeout(400)
                    pg.evaluate("() => document.querySelectorAll('.card.closed').forEach(c => c.classList.remove('closed'))")
                    blocks = pg.evaluate(PROBE_V)
                    # Переключатель сегментами («Тип объекта имущества» —
                    # .lk-seg, role=radiogroup) разбор справочника не видит:
                    # он знает списки и поля ввода. Добавляем такие поля к
                    # блоку, где они стоят.
                    seg = pg.evaluate(SEGMENTS)
                    for g in seg:
                        blk = next((b for b in blocks if b['заголовок'] == g['блок']), None)
                        if blk is None:
                            blk = {'заголовок': g['блок'], 'номер': '', 'поля': []}
                            blocks.append(blk)
                        if not any(f['подпись'] == g['подпись'] for f in blk['поля']):
                            blk['поля'].append({'подпись': g['подпись'], 'значения': g['значения'], 'значение': g['значение'],
                                                'вид': 'переключатель', 'ключ': '', 'атрибут': ''})
                it['typeTitle'] = title
                it['cardBlocks'] = blocks
                data.append(it)
                print('  %s · %s (%s): полей осмотра %d, блоков карточки %d' % (
                    title, it['oiId'], it['card'], sum(len(s['fields']) for s in it['secs']), len(blocks)))
            b.close()
    finally:
        srv.kill()
    return data


KIND_RU = {'building': 'Строение', 'apartment': 'Квартира', 'land': 'Земельный участок', 'mech': 'Механизмы',
           'movable': 'Механизмы (прежнее)', 'vehicle': 'ТС'}


def rows_of(data):
    sverka, absent = [], []
    for it in data:
        cfields = []
        for b in it['cardBlocks']:
            for f in b['поля']:
                if f['подпись']:
                    cfields.append(dict(f, блок=b['заголовок']))
        kind = KIND_RU.get(it['card'], it['card'])
        oi = '%s — %s' % (kind, it['name']) if it['name'] else kind
        used = set()
        # Механизмы: форма — по единицам записи (oi.mechanisms); карточка
        # показывает единицы и у прежнего вида «движимое» (переводит на лету).
        units_card = [b['заголовок'] for b in it['cardBlocks'] if b['заголовок'] not in ('Состав',)]
        units_form = [s for s in it['secs'] if s['title'] != 'Комментарии']
        if it['card'] in ('mech', 'movable') and units_card and not units_form:
            sverka.append([it['typeTitle'], oi, '—', '(единицы механизмов)', '', 'единицы: ' + '; '.join(units_card),
                           'нет', '', '', '', '', '', '',
                           'единиц в форме осмотра нет — в записи ОИ нет списка единиц (вид «%s»), карточка показывает %d' % (it['card'], len(units_card))])
        if not it['secs']:
            sverka.append([it['typeTitle'], oi, '—', '(формы осмотра нет)', '', '', '', '', '', '', '', '', '',
                           'Форма осмотра для этого вида не задана'])
        for s in it['secs']:
            for f in s['fields']:
                if f['long'] or f['label'].endswith('свой вариант'):
                    continue
                # Вариант поля для другого вида строения (класс капитальности:
                # у гражданского, производственно-складского и прочего свои
                # перечни) этому ОИ не показывается — сравнивать его с
                # карточкой значит ловить чужой перечень (сверка 15.10.2026).
                if not f.get('shown', True):
                    continue
                c, how = match(f['label'], cfields, s['title'])
                if c:
                    used.add(id(c))
                io_ = vals(f['opts'])
                ic = vals(c['значения']) if c else []
                only_i = [v for v in io_ if not any(norm(v) == norm(x) for x in ic)] if c and ic else []
                only_c = [v for v in ic if not any(norm(v) == norm(x) for x in io_)] if c and io_ else []
                doc = f['doc']
                # Поле карточки, из которого на самом деле берётся «по документам».
                want = card_attr(f['docPath'])
                src = next((x for x in cfields if want and x.get('атрибут') == want), None)
                src_label = src['подпись'] if src else (f['docPath'] or '')
                cval = (c.get('значение') or '') if c else ''
                # У мультивыбора к тексту приклеен счётчик выбранного.
                if c and c['вид'] == 'мультивыбор':
                    cval = re.sub(r'\d+$', '', cval).strip()
                if doc is None:
                    dv = ''
                elif not c:
                    dv = '—'
                else:
                    dv = 'да' if same_value(doc, cval) else 'НЕТ'
                problems = []
                if not c:
                    problems.append('нет в карточке')
                else:
                    if how in ('похоже', 'другая подпись'):
                        problems.append('подпись другая')
                    if only_i or only_c:
                        problems.append('перечни расходятся')
                    if io_ and not ic and c['вид'] not in ('список', 'мультивыбор'):
                        problems.append('в карточке не список (%s)' % c['вид'])
                    if not io_ and ic:
                        problems.append('в карточке список, в осмотре нет')
                    if dv == 'НЕТ':
                        problems.append('«по документам» не равно карточке')
                    # Замер по факту сверяется с документами — это задумано: поле
                    # «по документам» и должно быть другим, если оно по
                    # правоустанавливающим документам.
                    if src and src is not c and doc is not None and 'правоуст' not in src['подпись']:
                        problems.append('«по документам» — из поля «%s»' % src['подпись'])
                sverka.append([
                    it['typeTitle'], oi, s['title'], f['label'] + (', ' + f['unit'] if f['unit'] else ''),
                    'да' if f['req'] else '',
                    ('%s · %s' % (c['блок'], c['подпись'])) if c else '',
                    how or 'нет',
                    '; '.join(only_i), '; '.join(only_c),
                    src_label if doc is not None else '', '' if doc is None else doc, cval, dv,
                    ', '.join(problems) or 'совпадает',
                ])
        for c in cfields:
            if id(c) in used or c['вид'] == 'только чтение':
                continue
            absent.append([it['typeTitle'], oi, c['блок'], c['подпись'], c['вид'],
                           '; '.join(vals(c['значения'])[:12]) + (' …' if len(vals(c['значения'])) > 12 else '')])
    return sverka, absent


def safe_out(path):
    """Правило пользователя 14.10.2026: файл, который человек правил после
    сборки или держит открытым, не перезаписывается."""
    lock = os.path.join(os.path.dirname(path), '~$' + os.path.basename(path))
    if os.path.exists(lock):
        raise SystemExit('файл открыт (%s) — закройте и запустите снова' % lock)
    if os.path.exists(path):
        from openpyxl import load_workbook
        props = load_workbook(path, read_only=True).properties
        if props.lastModifiedBy and props.lastModifiedBy != 'Сборщик макета':
            raise SystemExit('файл правил %s — не перезаписываю; переименуйте его или удалите сами' % props.lastModifiedBy)
    return path


def build(sverka, absent, path):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    wb = Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill('solid', fgColor='DCE9F7')
    bad_fill = PatternFill('solid', fgColor='FDECEC')
    warn_fill = PatternFill('solid', fgColor='FFF4DD')
    wrap = Alignment(vertical='top', wrap_text=True)

    def sheet(title, cols, rows, widths, mark=None):
        ws = wb.create_sheet(title)
        ws.append(cols)
        for i in range(1, len(cols) + 1):
            c = ws.cell(row=1, column=i)
            c.fill = head_fill
            c.font = Font(bold=True, color='1F3A5F')
            c.alignment = Alignment(vertical='center', wrap_text=True)
        for r in rows:
            ws.append(r)
            if mark:
                fill = mark(r)
                if fill:
                    for i in range(1, len(cols) + 1):
                        ws.cell(row=ws.max_row, column=i).fill = fill
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = wrap
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        return ws

    def mark(r):
        t = r[-1]
        if t == 'совпадает':
            return None
        return bad_fill if ('нет в карточке' in t or 'не равно' in t or 'не задана' in t or 'из поля' in t or 'единиц в форме' in t) else warn_fill

    # Сводка — по виду ОИ: сколько полей совпало и что не так.
    summary = {}
    for r in sverka:
        k = (r[0], r[1])
        s = summary.setdefault(k, {'всего': 0, 'ok': 0, 'нет': 0, 'подпись': 0, 'перечни': 0, 'док': 0})
        s['всего'] += 1
        t = r[-1]
        s['ok'] += t == 'совпадает'
        s['нет'] += 'нет в карточке' in t
        s['подпись'] += 'подпись другая' in t
        s['перечни'] += 'перечни расходятся' in t
        s['док'] += 'не равно' in t
    sheet('Сводка', ['Тип ОЦ', 'ОИ', 'Полей осмотра', 'Совпадают', 'Нет в карточке', 'Подпись другая',
                     'Перечни расходятся', '«По документам» не равно карточке'],
          [[k[0], k[1], v['всего'], v['ok'], v['нет'], v['подпись'], v['перечни'], v['док']] for k, v in summary.items()],
          [24, 34, 12, 12, 14, 14, 14, 18])
    sheet('Сверка', ['Тип ОЦ', 'ОИ', 'Раздел осмотра', 'Поле осмотра', 'Обяз.', 'Поле карточки (блок · подпись)',
                     'Подпись', 'Значения только в осмотре', 'Значения только в карточке',
                     '«По документам» — поле карточки', '«По документам» в осмотре', 'Значение в карточке',
                     'Сходится', 'Итог'],
          sverka, [20, 26, 18, 30, 7, 36, 9, 34, 34, 18, 18, 18, 9, 30], mark)
    sheet('Нет в осмотре', ['Тип ОЦ', 'ОИ', 'Блок карточки', 'Поле карточки', 'Вид', 'Значения (первые 12)'],
          absent, [20, 26, 30, 36, 14, 60])
    wb.properties.creator = 'Сборщик макета'
    wb.properties.lastModifiedBy = 'Сборщик макета'
    wb.save(path)


def main():
    path = safe_out(OUT)
    data = collect()
    sverka, absent = rows_of(data)
    build(sverka, absent, path)
    print('готово: %s — строк сверки %d, полей только в карточке %d' % (os.path.relpath(path, ROOT), len(sverka), len(absent)))


if __name__ == '__main__':
    main()
