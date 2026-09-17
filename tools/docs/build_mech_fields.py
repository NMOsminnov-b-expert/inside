# -*- coding: utf-8 -*-
"""Состав полей карточки «Механизмы и оборудование» — таблицей на согласование.

Зачем скрипт, а не документ руками. Поля заданы в справочнике
app/modules/civil/data/mechFields.js и меняются вместе с карточкой; написанная
однажды таблица разошлась бы с макетом за неделю. Здесь состав снимается с
самого справочника, поэтому пересобирается одной командой.

    python tools/docs/build_mech_fields.py

На выходе — docs/mehanizmy-polya.xlsx (локальный файл, никуда не выгружается):

    Обзор              классы и подгруппы: сколько типов и полей, нужна ли страна
    Поля по подгруппам состав полей каждой подгруппы: подпись, вид значения,
                       единицы измерения, варианты выбора
    Уточнения по типам чем отдельный тип отличается от своей подгруппы

Формат — по требованию пользователя 15.09.2026: то, что читает человек,
отдаётся в таблицах и docx, а не в markdown.
"""
import io
import json
import os
import subprocess
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
DOCS = os.path.join(ROOT, 'docs')
OUT = os.path.join(DOCS, 'mehanizmy-polya.xlsx')

HEAD_FILL = PatternFill('solid', fgColor='1F4E5F')
HEAD_FONT = Font(color='FFFFFF', bold=True, size=10)
CELL_FONT = Font(size=10)
THIN = Side(style='thin', color='D9D9D9')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

KIND = {
    'text': 'строка',
    'num': 'число',
    'int': 'целое число',
    'select': 'выбор из списка',
    'date': 'дата',
}


def dump():
    """Состав полей — из самого справочника, через node."""
    script = os.path.join(ROOT, 'tools', 'docs', 'mech_fields_dump.mjs')
    node = 'node.exe' if os.name == 'nt' else 'node'
    out = subprocess.run([node, script], cwd=os.path.dirname(script),
                         stdout=subprocess.PIPE, stderr=subprocess.PIPE)
    if out.returncode:
        sys.stderr.write(out.stderr.decode('utf-8', 'replace'))
        raise SystemExit('не удалось прочитать справочник полей')
    return json.loads(out.stdout.decode('utf-8'))


def sheet(wb, title, cols):
    ws = wb.create_sheet(title)
    ws.append([c[0] for c in cols])
    for i, (_, width) in enumerate(cols, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
        cell = ws.cell(row=1, column=i)
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical='center', wrap_text=True)
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = 'A2'
    return ws


def finish(ws):
    ws.auto_filter.ref = 'A1:%s%d' % (get_column_letter(ws.max_column), ws.max_row)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.font = CELL_FONT
            cell.border = BORDER
            cell.alignment = Alignment(vertical='top', wrap_text=True)


def field_row(f):
    # Варианты разделяем точкой с запятой: в самих вариантах встречается
    # десятичная запятая («0,2S»), и по запятой список читался бы неверно.
    return [f['label'], KIND.get(f['type'], f['type']),
            ' / '.join(f['units']), '; '.join(f['options']), f['key']]


def build(data):
    wb = Workbook()
    wb.remove(wb.active)

    over = sheet(wb, 'Обзор', [
        ('Класс', 34), ('Подгруппа', 38), ('Типов', 8),
        ('Основных полей', 15), ('Дополнительных полей', 18),
        ('Страна происхождения', 18), ('Источник', 26),
    ])
    for c in data['classes']:
        src = 'лист «Схема» (черновик)' if c['fromSchema'] else 'лист «Классификатор»'
        if not c['subgroups']:
            over.append([c['name'], '— без подгрупп —', 0, len(c['main']), len(c['extra']),
                         'да' if c['country'] else 'нет', src])
            continue
        for s in c['subgroups']:
            over.append([c['name'], s['name'], len(s['types']), len(s['main']), len(s['extra']),
                         'да' if s['country'] else 'нет', src])
    finish(over)

    cols = [('Класс', 30), ('Подгруппа', 34), ('Раздел', 16), ('Поле', 40),
            ('Вид значения', 15), ('Единицы измерения', 22), ('Варианты выбора', 52),
            ('Ключ в данных', 18)]
    subs = sheet(wb, 'Поля по подгруппам', cols)
    for c in data['classes']:
        groups = c['subgroups'] or [{'name': '— без подгрупп —', 'main': c['main'], 'extra': c['extra']}]
        for s in groups:
            for part, name in (('main', 'основные'), ('extra', 'дополнительные')):
                for f in s[part]:
                    subs.append([c['name'], s['name'], name] + field_row(f))
    finish(subs)

    types = sheet(wb, 'Уточнения по типам', [
        ('Класс', 26), ('Подгруппа', 30), ('Тип', 36), ('Что меняется', 16),
        ('Раздел', 16), ('Поле', 38), ('Вид значения', 15), ('Единицы измерения', 20),
        ('Варианты выбора', 46), ('Ключ в данных', 18),
    ])
    for c in data['classes']:
        for s in c['subgroups']:
            for bt in s['byType']:
                for part, name in (('main', 'основные'), ('extra', 'дополнительные')):
                    for f in bt[part]:
                        types.append([c['name'], s['name'], bt['type'], 'добавлено', name] + field_row(f))
                for key in bt['omit']:
                    label = next((x['label'] for x in s['main'] + s['extra'] if x['key'] == key), key)
                    types.append([c['name'], s['name'], bt['type'], 'убрано', '', label, '', '', '', key])
    finish(types)

    wb.save(OUT)
    return over.max_row - 1, subs.max_row - 1, types.max_row - 1


if __name__ == '__main__':
    if not os.path.isdir(DOCS):
        os.makedirs(DOCS)
    a, b, c = build(dump())
    io.open(sys.stdout.fileno(), 'w', encoding='utf-8', closefd=False).write(
        'собрано: %s\n  обзор — строк %d\n  поля по подгруппам — строк %d\n'
        '  уточнения по типам — строк %d\n' % (os.path.relpath(OUT, ROOT), a, b, c))
