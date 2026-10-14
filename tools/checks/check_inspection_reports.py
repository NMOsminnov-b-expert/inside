# -*- coding: utf-8 -*-
"""Отчёты за день: осмотрщик сдаёт, регионал видит, бухгалтерия выгружает.

ТЗ рабочего дня осмотрщика, раздел 3 (запись графа tz-rabota-osmotrschika).

Что ловит сценарий:
  * в шапке задачи осмотрщика нет ЕНИ (указание пользователя 12.10.2026);
  * «Сдать отчёт» без показания одометра или без снимка проходит молча
    (кнопка должна назвать, чего не хватает, — практика APCS);
  * два выезда дня с одинаковой подписью отмечаются вместе (отметки шли
    по тексту места, а не по назначению);
  * снимок одометра не сохраняется с отчётом;
  * вечерний отчёт без одометра не даёт пробег по маршруту, литры не
    считаются по расходу авто;
  * сданный отчёт не виден регионалу в «Все осмотрщики → Отчёты» и в
    правой панели;
  * выгрузка для бухгалтерии — не xlsx, без закреплённой шапки или
    автофильтра;
  * в консоли ошибки.
"""
import base64
import os
import tempfile

NAME = 'отчёты за день'

TOUCHES = ('app/pages/inspector/*', 'app/pages/inspections/*', 'app/kernel/inspections.js',
           'app/kernel/inspectionsDemo.js', 'app/kernel/xlsx.js')

# Снимок одометра — PNG 1×1: форме важен сам факт снимка.
PNG = base64.b64decode('iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mP8z8BQDwAEhQGAhKmMIQAAAABJRU5ErkJggg==')
WHO = 'Абдылдаев Айсултан Марсович'


def run(t):
    pg = t.page
    today = pg.evaluate("() => { const d = new Date(); return d.getFullYear() + '-' + String(d.getMonth() + 1).padStart(2, '0') + '-' + String(d.getDate()).padStart(2, '0'); }")
    kernel = "import('./app/kernel/inspections.js')"

    # --- ЕНИ в шапке задачи (указание пользователя 12.10.2026) -----------------
    t.open('#/insp?as=' + WHO + '&tab=tasks', wait='.ins-card-a')
    pg.locator('.ins-card-a').first.click()
    t.ck(t.wait_until("() => /\d-\d+-\d+-\d+/.test(document.querySelector('.ins-head [data-ins-eni]')?.textContent || '')"),
         'в шапке задачи осмотрщика нет ЕНИ')

    # --- утренний ---------------------------------------------------------------
    t.open('#/insp?as=' + WHO, wait='.ins-rep')
    if not t.ck(pg.locator('.ins-rep [data-rep-open="morning"]').count() == 1, 'на экране осмотрщика нет плашки «Отчёты за сегодня»'):
        return
    pg.click('[data-rep-open="morning"]')
    t.wait_for('[data-rp-send]')
    rows = pg.locator('[data-rp-place]')
    n = rows.count()
    if n >= 2:
        rows.nth(0).click()
        t.ck(pg.locator('[data-rp-place][aria-pressed="true"]').count() == n - 1,
             'снятие одной отметки «Куда выезжаю» сняло и другие')
        rows.nth(0).click()
    pg.click('[data-rp-send]')
    t.ck(t.wait_until("() => /одометр/i.test(document.querySelector('[data-rp-err]')?.textContent || '')"),
         'без показания одометра «Сдать отчёт» не сказал, чего не хватает')
    pg.fill('[data-rp="odometer"]', '120500')
    pg.click('[data-rp-send]')
    t.ck(t.wait_until("() => /Сфотографируйте/.test(document.querySelector('[data-rp-err]')?.textContent || '')"),
         'без снимка одометра отчёт не остановился')
    with pg.expect_file_chooser() as fc:
        pg.click('[data-rp-shoot]')
    fc.value.set_files({'name': 'odo.png', 'mimeType': 'image/png', 'buffer': PNG})
    t.ck(t.wait_for('.ins-rp-photo img'), 'снимок одометра не показан в форме')
    pg.click('[data-rp-send]')
    t.ck(t.wait_for('.ins-rep [data-rep-open="morning"] .ins-pill-done'), 'после сдачи утренний не отмечен «сдан»')
    m = pg.evaluate("async ([w, d]) => { const k = await %s; const r = k.reportOf(w, d, 'morning'); return r && [r.odometer, /^data:image/.test(r.photo || ''), (r.places || []).length]; }" % kernel, [WHO, today])
    t.ck(bool(m) and m[0] == '120500' and m[1], 'утренний отчёт сохранён неполно: %s' % m)

    # --- вечерний без одометра ----------------------------------------------------
    pg.click('[data-rep-open="evening"]')
    t.wait_for('[data-rp-send]')
    pg.click('[data-rp-noodo]')
    t.wait_for('[data-rp-pt]')
    pg.locator('[data-rp-pt]').nth(1).fill('ул. Советская')
    pg.click('[data-rp-send]')
    t.ck(t.wait_until("() => /пробег/i.test(document.querySelector('[data-rp-err]')?.textContent || '')"),
         'без пробега по маршруту отчёт не остановился')
    pg.fill('[data-rp="km"]', '40')
    t.ck(t.wait_until("() => /40 км/.test(document.querySelector('.ins-rp-sum')?.textContent || '')"),
         'в форме не показан пробег за день по маршруту')
    pg.click('[data-rp-send]')
    t.ck(t.wait_for('.ins-rep [data-rep-open="evening"] .ins-pill-done'), 'после сдачи вечерний не отмечен «сдан»')
    d = pg.evaluate("async ([w, d]) => { const r = (await %s).dayReport(w, d); return [r.km, r.kmBy, r.litres]; }" % kernel, [WHO, today])
    t.ck(d[0] == 40 and d[1] == 'route' and d[2] is not None and abs(d[2] - 3.4) < 0.01,
         'пробег или литры дня посчитаны неверно (ждали 40 км по маршруту, 3,4 л при 8,5 л/100): %s' % d)

    # --- регионал ---------------------------------------------------------------
    t.open('#/inspections', wait='[data-ip-section]')
    pg.click('[data-ip-section="people"]')
    pg.locator('.ip-person-main[data-ip-person=""]').click()
    pg.click('[data-ip-atab="reports"]')
    t.wait_for('[data-ip-rday-set]')
    pg.fill('[data-ip-rday-set]', today)
    pg.dispatch_event('[data-ip-rday-set]', 'change')
    row = '.ip-reptbl tr[data-ip-dp-open="report:%s|%s"]' % (WHO, today)
    if not t.ck(t.wait_for(row), 'в «Отчётах» за сегодня нет строки осмотрщика'):
        return
    t.ck(pg.locator(row + ' .ip-rtime, ' + row + ' .ip-rmark.warn').count() >= 2, 'в строке не показано время сдачи обоих отчётов')
    pg.click(row)
    t.ck(t.wait_until("() => /40 км/.test(document.querySelector('.ip-detail-pane')?.innerText || '') && /по маршруту/.test(document.querySelector('.ip-detail-pane')?.innerText || '')"),
         'правая панель не показала пробег дня по маршруту')
    t.ck(pg.locator('.ip-detail-pane .ip-odo img').count() == 1, 'в правой панели нет снимка одометра')
    t.ck(pg.evaluate('document.documentElement.scrollWidth <= innerWidth'), 'страница панели прокручивается вбок')

    # --- выгрузка для бухгалтерии --------------------------------------------------
    pg.fill('[data-ip-rexp="to"]', today)
    pg.dispatch_event('[data-ip-rexp="to"]', 'change')
    with pg.expect_download(timeout=15000) as dl:
        pg.click('[data-ip-rexp-go]')
    path = os.path.join(tempfile.gettempdir(), '_tmp_check_probeg.xlsx')
    dl.value.save_as(path)
    try:
        import openpyxl
        wb = openpyxl.load_workbook(path)
        ws = wb['По дням']
        t.ck(ws.freeze_panes == 'A2' and bool(ws.auto_filter.ref), 'у листа выгрузки нет закреплённой шапки или автофильтра')
        found = [r for r in ws.iter_rows(min_row=2, values_only=True) if r[0] == WHO and r[1] == '.'.join(reversed(today.split('-')))]
        t.ck(len(found) == 1 and found[0][8] == 40 and found[0][9] == 'по маршруту',
             'в выгрузке нет сегодняшней строки с пробегом 40 км по маршруту: %s' % (found,))
        t.ck('Итог' in wb.sheetnames, 'в выгрузке нет листа «Итог»')
    finally:
        try:
            os.remove(path)
        except OSError:
            pass
