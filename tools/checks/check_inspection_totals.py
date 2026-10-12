# -*- coding: utf-8 -*-
"""Итоги по сотрудникам за период и их выгрузка.

ТЗ рабочего дня осмотрщика, раздел 9 (запись графа tz-rabota-osmotrschika).

Что ловит сценарий:
  * итоговая строка не равна сумме строк (пробег);
  * пробег в итогах расходится с отчётами за день за тот же период;
  * смена периода не пересчитывает итоги;
  * строка не открывает отчёты осмотрщика;
  * выгрузка — не xlsx с листами «Итоги» и «По дням», без закреплённой
    шапки;
  * в консоли ошибки.
"""
import os
import tempfile

NAME = 'итоги по сотрудникам'

TOUCHES = ('app/pages/inspections/*', 'app/kernel/inspections.js', 'app/kernel/xlsx.js')


def num(s):
    s = (s or '').replace(' ', '').replace(' ', '').replace(',', '.')
    return float(s) if s and s != '—' else 0.0


def run(t):
    pg = t.page
    t.open('#/inspections', wait='[data-ip-section]')
    pg.click('[data-ip-section="people"]')
    pg.locator('.ip-person-main[data-ip-person=""]').click()
    pg.click('[data-ip-atab="totals"]')
    if not t.ck(t.wait_for('.ip-sumtbl tbody tr'), 'нет таблицы итогов'):
        return
    km_col = 6   # Осмотрщик, Дней, Объектов, ОИ, Снимков, Км
    rows = pg.locator('.ip-sumtbl tbody tr')
    total = sum(num(rows.nth(i).locator('td').nth(km_col - 1).inner_text()) for i in range(rows.count()))
    foot = num(pg.locator('.ip-sumtbl tfoot td').nth(km_col - 2).inner_text())
    t.ck(abs(total - foot) < 0.5, 'итог пробега %s не равен сумме строк %s' % (foot, total))

    frm, to = pg.input_value('[data-ip-sum="from"]'), pg.input_value('[data-ip-sum="to"]')
    who = rows.first.get_attribute('data-ip-sum-open')
    km_model = pg.evaluate("""async ([w, f, t]) => { const m = await import('./app/pages/inspections/model.js');
      return m.reportTotals(m.reportDays(w, f, t)).km; }""", [who, frm, to])
    t.ck(abs(num(rows.first.locator('td').nth(km_col - 1).inner_text()) - km_model) < 0.5, 'пробег в итогах расходится с отчётами за день')

    # период — один день: пробег меньше или равен
    pg.fill('[data-ip-sum="from"]', to)
    pg.dispatch_event('[data-ip-sum="from"]', 'change')
    t.ck(t.wait_until("(v) => document.querySelector('[data-ip-sum=\"from\"]').value === v", to), 'период не применился')
    foot1 = num(pg.locator('.ip-sumtbl tfoot td').nth(km_col - 2).inner_text())
    t.ck(foot1 <= foot, 'смена периода не пересчитала итоги: %s → %s' % (foot, foot1))

    # выгрузка
    with pg.expect_download(timeout=15000) as dl:
        pg.click('[data-ip-sum-xlsx]')
    out = os.path.join(tempfile.gettempdir(), '_tmp_check_itogi.xlsx')
    dl.value.save_as(out)
    import openpyxl
    wb = openpyxl.load_workbook(out)
    t.ck(wb.sheetnames == ['Итоги', 'По дням'] and wb['Итоги'].freeze_panes == 'A2' and wb['Итоги'].max_row >= 2,
         'выгрузка итогов неверна: %s' % wb.sheetnames)
    os.remove(out)

    # строка — к отчётам осмотрщика
    pg.locator('.ip-sumtbl tbody tr').first.click()
    t.ck(t.wait_until("() => document.querySelector('[data-ip-ptab=\"reports\"].on') !== null"), 'строка итогов не открыла отчёты осмотрщика')
    t.ck(pg.evaluate('document.documentElement.scrollWidth <= innerWidth'), 'страница прокручивается вбок')
