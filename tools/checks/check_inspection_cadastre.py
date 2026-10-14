# -*- coding: utf-8 -*-
"""Запросы документов: регионал просит, осмотрщик отвечает, файл — в карточке.

ТЗ рабочего дня осмотрщика, раздел 5 (запись графа tz-rabota-osmotrschika).

Что ловит сценарий:
  * запрос в кадастр можно отдать осмотрщику без доступа в кадастр;
  * запрос без строк или без осмотрщика создаётся молча;
  * «Нет в кадастре» без причины и «Получен» без PDF или без даты
    сканирования сохраняются молча;
  * полученный PDF не ложится в документы объекта или ложится без даты
    сканирования; имя файла — не «№ + ЕНИ + название»;
  * регионал не видит итог строки; отчёт по запросу — не xlsx;
  * в консоли ошибки.
"""
import os
import tempfile

from check_archive import make_pdf

NAME = 'запросы документов'

TOUCHES = ('app/pages/inspector/*', 'app/pages/inspections/*', 'app/kernel/inspections.js',
           'app/kernel/inspectionsDemo.js', 'app/kernel/institutions.js')

K = "import('./app/kernel/inspections.js')"


def run(t):
    pg = t.page
    pdf = os.path.join(tempfile.gettempdir(), '_tmp_check_cad.pdf')
    make_pdf(pdf, pages=1)
    try:
        t.open('#/inspections', wait='[data-ip-section]')
        pg.click('[data-ip-section="cadastre"]')
        pg.click('[data-ip-cad-new]')
        t.wait_for('[data-ip-cad-create]')
        pg.click('[data-ip-cad-create]')
        t.ck(t.wait_until("() => /объект/.test(document.querySelector('[data-ip-cad-err]')?.textContent || '')"),
             'запрос без строк не остановился')

        with_access = pg.evaluate("async () => { const k = await %s; return k.profileNames().filter((n) => k.profileOf(n).cadastre); }" % K)
        without = pg.evaluate("async () => { const k = await %s; return k.profileNames().filter((n) => !k.profileOf(n).cadastre); }" % K)
        if not t.ck(with_access and without, 'в демо нет осмотрщиков с доступом в кадастр и без него: %s / %s' % (with_access, without)):
            return
        t.ck(pg.locator('[data-ip-cadf="person"] option[value="%s"][disabled]' % without[0]).count() == 1,
             'осмотрщика без доступа в кадастр можно выбрать для запроса в кадастр')
        who = with_access[0]

        pg.fill('[data-ip-cad-q]', '14756')
        t.wait_for('[data-ip-cad-add]')
        pg.locator('[data-ip-cad-add]').first.click()
        t.ck(t.wait_for('[data-ip-cad-row="0"][data-k="doc"]'), 'объект не добавился строкой запроса')
        pg.click('[data-ip-cad-create]')
        t.ck(t.wait_until("() => /осмотрщик/i.test(document.querySelector('[data-ip-cad-err]')?.textContent || '')"),
             'запрос без осмотрщика не остановился')
        pg.select_option('[data-ip-cadf="person"]', who)
        pg.fill('[data-ip-cad-row="0"][data-k="note"]', 'проверка')
        pg.click('[data-ip-cad-create]')
        t.ck(t.wait_until("() => /№ \\d+/.test(document.querySelector('.ip-card h1')?.textContent || '')"), 'запрос не создан')
        req = pg.evaluate("async () => { const r = (await %s).allCadRequests()[0]; return [r.id, r.rows[0].id, r.rows[0].typeId, r.rows[0].ocId, r.rows[0].eni, r.rows[0].doc]; }" % K)
        rid, rowid, type_id, oc_id, eni, doc = req

        # --- осмотрщик ------------------------------------------------------------
        # Запросы осмотрщика — свой раздел нижней панели (14.10.2026).
        t.open('#/insp?as=' + who + '&tab=cad', wait='.ins-body')
        t.ck(t.wait_for('[data-cad-open="%s"]' % rid), 'у осмотрщика нет нового запроса в «Запросах документов»')
        pg.click('[data-cad-open="%s"]' % rid)
        t.wait_for('[data-cad-row="%s"]' % rowid)
        pg.click('[data-cad-res="%s|none"]' % rowid)
        pg.click('[data-cad-save="%s"]' % rowid)
        t.ck(t.wait_until("() => /причину/.test(document.querySelector('.ins-rp-err')?.textContent || '')"), '«Нет в кадастре» без причины не остановилось')
        pg.click('[data-cad-res="%s|got"]' % rowid)
        pg.click('[data-cad-save="%s"]' % rowid)
        t.ck(t.wait_until("() => /PDF/.test(document.querySelector('.ins-rp-err')?.textContent || '')"), '«Получен» без PDF не остановилось')
        with pg.expect_file_chooser() as fc:
            pg.click('[data-cad-file="%s"]' % rowid)
        fc.value.set_files(pdf)
        t.wait_for('.ins-cad-fname')
        pg.click('[data-cad-save="%s"]' % rowid)
        t.ck(t.wait_until("() => /сканирования/.test(document.querySelector('.ins-rp-err')?.textContent || '')"), '«Получен» без даты сканирования не остановилось')
        pg.fill('[data-cad-scan="%s"]' % rowid, '2026-10-01')
        pg.click('[data-cad-save="%s"]' % rowid)
        t.ck(t.wait_for('[data-cad-row="%s"] .ins-ro' % rowid), 'ответ не сохранён')

        got = pg.evaluate("""async ([rid, rowid, tid, oid]) => {
          const k = await %s; const reg = await import('./app/kernel/registry.js');
          const row = k.cadRequest(rid).rows.find((x) => x.id === rowid);
          const rec = reg.getType(tid).records.loadRecord(oid);
          const d = (rec.docs || []).find((x) => x.id === (row.file || {}).docId);
          return [row.result, row.scanDate, (row.file || {}).name, d ? [d.date, d.type, d.name, !!d.file] : null];
        }""" % K, [rid, rowid, type_id, oc_id])
        t.ck(got[0] == 'got' and got[1] == '2026-10-01', 'итог строки сохранён неверно: %s' % got)
        t.ck(bool(got[2]) and eni in got[2] and doc in got[2] and got[2].endswith('.pdf'),
             'имя файла не «№ + ЕНИ + название»: %s' % got[2])
        t.ck(got[3] is not None and got[3][0] == '2026-10-01' and got[3][1] == doc and got[3][3],
             'файл не лёг в документы объекта с датой сканирования: %s' % (got[3],))

        # --- регионал видит итог и выгружает отчёт ------------------------------------
        t.open('#/inspections', wait='[data-ip-section]')
        pg.click('[data-ip-section="cadastre"]')
        pg.click('[data-ip-cad="%s"]' % rid)
        t.ck(t.wait_until("() => [...document.querySelectorAll('.ip-cadtbl .ip-rmark.ok')].length === 1"), 'регионал не видит «Получен» в строке запроса')
        with pg.expect_download(timeout=15000) as dl:
            pg.click('[data-ip-cad-xlsx="%s"]' % rid)
        out = os.path.join(tempfile.gettempdir(), '_tmp_check_cad.xlsx')
        dl.value.save_as(out)
        import openpyxl
        ws = openpyxl.load_workbook(out).active
        t.ck(ws.freeze_panes == 'A2' and ws.max_row == 2 and ws.cell(2, 7).value == 'Получен', 'отчёт по запросу неверен')
        os.remove(out)
        t.ck(pg.evaluate('document.documentElement.scrollWidth <= innerWidth'), 'страница прокручивается вбок')
    finally:
        try:
            os.remove(pdf)
        except OSError:
            pass
