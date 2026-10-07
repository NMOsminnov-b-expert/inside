# -*- coding: utf-8 -*-
"""Поиск «Тип ТС, вид кузова» карточки ТС — все сценарии подряд.

Просьба пользователя 07.10.2026: «Мб тестики? На все сценарии. Даже самые
конченые…» — после замечаний «вбил седан, а мне заместо „легковой, седан“
выдало „легковой“» и «на кабриолет внедорожник предложило».

Поиск гоняется в браузере на тех же функциях, что у карточки: выдача —
templates.js searchLeaves, отбор — kernel/treeSearch.js findLeaves, выбор —
templates.js pickLeaf (что встаёт в «Тип ТС» и в категорию). Выбор — первый
пункт выдачи, как по Enter.

Что ловит сценарий:
  * запись техпаспорта по слову кузова («кабриолет», «седан», «самосвал») —
    после выбора первого пункта в «Тип ТС» нет этого кузова (встала голая
    запись базы или чужая);
  * полная запись («легковой, седан») — после выбора в поле не она;
  * шаблон по своему названию — не первый; другое название шаблона («воровайка»)
    не находит его в верхних строках;
  * база или вид спецтехники по названию — не первые или ставят не ту запись;
  * модель по «марка модель» — не первая;
  * запись из обезличенной сверки техпаспортов — после выбора поле пустое;
  * мусорный ввод (пусто, пробелы, знаки, латиница, раскладка, опечатки,
    разметка, 500 символов, цифры, повторы, капслок, «ё», кавычки) — ошибка,
    выдача длиннее предела, выбор ломает карточку;
  * живое поле: вставка разметки создаёт элементы в выдаче, Enter без выдачи,
    стрелки и Escape — ошибка в консоли или поле не работает.
"""
import json
import os

NAME = 'поиск «Тип ТС»: все сценарии'

TOUCHES = ('app/kernel/treeSearch.js', 'app/modules/vehicle/templates.js', 'app/modules/vehicle/tsModel.js',
           'app/modules/vehicle/ctrl.js', 'app/modules/vehicle/data/tsCatalog.js', 'app/modules/vehicle/data/tsModels.js',
           'tools/data/ts_templates.py', 'tools/data/build_ts_catalog.py')

ROOT = os.path.dirname(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

# Реальные записи из сверки (столбец I обезличенной книги), если книга есть.
def sverka_values():
    p = os.path.join(ROOT, 'docs', 'sverka-tipov-ts.xlsx')
    if not os.path.exists(p):
        return []
    try:
        from openpyxl import load_workbook
        ws = load_workbook(p, read_only=True)['Вид кузова (I)']
        return [str(r[0]) for r in ws.iter_rows(min_row=2, values_only=True) if r[0]]
    except Exception:  # noqa: BLE001 — книги может не быть на машине
        return []


GARBAGE = ['', '   ', 'ё', 'Ё', 'СЕДАН', 'седан!!!', '«седан»', '"седан"', 'седан седан седан', 'сидан', 'sedan',
           'ctlfy', 'cfvjcdfk', '<img src=x onerror=alert(1)>', '</div><b>x', 'а' * 500, '123', '0', '%', '()', '\\',
           '[', '*', '+', '?', '.', '-', '—', ',', ', , ,', 'кабриолет кабриолет кабриолет', 'грузовой самосвал прицеп',
           'камаз камаз', 'KAMAZ 65115', 'камаз65115', '65115-62', 'газ-53', 'газ 53 самосвал', 'экскаватор колесный гусеничный',
           'трал полуприцеп трал', 'мото коляска', 'х', 'й', 'ъ', 'ь', '  седан  ', '\tседан', 'седан\n', '😀', '№', 'кр № 123']

JS = r"""async (input) => {
  const [T, M, C, K] = await Promise.all([import('./app/modules/vehicle/templates.js'), import('./app/modules/vehicle/tsModel.js'),
    import('./app/modules/vehicle/data/tsCatalog.js'), import('./app/kernel/treeSearch.js')]);
  const fresh = () => ({ kind: '', f: {}, extra: [], modules: [] });
  const words = (s) => String(s || '').toLowerCase().replace(/ё/g, 'е').split(/[^a-zа-я0-9]+/).filter(Boolean);
  const has = (outcome, query) => { const ow = words(outcome); return words(query).every((w) => ow.some((x) => x.startsWith(w.slice(0, Math.max(4, w.length - 2))))); };
  const search = (v, q) => K.findLeaves(T.searchLeaves(v), q);
  const pickFirst = (q) => { const v = fresh(); const r = search(v, q); const l = r.list[0]; if (!l) return { v, l: null, r };
    T.pickLeaf(v, l, q); return { v, l, r }; };
  const fail = {}; const n = {};
  const bad = (g, x) => { (fail[g] = fail[g] || []).push(x); };
  const count = (g) => { n[g] = (n[g] || 0) + 1; };

  // Записи техпаспорта: по слову кузова и полной записью.
  const recs = [];
  Object.entries(C.TS_VTYPE_BY_CATEGORY).forEach(([cat, list]) => { if (cat !== 'По техпаспорту') list.forEach((r) => recs.push([cat, r])); });
  const seenBody = new Set();
  recs.forEach(([cat, r]) => {
    const body = r.includes(', ') ? r.slice(r.indexOf(', ') + 2) : r;
    if (!seenBody.has(body)) {
      seenBody.add(body); count('кузов');
      // Шаблон ставит свою запись (решение пользователя 07.10.2026: «шаблон
      // выбран — запись шаблона»), и пункт это подписывает; база и запись —
      // запись с набранным кузовом.
      const o = pickFirst(body);
      const ok = o.l && (o.l.tpl ? o.v.f.vtype === o.l.tpl.vtype : has(o.v.f.vtype, body));
      if (!ok) bad('кузов', `«${body}» → ${o.l ? o.l.name : 'ничего'} → «${o.v.f.vtype || ''}»`);
    }
    count('запись');
    const o = pickFirst(r);
    const ok = o.l && (o.l.tpl ? o.v.f.vtype === o.l.tpl.vtype : has(o.v.f.vtype, r));
    if (!ok) bad('запись', `«${r}» → ${o.l ? o.l.name : 'ничего'} → «${o.v.f.vtype || ''}»`);
  });

  // Шаблоны по названию и по другим названиям.
  C.TS_TEMPLATES.forEach((t) => {
    count('шаблон');
    const l = search(fresh(), t.name).list[0];
    if (!l || l.name !== t.name) bad('шаблон', `«${t.name}» → ${l ? l.name : 'ничего'}`);
  });
  const aliases = {};
  C.TS_TEMPLATES.forEach((t) => (t.aliases || []).forEach((a) => { (aliases[a] = aliases[a] || []).push(t.name); }));
  Object.entries(aliases).forEach(([a, names]) => {
    count('другое название');
    const top = search(fresh(), a).list.slice(0, 10).map((l) => l.name);
    if (!top.some((x) => names.includes(x))) bad('другое название', `«${a}» → ${top.slice(0, 3).join(' | ')}`);
  });

  // Базы и виды спецтехники.
  M.kindLeaves().forEach((k) => {
    count('база и вид');
    const o = pickFirst(k.name);
    // Первым может быть и запись техпаспорта того же текста («Автобус» →
    // «автобус») — важно, что встаёт запись базы или вида.
    const want = M.kindVtype(k);
    if (!o.l || (want ? o.v.f.vtype !== want : o.l.name !== k.name))
      bad('база и вид', `«${k.name}» → ${o.l ? o.l.name : 'ничего'} → «${o.v.f.vtype || ''}» (ждали «${want}»)`);
  });

  // Модели.
  const models = [...new Set(T.modelLeaves().map((l) => l.name))];
  // Одна модель в данных бывает записана по-разному («82.1» и «82,1») — без
  // учёта знаков это одно имя.
  const same = (a, b) => words(a).join(' ') === words(b).join(' ');
  models.forEach((name) => {
    count('модель');
    const l = search(fresh(), name).list[0];
    if (!l || !same(l.name, name)) bad('модель', `«${name}» → ${l ? l.name : 'ничего'}`);
  });

  // Записи из сверки техпаспортов: выбор первого не оставляет поле пустым.
  input.sverka.forEach((q) => {
    count('сверка');
    const o = pickFirst(q);
    if (o.l && !String(o.v.f.vtype || '').trim()) bad('сверка', `«${q}» → ${o.l.name} → пусто`);
  });

  // Мусорный ввод: без ошибок, выдача не длиннее предела, выбор не ломает карточку.
  input.garbage.forEach((q) => {
    count('мусор');
    try {
      const r = search(fresh(), q);
      if (r.list.length > 40) bad('мусор', `«${q.slice(0, 20)}» → ${r.list.length} пунктов`);
      if (r.list[0]) { const v = fresh(); T.pickLeaf(v, r.list[0], q); if (typeof v.f.vtype !== 'string' && v.f.vtype !== undefined) bad('мусор', `«${q.slice(0, 20)}» → тип ${typeof v.f.vtype}`); }
    } catch (e) { bad('мусор', `«${q.slice(0, 20)}» → ошибка ${e.message}`); }
  });
  return { fail, n };
}"""


def run(t):
    pg = t.page
    t.open('', wait='.reg-tr')
    pg.click('.reg-create [data-dd-toggle]')
    pg.click('.reg-create [data-create="vehicle"]')
    t.wait_for('#ts-find-q')

    res = pg.evaluate(JS, {'garbage': GARBAGE, 'sverka': sverka_values()})
    for group, total in res['n'].items():
        bad = res['fail'].get(group, [])
        t.ck(not bad, '%s: %d из %d — %s' % (group, len(bad), total, '; '.join(bad[:8])))

    # Живое поле: разметка не превращается в элементы, Enter без выдачи, стрелки, Escape.
    q = '#ts-find-q'
    pg.fill(q, '<img src=x onerror="window.__xss=1">')
    t.wait_until("() => !document.querySelector('#ts-find-list').hidden")
    t.ck(pg.locator('#ts-find-list img').count() == 0 and not pg.evaluate('() => window.__xss'),
         'разметка из поля попала в выдачу элементами')
    pg.fill(q, 'щщщщщщ')
    t.wait_until("() => !document.querySelector('#ts-find-list').hidden")
    pg.keyboard.press('Enter')
    t.ck(pg.input_value(q) == 'щщщщщщ', 'Enter без выдачи изменил поле')
    pg.fill(q, 'самосвал')
    t.wait_until("() => document.querySelectorAll('#ts-find-list [role=option]').length > 2")
    pg.keyboard.press('ArrowDown')
    pg.keyboard.press('ArrowUp')
    pg.keyboard.press('ArrowUp')
    sel = pg.evaluate("() => [...document.querySelectorAll('#ts-find-list [role=option]')].findIndex((e) => e.getAttribute('aria-selected') === 'true')")
    n = pg.locator('#ts-find-list [role=option]').count()
    t.ck(sel == n - 1, 'стрелка вверх с первого пункта не ушла на последний: %s из %s' % (sel, n))
    pg.keyboard.press('Escape')
    t.ck(pg.locator('#ts-find-list').is_hidden(), 'Escape не закрыл выдачу')
    pg.fill(q, 'а' * 500)
    pg.keyboard.press('Tab')
    t.ck(len(pg.input_value(q)) == 500, 'длинный ввод обрезан или потерян')
