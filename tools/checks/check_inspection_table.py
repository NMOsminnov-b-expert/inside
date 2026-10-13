# -*- coding: utf-8 -*-
"""Табличный вид осмотрщиков «как в Excel» и загрузка xlsx со сверкой.

ТЗ рабочего дня осмотрщика, раздел 10 (запись графа tz-rabota-osmotrschika);
уточнение пользователя 12.10.2026 — на весь экран и работа как в Excel.

Что ловит сценарий:
  * ввод в ячейку не сохраняется в профиль; Enter не переходит вниз;
  * Esc не отменяет правку; негодная дата сохраняется;
  * блок из Excel (табуляция и перевод строки) не вставляется блоком;
    копирование отдаёт не табличный текст;
  * Ctrl+Z не отменяет; сортировка и фильтр по столбцу не работают;
  * «На весь экран» не разворачивает, Esc не сворачивает;
  * выгрузка и загрузка обратно: изменения применяются без показа или
    «изменить» применяется без отметки;
  * «+ поле» не заводит столбец своего поля прямо в таблице; перетаскивание
    заголовка не меняет порядок столбцов или порядок теряется после
    перезагрузки (указание пользователя 13.10.2026 «свои поля добавляем
    прямо в таблице. Менять поля местами можем там же»);
  * в консоли ошибки.
"""
import os
import tempfile

NAME = 'табличный вид осмотрщиков'

TOUCHES = ('app/kernel/grid.js', 'app/kernel/xlsx.js', 'app/pages/inspections/*', 'app/kernel/inspections.js')

K = "import('./app/kernel/inspections.js')"


def col_index(pg, label):
    return pg.evaluate("(l) => [...document.querySelectorAll('.kg thead tr:first-child th')].findIndex((th) => th.textContent.trim().replace(/ [↑↓]$/, '') === l)", label)


def prof(pg, name, expr):
    return pg.evaluate("async (n) => { const p = (await %s).profileOf(n); return %s; }" % (K, expr), name)


def run(t):
    pg = t.page
    t.open('#/inspections', wait='[data-ip-section]')
    pg.click('[data-ip-section="people"]')
    pg.locator('.ip-person-main[data-ip-person=""]').click()
    pg.click('[data-ip-atab="table"]')
    if not t.ck(t.wait_for('.kg td[data-r="0"]'), 'табличный вид не открылся'):
        return
    name0 = pg.locator('.kg td[data-r="0"][data-c="0"]').inner_text()
    name1 = pg.locator('.kg td[data-r="1"][data-c="0"]').inner_text()
    ce = col_index(pg, 'Почта')
    cp = col_index(pg, 'Телефон')
    cb = col_index(pg, 'Дата рождения')

    # ввод — с первого знака, Enter — сохранить и вниз
    pg.click('.kg td[data-r="0"][data-c="%d"]' % ce)
    pg.keyboard.type('q@test.kg')
    pg.keyboard.press('Enter')
    t.ck(t.wait_until("async (n) => (await %s).profileOf(n).email === 'q@test.kg'" % K, name0), 'ввод в ячейку не сохранился в профиль')
    t.ck(pg.locator('.kg td.act[data-r="1"][data-c="%d"]' % ce).count() == 1, 'Enter не перевёл на строку ниже')

    # F2 и Esc — без изменений
    pg.keyboard.press('ArrowUp')
    pg.keyboard.press('F2')
    pg.keyboard.type('zzz')
    pg.keyboard.press('Escape')
    t.ck(prof(pg, name0, 'p.email') == 'q@test.kg', 'Esc не отменил правку')

    # негодная дата — не сохраняется
    pg.click('.kg td[data-r="0"][data-c="%d"]' % cb)
    pg.keyboard.type('31.31.2026')
    pg.keyboard.press('Enter')
    t.ck(t.wait_until("() => /дата как/.test(document.querySelector('.kg-status')?.textContent || '')"), 'негодная дата не остановилась')
    pg.keyboard.press('Escape')
    t.ck(prof(pg, name0, 'p.birthDate') != '2026-31-31', 'негодная дата сохранилась')

    # вставка блока 2×2 из Excel: Телефон и WhatsApp в строках 0–1
    pg.click('.kg td[data-r="0"][data-c="%d"]' % cp)
    pg.evaluate("""() => { const dt = new DataTransfer(); dt.setData('text/plain', '+996 700 000 001\\t+996 700 000 011\\r\\n+996 700 000 002\\t+996 700 000 022\\r\\n');
      document.querySelector('.kg').dispatchEvent(new ClipboardEvent('paste', { clipboardData: dt, bubbles: true, cancelable: true })); }""")
    t.ck(t.wait_until("async ([a, b]) => { const k = await %s; return k.profileOf(a).phone === '+996 700 000 001' && k.profileOf(b).whatsapp === '+996 700 000 022'; }" % K, [name0, name1]),
         'блок из Excel не вставился блоком 2×2')
    copied = pg.evaluate("""() => { const dt = new DataTransfer();
      document.querySelector('.kg').dispatchEvent(new ClipboardEvent('copy', { clipboardData: dt, bubbles: true, cancelable: true })); return dt.getData('text/plain'); }""")
    t.ck(copied.split('\r\n')[1].split('\t') == ['+996 700 000 002', '+996 700 000 022'], 'копирование отдаёт не таблицу: %r' % copied)

    # Ctrl+Z — отменяет вставку
    pg.keyboard.press('Control+z')
    t.ck(t.wait_until("async (a) => (await %s).profileOf(a).phone !== '+996 700 000 001'" % K, name0), 'Ctrl+Z не отменил вставку')

    # сортировка и фильтр
    pg.click('.kg [data-kg-sort="name"]')
    pg.click('.kg [data-kg-sort="name"]')
    # Сортировка — внутри раздела (разделы — по ответственному лицу).
    first = pg.evaluate("() => { const out = []; let n = 0; for (const tr of document.querySelectorAll('.kg tbody tr')) { if (tr.classList.contains('kg-g1')) { n++; if (n > 1) break; } else if (!tr.classList.contains('kg-g')) out.push(tr.querySelector('td[data-c=\"0\"]').textContent); } return out; }")
    t.ck(len(first) >= 2 and first == sorted(first, reverse=True), 'сортировка по убыванию внутри раздела не сработала: %s' % first)
    pg.fill('.kg [data-kg-filter="name"]', name1[:4])
    t.ck(t.wait_until("() => document.querySelectorAll('.kg tbody td[data-c=\"0\"]').length === 1"), 'фильтр по столбцу не сработал')
    pg.fill('.kg [data-kg-filter="name"]', '')

    # разделы по ответственным лицам, подразделы по областям, разные цвета
    sections = pg.locator('.kg tr.kg-g1').count()
    t.ck(sections >= 2 and pg.locator('.kg tr.kg-g2').count() >= sections, 'таблица не разделена на разделы и подразделы: %s' % sections)
    colors = pg.evaluate("() => [...document.querySelectorAll('.kg tr.kg-g1')].map((tr) => getComputedStyle(tr.querySelector('th')).backgroundColor)")
    t.ck(len(set(colors)) == len(colors), 'разделы не различаются цветом: %s' % colors)
    mgr = pg.locator('.kg tr.kg-g1 b').first.inner_text()
    pg.select_option('.kg [data-kg-filter="manager"]', mgr)
    t.ck(t.wait_until("() => document.querySelectorAll('.kg tr.kg-g1').length === 1"), 'фильтр «Ответственное лицо» не оставил один раздел')
    pg.select_option('.kg [data-kg-filter="manager"]', '')
    # свернуть раздел — строки уходят, заголовок остаётся
    rows0 = pg.locator('.kg tbody td[data-c="0"]').count()
    pg.locator('.kg tr.kg-g1 [data-kg-group]').first.click()
    t.ck(pg.locator('.kg tbody td[data-c="0"]').count() < rows0 and pg.locator('.kg tr.kg-g1').count() == sections, 'раздел не свернулся')
    pg.locator('.kg tr.kg-g1 [data-kg-group]').first.click()
    # смена ответственного — выбором из списка; строка уходит в другой раздел
    cm = col_index(pg, 'Ответственное лицо')
    who = pg.locator('.kg td[data-r="0"][data-c="0"]').inner_text()
    other = pg.locator('.kg tr.kg-g1 b').nth(1).inner_text()
    pg.click('.kg td[data-r="0"][data-c="%d"]' % cm)
    pg.keyboard.press('F2')
    t.ck(t.wait_for('.kg-pick li'), 'у поля-выбора не открылся список')
    # Значение не из списка не принимается (замечание пользователя: «где поля —
    # выбор, то в таблице тоже должен быть выбор, а не текст»).
    pg.keyboard.type('zzz')
    t.ck(pg.locator('.kg-pick .kg-pick-none').count() == 1, 'список не сузился по вводу')
    pg.keyboard.press('Enter')
    t.ck(prof(pg, who, 'p.manager') != 'zzz', 'в поле-выбор попало значение не из списка')
    pg.click('.kg td[data-r="0"][data-c="%d"]' % cm)
    pg.keyboard.press('F2')
    pg.keyboard.type(other[:4])
    pg.keyboard.press('Enter')
    t.ck(t.wait_until("async ([w, m]) => (await %s).profileOf(w).manager === m" % K, [who, other]), 'ответственное лицо не выбралось из списка')
    t.ck(t.wait_until("(w) => { let sec = ''; for (const tr of document.querySelectorAll('.kg tbody tr')) { if (tr.classList.contains('kg-g1')) sec = tr.textContent; else if (tr.querySelector('td[data-c=\"0\"]')?.textContent === w) return sec; } return ''; }", who),
         'строка не перешла в раздел нового ответственного')
    # области — выбор нескольких: пробел отмечает, Enter сохраняет
    cr = col_index(pg, 'Области')
    pg.click('.kg td[data-r="0"][data-c="%d"]' % cr)
    who2 = pg.locator('.kg td[data-r="0"][data-c="0"]').inner_text()
    before = prof(pg, who2, 'p.regions.length')
    pg.keyboard.press('F2')
    pg.keyboard.type('Нарын')
    pg.keyboard.press(' ')
    pg.keyboard.press('Enter')
    t.ck(t.wait_until("async ([w, n]) => { const r = (await %s).profileOf(w).regions; return r.includes('Нарын') && r.length === n + 1; }" % K, [who2, before]),
         'выбор нескольких областей не сохранился')
    # подытогов-дублей нет; подпись раздела держится слева при прокрутке
    t.ck(pg.locator('.kg tr.kg-sub').count() == 0, 'остались строки «Итого» — дубль числа в заголовке')
    pg.evaluate("document.querySelector('.kg-scroll').scrollLeft = 900")
    stays = pg.evaluate("() => { const g = document.querySelector('.kg tr.kg-g1 .kg-gh').getBoundingClientRect(); const s = document.querySelector('.kg-scroll').getBoundingClientRect(); return g.left - s.left < 20; }")
    t.ck(stays, 'подпись раздела уехала при прокрутке вбок')
    pg.evaluate("document.querySelector('.kg-scroll').scrollLeft = 0")

    # на весь экран и обратно
    pg.click('[data-ip-grid-full]')
    t.ck(pg.locator('.ip-gridwrap.full').count() == 1, '«На весь экран» не развернул таблицу')
    pg.locator('.kg td[data-r="0"][data-c="1"]').click()
    pg.keyboard.press('Escape')
    t.ck(t.wait_until("() => !document.querySelector('.ip-gridwrap.full')"), 'Esc не свернул таблицу')

    # выгрузка → правка файла → загрузка со сверкой
    with pg.expect_download(timeout=15000) as dl:
        pg.click('[data-ip-grid-xlsx]')
    path = os.path.join(tempfile.gettempdir(), '_tmp_check_grid.xlsx')
    dl.value.save_as(path)
    import openpyxl
    wb = openpyxl.load_workbook(path)
    ws = wb.active
    t.ck(ws.freeze_panes == 'A2' and ws.cell(1, 1).value == 'ФИО', 'выгрузка таблицы без шапки или закрепления')
    head = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, 1).value == name0:
            ws.cell(r, head.index('Почта') + 1).value = 'new@test.kg'          # изменить
        if ws.cell(r, 1).value == name1:
            ws.cell(r, head.index('Устройство') + 1).value = 'Redmi Note 12'  # дописать
    wb.save(path)
    with pg.expect_file_chooser() as fc:
        pg.click('[data-ip-grid-import]')
    fc.value.set_files(path)
    t.ck(t.wait_for('[data-ip-imp-apply]'), 'после загрузки нет разбора изменений')
    t.ck(pg.locator('[data-ip-imp]:checked').count() == 1 and pg.locator('[data-ip-imp]').count() == 2,
         '«дописать» должно быть отмечено, «изменить» — нет: отмечено %s из %s' % (pg.locator('[data-ip-imp]:checked').count(), pg.locator('[data-ip-imp]').count()))
    pg.click('[data-ip-imp-apply]')
    t.ck(t.wait_until("async ([a, b]) => { const k = await %s; return k.profileOf(b).device === 'Redmi Note 12' && k.profileOf(a).email === 'q@test.kg'; }" % K, [name0, name1]),
         'загрузка применила не то: дописать — да, изменить без отметки — нет')
    os.remove(path)
    t.ck(pg.evaluate('document.documentElement.scrollWidth <= innerWidth'), 'страница прокручивается вбок')

    # --- своё поле прямо из таблицы и порядок столбцов перетаскиванием ---------------
    pg.click('[data-kg-add]')
    t.wait_for('.ip-cfadd [data-ip-cfn="label"]')
    pg.fill('.ip-cfadd [data-ip-cfn="label"]', 'Проверка: размер')
    pg.keyboard.press('Enter')
    t.ck(t.wait_until("() => [...document.querySelectorAll('.kg thead tr:first-child th')].some((th) => th.textContent.trim() === 'Проверка: размер')"),
         '«+ поле» не завело столбец своего поля')
    i_mail, i_tel = col_index(pg, 'Почта'), col_index(pg, 'Телефон')
    pg.locator('.kg thead tr:first-child th').nth(i_mail).drag_to(pg.locator('.kg thead tr:first-child th').nth(i_tel), target_position={'x': 4, 'y': 10})
    t.ck(t.wait_until("() => { const h = [...document.querySelectorAll('.kg thead tr:first-child th')].map((th) => th.textContent.trim().replace(/ [↑↓]$/, '')); return h.indexOf('Почта') < h.indexOf('Телефон'); }"),
         'перетаскивание заголовка не поставило «Почту» перед «Телефоном»')
    pg.reload()
    t.wait_for('.kg thead')
    t.ck(col_index(pg, 'Почта') < col_index(pg, 'Телефон'), 'порядок столбцов не пережил перезагрузку')
